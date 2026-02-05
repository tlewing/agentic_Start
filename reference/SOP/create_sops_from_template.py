# -*- coding: utf-8 -*-
"""
Create SOPs from GSL Master Template
Generates all 19 SOPs using the exact template structure
"""

import os
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Paths
TEMPLATE_PATH = r"C:\Users\tewing\Desktop\Claude Projects\GSL_SOP_Template.docx"
OUTPUT_FOLDER = r"C:\Users\tewing\Desktop\Claude Projects\SOP_Revisions\Revised"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

def create_sop_from_template(sop_data):
    """Create an SOP document based on the GSL template structure"""
    doc = Document(TEMPLATE_PATH)

    # Clear all existing content except the RACI table
    # We'll rebuild the document paragraphs
    paragraphs_to_keep = []
    raci_table = None

    # Keep the RACI table if it exists
    if doc.tables:
        raci_table = doc.tables[0]

    # Clear all paragraphs
    for para in doc.paragraphs:
        para.clear()

    # Now rebuild the document
    # We'll use the first paragraph and add more

    paras = doc.paragraphs
    para_idx = 0

    def get_or_add_para():
        nonlocal para_idx
        if para_idx < len(paras):
            p = paras[para_idx]
            para_idx += 1
            return p
        else:
            return doc.add_paragraph()

    def add_bold_text(para, text):
        para.clear()
        run = para.add_run(text)
        run.bold = True

    def add_normal_text(para, text):
        para.clear()
        para.add_run(text)

    def add_list_item(text):
        p = doc.add_paragraph(text, style='List Paragraph')
        return p

    # === BUILD DOCUMENT ===

    # Title
    p = get_or_add_para()
    add_bold_text(p, f"SOP: {sop_data['sop_id']} - {sop_data['title']}")

    # Department
    p = get_or_add_para()
    add_bold_text(p, f"Department: {sop_data['department']}")

    # Related SOPs
    p = get_or_add_para()
    add_bold_text(p, "Related SOPs:")

    for related in sop_data['related_sops']:
        p = get_or_add_para()
        add_normal_text(p, related)

    # Blank line
    get_or_add_para()

    # Purpose
    p = get_or_add_para()
    add_bold_text(p, "Purpose")
    p = get_or_add_para()
    add_normal_text(p, sop_data['purpose'])

    # Blank lines
    get_or_add_para()
    get_or_add_para()

    # Scope
    p = get_or_add_para()
    add_bold_text(p, "Scope")
    p = get_or_add_para()
    add_normal_text(p, "Applies to all projects / situations where:")

    for scope_item in sop_data['scope_items']:
        add_list_item(scope_item)

    # Blank lines
    doc.add_paragraph()
    doc.add_paragraph()

    # Roles & Responsibilities
    p = doc.add_paragraph()
    run = p.add_run("Roles & Responsibilities")
    run.bold = True

    for role_name, responsibilities in sop_data['roles'].items():
        p = doc.add_paragraph()
        p.add_run(role_name)
        for resp in responsibilities:
            add_list_item(resp)

    # Blank line
    doc.add_paragraph()

    # Requirements
    p = doc.add_paragraph()
    run = p.add_run("Requirements")
    run.bold = True

    p = doc.add_paragraph()
    p.add_run("Required inputs, tools, logs, templates, and references:")

    for req in sop_data['requirements']:
        add_list_item(req)

    # Blank line
    doc.add_paragraph()

    # Procedure
    p = doc.add_paragraph()
    run = p.add_run("Procedure")
    run.bold = True

    for i, (step_name, step_desc) in enumerate(sop_data['procedure'], 1):
        p = doc.add_paragraph()
        run = p.add_run(f"Step {i} - {step_name}")
        run.bold = True

        p = doc.add_paragraph()
        p.add_run(step_desc)

        doc.add_paragraph()

    # Appendix
    p = doc.add_paragraph()
    run = p.add_run("Appendix")
    run.bold = True

    for item in sop_data['appendix']:
        add_list_item(item)

    return doc


# ============================================================
# SOP DATA DEFINITIONS
# ============================================================

ALL_SOPS = [
    {
        "sop_id": "9.2.060",
        "title": "Create Issue List and Begin RFI Process",
        "department": "Project Management",
        "related_sops": [
            "9.2.040 - PM Reviews Plans, Specifications & Schedule",
            "9.2.050 - FS Reviews Plans, Specifications & Schedule",
            "9.4.205 - Manage RFIs (Requests for Information)"
        ],
        "purpose": "To establish a standardized process for documenting project issues identified during pre-construction review and initiating the Request for Information (RFI) process. This ensures timely resolution of design conflicts, ambiguities, and missing information before they impact construction activities. The RFI folder is initiated during estimating by the Estimator and handed off to the project team at turnover.",
        "scope_items": [
            "Pre-construction reviews have identified issues requiring clarification",
            "Contract document review by PM and Foreman is in progress",
            "Design conflicts, errors, or ambiguities have been discovered",
            "Scope clarification is needed from the design team"
        ],
        "roles": {
            "Estimator": [
                "Initiates/creates the project RFI folder during estimating",
                "Hands off the RFI folder to the project team at turnover"
            ],
            "Project Manager": [
                "During contract document review, creates and maintains the RFI Issue List",
                "Prioritizes issues and submits RFIs promptly to support planning",
                "Maintains the RFI log and ensures filing in the project repository",
                "Tracks RFI responses and follows up on overdue items",
                "Assesses impact of responses on schedule/cost"
            ],
            "Foreman": [
                "During contract document review, identifies constructability/field conflicts",
                "Adds field issues to the RFI Issue List",
                "Assists with clarifying and validating RFI questions with field perspective",
                "Reviews RFI responses for field impact"
            ],
            "General Superintendent": [
                "Validates issue priorities for major items",
                "Reviews significant design clarifications"
            ]
        },
        "requirements": [
            "Contract documents (plans, specs, addenda)",
            "Existing RFI folder (initiated during estimating)",
            "RFI Issue List / RFI log template",
            "Project document repository / project folder location",
            "Contract RFI procedures and submission requirements",
            "Design team contact information"
        ],
        "procedure": [
            ("Confirm RFI Folder Exists (Estimator Handoff)",
             "PM confirms the RFI folder exists from estimating. If missing, PM creates it immediately and notifies Estimator/Branch Manager. Verify folder location and access permissions are correct."),
            ("Build RFI Issue List During Contract Document Review",
             "PM and Foreman review contract documents (plans, specs, addenda). Capture discrepancies, conflicts, missing info, and constructability issues in the RFI Issue List. Include drawing/spec references for each issue. Categorize by type and assign priority."),
            ("Prioritize and Draft RFIs",
             "PM prioritizes issues affecting schedule, procurement, prefab, and layout/sequencing. Draft RFIs clearly with references and the exact clarification/decision needed. Use standard RFI format per contract requirements. Provide proposed solution if applicable."),
            ("Issue RFIs ASAP to Support Planning",
             "PM submits RFIs as soon as possible (do not wait to bundle). Route through proper channels (GC if subcontractor). Log submission date and recipient. Track due dates and status. Escalate overdue items per project controls process."),
            ("File Responses and Update the Log",
             "File RFIs, responses, and supporting exhibits in the RFI folder. Update the log with responses and impacts. Review responses for completeness. Assess impact on schedule/cost and initiate change process if needed. Distribute relevant answers to the team.")
        ],
        "appendix": [
            "Issue Log Template",
            "Standard RFI Form",
            "RFI Priority Matrix",
            "Response Follow-up Checklist"
        ]
    },
    {
        "sop_id": "9.2.070",
        "title": "Conduct Site Visit",
        "department": "Project Management / Field Operations",
        "related_sops": [
            "9.2.050 - FS Reviews Plans, Specifications & Schedule",
            "9.2.060 - Create Issue List and Begin RFI Process",
            "9.2.080 - Prepare Material Handling Plan",
            "9.3.010 - Setup Office Trailer",
            "9.3.470 - Develop Site-Specific Safety Plan"
        ],
        "purpose": "To establish a standardized process for conducting thorough site visits during the preconstruction phase to gather critical information for project planning, identify site conditions, and validate assumptions made during estimating.",
        "scope_items": [
            "Project has been awarded or LOI/NTP received",
            "On-site evaluation is required for planning",
            "Material handling and safety plans need site-specific information",
            "Mobilization requirements need to be assessed"
        ],
        "roles": {
            "Project Manager": [
                "Schedules and coordinates site visit",
                "Documents site conditions and access constraints",
                "Identifies material handling requirements",
                "Coordinates with GC/Owner for site access"
            ],
            "Foreman / Field Supervisor": [
                "Participates in site visit",
                "Evaluates constructability and work sequence",
                "Identifies safety concerns and hazards",
                "Assesses staging and laydown areas"
            ],
            "General Superintendent": [
                "Participates in site visits for major projects",
                "Validates resource and equipment needs",
                "Reviews site logistics feasibility"
            ],
            "Safety Coordinator": [
                "Identifies site-specific safety requirements",
                "Documents hazards for safety plan development"
            ]
        },
        "requirements": [
            "Contract documents (plans, specs)",
            "Site visit checklist",
            "Camera/documentation equipment",
            "PPE appropriate for site conditions",
            "Site access authorization"
        ],
        "procedure": [
            ("Schedule and Prepare for Site Visit",
             "Coordinate site access with GC/Owner. Review contract documents before visit. Prepare site visit checklist. Confirm attendees and PPE requirements."),
            ("Evaluate Site Access and Logistics",
             "Document access routes and restrictions. Identify delivery areas and constraints. Assess parking and staging availability. Note traffic patterns and restrictions."),
            ("Assess Work Areas and Conditions",
             "Walk all work areas. Document existing conditions. Identify coordination requirements with other trades. Note ceiling heights, clearances, and access constraints."),
            ("Identify Safety and Environmental Factors",
             "Document hazards and safety concerns. Identify required permits or special precautions. Note environmental conditions. Assess emergency access and egress."),
            ("Document and Distribute Findings",
             "Complete site visit report with photos. Distribute findings to project team. Update planning documents based on findings. Add issues to RFI list as needed.")
        ],
        "appendix": [
            "Site Visit Checklist",
            "Site Visit Report Template"
        ]
    },
    {
        "sop_id": "9.2.080",
        "title": "Prepare Material Handling Plan",
        "department": "Project Management / Field Operations",
        "related_sops": [
            "9.2.020 - Procurement of Large Feeder Wire",
            "9.2.070 - Conduct Site Visit",
            "9.2.100 - Prepare Layout & Sequencing Plan"
        ],
        "purpose": "To establish a standardized process for developing a comprehensive material handling plan that addresses delivery, staging, storage, and distribution of materials throughout the project.",
        "scope_items": [
            "Project site visit has been completed",
            "Material lists and delivery schedules are available",
            "Staging and laydown areas need to be planned",
            "Coordination with other trades is required"
        ],
        "roles": {
            "Project Manager": [
                "Coordinates development of material handling plan",
                "Ensures plan aligns with project schedule",
                "Communicates material delivery requirements to vendors"
            ],
            "Foreman / Field Supervisor": [
                "Identifies field material handling requirements",
                "Validates staging and storage locations",
                "Plans material distribution to work areas"
            ],
            "General Superintendent": [
                "Reviews and approves material handling plan",
                "Coordinates equipment and resource needs"
            ]
        },
        "requirements": [
            "Site visit findings and documentation",
            "Project plans and material lists",
            "Delivery schedules and lead times",
            "Site logistics plan from GC (if applicable)"
        ],
        "procedure": [
            ("Identify Material Categories and Quantities",
             "Review project material lists. Categorize materials by size, weight, and handling requirements. Identify long-lead items requiring early coordination."),
            ("Assess Site Logistics",
             "Review site visit documentation. Identify delivery access points and restrictions. Determine staging and laydown areas. Assess storage requirements."),
            ("Develop Delivery Schedule",
             "Coordinate deliveries with project schedule. Sequence deliveries to avoid congestion. Plan for just-in-time delivery where feasible."),
            ("Plan Material Distribution",
             "Map material flow from receiving to installation areas. Identify equipment needs (forklifts, carts, etc.). Plan vertical distribution for multi-story projects."),
            ("Document and Communicate Plan",
             "Complete material handling plan document. Distribute to project team and subcontractors. Update as project conditions change.")
        ],
        "appendix": [
            "Material Handling Plan Template",
            "Delivery Schedule Template"
        ]
    },
    {
        "sop_id": "9.2.090",
        "title": "Develop Labor Budget",
        "department": "Project Management",
        "related_sops": [
            "9.2.080 - Prepare Material Handling Plan",
            "9.2.110 - Develop Project Schedule"
        ],
        "purpose": "To establish a standardized process for developing an accurate labor budget that supports project planning, resource allocation, and cost control.",
        "scope_items": [
            "Project estimate with labor hours is available",
            "Material handling plan has been developed",
            "Resource planning is required for scheduling",
            "Cost control baseline needs to be established"
        ],
        "roles": {
            "Project Manager": [
                "Develops and maintains labor budget",
                "Validates labor hours against estimate",
                "Monitors labor budget throughout project"
            ],
            "General Superintendent": [
                "Reviews labor budget for accuracy",
                "Validates productivity assumptions",
                "Supports resource planning"
            ],
            "Estimator": [
                "Provides estimate labor hours and assumptions",
                "Clarifies scope and productivity factors"
            ]
        },
        "requirements": [
            "Project estimate with labor hours by phase/system",
            "Material handling plan (9.2.080)",
            "Historical productivity data",
            "Project schedule milestones"
        ],
        "procedure": [
            ("Review Estimate Labor Hours",
             "Obtain labor hour breakdown from estimating. Review productivity assumptions. Identify any scope changes since estimate."),
            ("Validate Labor Quantities",
             "Compare estimate quantities to takeoff. Adjust for field conditions identified in site visit. Factor in project-specific complexities."),
            ("Develop Phased Labor Budget",
             "Allocate labor hours to schedule phases. Plan labor loading by week/month. Identify peak manpower requirements."),
            ("Establish Tracking Metrics",
             "Set up labor tracking codes. Define reporting requirements. Establish productivity targets."),
            ("Review and Approve",
             "Present labor budget to management. Incorporate feedback and adjustments. Finalize and distribute to project team.")
        ],
        "appendix": [
            "Labor Budget Template",
            "Manpower Loading Chart Template"
        ]
    },
    {
        "sop_id": "9.2.100",
        "title": "Prepare Layout & Sequencing Plan",
        "department": "Project Management / Field Operations",
        "related_sops": [
            "9.2.090 - Develop Labor Budget",
            "9.2.110 - Develop Project Schedule"
        ],
        "purpose": "To establish a standardized process for developing a comprehensive layout and work sequencing plan that optimizes construction efficiency and coordinates with other trades.",
        "scope_items": [
            "Labor budget has been developed",
            "Project schedule development is in progress",
            "Work area breakdown is needed",
            "Trade coordination is required"
        ],
        "roles": {
            "Project Manager": [
                "Coordinates development of layout and sequencing plan",
                "Ensures plan aligns with overall project schedule",
                "Communicates sequencing requirements to team"
            ],
            "Foreman / Field Supervisor": [
                "Develops detailed work sequence",
                "Identifies installation methods and crew assignments",
                "Coordinates with other trade sequences"
            ],
            "General Superintendent": [
                "Reviews and approves layout and sequencing plan",
                "Validates resource and equipment allocation",
                "Resolves sequencing conflicts"
            ]
        },
        "requirements": [
            "Project plans and specifications",
            "Labor budget (9.2.090)",
            "Material handling plan",
            "GC master schedule and sequence",
            "Trade coordination requirements"
        ],
        "procedure": [
            ("Review Project Requirements",
             "Analyze plans for work area breakdown. Identify critical path activities. Note coordination dependencies with other trades."),
            ("Develop Work Area Layout",
             "Break project into logical work areas/zones. Map installation sequence within each area. Identify access and work flow patterns."),
            ("Sequence Installation Activities",
             "Determine optimal installation sequence. Coordinate with other trade activities. Plan for inspections and testing."),
            ("Align with Resources",
             "Match sequence to labor availability. Plan material delivery to support sequence. Identify equipment and tool requirements."),
            ("Document and Distribute Plan",
             "Complete layout and sequencing document. Create visual sequence graphics/maps. Distribute to project team and GC.")
        ],
        "appendix": [
            "Layout & Sequencing Plan Template",
            "Work Area Breakdown Template"
        ]
    },
    {
        "sop_id": "9.2.110",
        "title": "Develop Project Schedule",
        "department": "Project Management",
        "related_sops": [
            "9.2.090 - Develop Labor Budget",
            "9.2.100 - Prepare Layout & Sequencing Plan"
        ],
        "purpose": "To establish a standardized process for developing a comprehensive project schedule that supports planning, resource allocation, and progress monitoring.",
        "scope_items": [
            "Labor budget and layout/sequencing plan are complete",
            "GC master schedule is available",
            "Contract milestones have been identified",
            "Resource loading is required"
        ],
        "roles": {
            "Project Manager": [
                "Develops and maintains project schedule",
                "Coordinates schedule with GC master schedule",
                "Monitors and updates schedule throughout project"
            ],
            "Foreman / Field Supervisor": [
                "Provides input on activity durations",
                "Validates installation sequences",
                "Reports progress for schedule updates"
            ],
            "General Superintendent": [
                "Reviews and approves project schedule",
                "Validates resource loading",
                "Supports schedule conflict resolution"
            ]
        },
        "requirements": [
            "GC master schedule",
            "Labor budget (9.2.090)",
            "Layout and sequencing plan (9.2.100)",
            "Material delivery schedules",
            "Contract milestones and deadlines"
        ],
        "procedure": [
            ("Obtain and Review GC Schedule",
             "Request GC master schedule. Identify milestone dates and constraints. Note coordination requirements."),
            ("Develop Activity List",
             "Break project into schedulable activities. Define activity durations based on labor budget. Identify dependencies and constraints."),
            ("Build Schedule Logic",
             "Establish predecessor/successor relationships. Apply resource loading and leveling. Identify critical path."),
            ("Align with Master Schedule",
             "Coordinate milestones with GC schedule. Resolve conflicts and constraints. Validate overall duration."),
            ("Review and Baseline",
             "Review schedule with project team. Incorporate feedback and adjustments. Establish baseline and distribute.")
        ],
        "appendix": [
            "Project Schedule Template",
            "Milestone Checklist"
        ]
    },
    {
        "sop_id": "9.2.120",
        "title": "Establish Tracking & Control Systems",
        "department": "Project Management",
        "related_sops": [
            "9.2.110 - Develop Project Schedule",
            "9.2.130 - Construction Execution Kickoff Meeting",
            "9.2.140 - Develop Project Budget"
        ],
        "purpose": "To establish the project's tracking and control systems early so cost, schedule, labor, procurement, RFIs/submittals, and change management can be monitored and controlled throughout project execution.",
        "scope_items": [
            "Project has been awarded (LOI/NTP received)",
            "Budget and schedule baselines are being established",
            "Project controls need to be set up before kickoff",
            "Reporting cadence needs to be defined"
        ],
        "roles": {
            "Project Manager": [
                "Establishes project controls structure (cost codes, reporting, logs, filing)",
                "Ensures systems are maintained and used consistently",
                "Defines update frequency and responsible parties",
                "Communicates expectations to the team"
            ],
            "Project Engineer / Project Coordinator": [
                "Maintains logs (RFI, submittal, issue)",
                "Supports document control",
                "Updates trackers and reports"
            ],
            "General Superintendent / Field Leadership": [
                "Provides field inputs for labor productivity tracking",
                "Contributes to lookaheads and progress reporting",
                "Reports on schedule progress and constraints"
            ],
            "Branch Manager": [
                "Confirms required controls are in place",
                "Supports escalation and resource needs"
            ]
        },
        "requirements": [
            "Project budget and cost codes (from 9.2.140)",
            "Baseline project schedule (from 9.2.110)",
            "Project folder / document repository structure",
            "Standard logs/templates (RFI, submittal, issue, change, procurement logs)",
            "Reporting cadence expectations"
        ],
        "procedure": [
            ("Establish Project File Structure and Access",
             "Confirm project repository/folder structure is created and permissions assigned. Define where controls/logs will live (single source of truth). Set up standard folder hierarchy per company standards."),
            ("Set Up Cost Tracking and Reporting Cadence",
             "Confirm budget, cost codes, and cost reporting format. Define update frequency (e.g., weekly cost review) and responsible parties. Establish cost-to-complete and forecast procedures."),
            ("Set Up Schedule Tracking",
             "Confirm baseline schedule is loaded/accessible. Define schedule update cadence and required inputs. Establish lookahead schedule process."),
            ("Set Up Core Logs and Trackers",
             "Create/confirm: RFI log, submittal log, issue list, procurement log, change log. Assign owners for each log and define minimum required fields. Set up safety items tracker if used."),
            ("Establish Labor/Productivity Tracking Method",
             "Define how labor hours, production quantities, and productivity will be tracked. Align with field reporting and superintendent inputs. Set up earned value or unit tracking as applicable."),
            ("Implement Controls and Communicate Expectations",
             "Communicate the systems, cadence, and responsibilities to the project team. Confirm controls will be reviewed during kickoff (9.2.130). Document expectations and maintain through execution.")
        ],
        "appendix": [
            "RFI Log Template",
            "Submittal Log Template",
            "Issue Log Template",
            "Change Log Template",
            "Procurement Log Template",
            "Weekly Cost Report Template"
        ]
    },
    {
        "sop_id": "9.2.130",
        "title": "Construction Execution Kickoff Meeting",
        "department": "Project Management / Field Operations",
        "related_sops": [
            "9.2.100 - Prepare Layout & Sequencing Plan",
            "9.2.120 - Establish Tracking & Control Systems",
            "Prefab Plan (SOP # TBD)"
        ],
        "purpose": "To establish a standardized process for conducting a construction execution kickoff meeting that aligns the project team on plans, responsibilities, and expectations before mobilization.",
        "scope_items": [
            "All preconstruction planning activities are complete",
            "Project team is assembled and ready for mobilization",
            "Budget, schedule, and plans need team alignment",
            "Roles and responsibilities need to be confirmed"
        ],
        "roles": {
            "Project Manager": [
                "Schedules and facilitates kickoff meeting",
                "Prepares and distributes meeting agenda",
                "Documents action items and decisions"
            ],
            "Foreman / Field Supervisor": [
                "Attends and participates in kickoff meeting",
                "Presents field execution plan",
                "Confirms understanding of project requirements"
            ],
            "General Superintendent": [
                "Attends kickoff meeting",
                "Validates resource commitments",
                "Confirms support for project execution"
            ],
            "Branch Manager": [
                "Attends kickoff meeting for major projects",
                "Confirms management support",
                "Addresses resource or escalation needs"
            ]
        },
        "requirements": [
            "Project plans and specifications",
            "Layout and sequencing plan (9.2.100)",
            "Labor budget and schedule",
            "Material handling plan",
            "Safety plan",
            "Tracking and control systems (9.2.120)"
        ],
        "procedure": [
            ("Schedule and Prepare Meeting",
             "Schedule meeting with all required attendees. Prepare agenda covering all key topics. Compile and distribute meeting materials in advance."),
            ("Review Project Scope and Plans",
             "Present project overview and key requirements. Review layout and sequencing plan. Confirm scope understanding with team."),
            ("Review Budget and Schedule",
             "Present labor budget and cost targets. Review project schedule and milestones. Discuss critical path and constraints."),
            ("Review Safety and Logistics",
             "Present site-specific safety plan. Review material handling and logistics plan. Confirm safety requirements and expectations."),
            ("Assign Responsibilities and Action Items",
             "Confirm roles and responsibilities. Document action items and owners. Set follow-up meeting schedule."),
            ("Document and Distribute Meeting Notes",
             "Complete meeting minutes. Distribute to all attendees and stakeholders. Track action item completion.")
        ],
        "appendix": [
            "Kickoff Meeting Agenda Template",
            "Meeting Minutes Template"
        ]
    },
    {
        "sop_id": "9.2.140",
        "title": "Develop Project Budget",
        "department": "Project Management",
        "related_sops": [
            "9.2.010 - Team Selection",
            "Manage Change Orders (SOP # TBD)",
            "Document Filing Standards (SOP # TBD)"
        ],
        "purpose": "To establish a standardized process for developing a comprehensive project budget that supports cost control, forecasting, and financial reporting throughout the project.",
        "scope_items": [
            "Contract has been awarded",
            "Project estimate is available for budget development",
            "Cost codes and tracking structure need to be established",
            "Budget approval is required before significant expenditures"
        ],
        "roles": {
            "Project Manager": [
                "Develops and maintains project budget",
                "Establishes cost codes and tracking structure",
                "Monitors budget throughout project execution"
            ],
            "Estimator": [
                "Provides estimate breakdown and assumptions",
                "Clarifies scope and pricing elements",
                "Supports budget validation"
            ],
            "Branch Manager": [
                "Reviews and approves project budget",
                "Monitors budget performance",
                "Addresses budget variances and issues"
            ]
        },
        "requirements": [
            "Contract value and scope",
            "Project estimate with detailed breakdown",
            "Standard cost code structure",
            "Budget template and reporting requirements"
        ],
        "procedure": [
            ("Review Contract and Estimate",
             "Obtain contract documents and value. Review estimate breakdown and assumptions. Identify any scope changes since estimate."),
            ("Establish Budget Structure",
             "Set up cost codes per company standards. Allocate contract value to cost categories. Establish contingency and allowances."),
            ("Develop Detailed Budget",
             "Break down labor, material, equipment, and subcontract costs. Validate pricing against current market conditions. Incorporate project-specific requirements."),
            ("Review and Approve Budget",
             "Present budget to management for review. Incorporate feedback and adjustments. Obtain budget approval."),
            ("Establish Budget Controls",
             "Set up cost tracking and reporting. Define variance thresholds and alerts. Communicate budget expectations to team. All budget documentation must be filed per Document Filing Standards SOP (SOP # TBD).")
        ],
        "appendix": [
            "Project Budget Template",
            "Cost Code Structure Guide"
        ]
    }
]

# Mobilization SOPs
MOBILIZATION_SOPS = [
    {
        "sop_id": "9.3.050",
        "title": "Setup Temporary Water",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.3.010 - Setup Office Trailer", "9.3.020 - Setup Storage Trailer"],
        "purpose": "To establish a standardized process for setting up temporary water services at the project site to support construction activities.",
        "scope_items": ["Temporary water is required at the project site", "Site visit has identified water service needs", "Permits and approvals are in place"]
    },
    {
        "sop_id": "9.3.060",
        "title": "Setup Sanitary Facilities",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.3.010 - Setup Office Trailer", "9.3.020 - Setup Storage Trailer", "9.3.050 - Setup Temporary Water"],
        "purpose": "To establish a standardized process for setting up sanitary facilities at the project site to meet OSHA requirements and support workforce needs.",
        "scope_items": ["Sanitary facilities are required per OSHA", "Site visit has identified facility locations", "Workforce size has been determined"]
    },
    {
        "sop_id": "9.3.070",
        "title": "Setup Temporary Lighting",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.3.040 - Setup Temporary Power", "9.3.050 - Setup Temporary Water", "9.3.060 - Setup Sanitary Facilities"],
        "purpose": "To establish a standardized process for setting up temporary lighting at the project site to support safe work conditions.",
        "scope_items": ["Temporary lighting is required for work areas", "Temporary power is available", "Safety requirements have been identified"]
    },
    {
        "sop_id": "9.3.080",
        "title": "Setup Signage",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.3.030 - Setup Site Fencing and Access Control", "9.3.060 - Setup Sanitary Facilities"],
        "purpose": "To establish a standardized process for setting up required signage at the project site to meet safety, identification, and regulatory requirements.",
        "scope_items": ["Safety signage is required per OSHA", "Site identification is needed", "Regulatory requirements must be met"]
    },
    {
        "sop_id": "9.3.090",
        "title": "Setup Security",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.3.030 - Setup Site Fencing and Access Control", "9.3.080 - Setup Signage"],
        "purpose": "To establish a standardized process for setting up security measures at the project site to protect materials, equipment, and personnel.",
        "scope_items": ["Security measures are required for the site", "Site fencing and access control are in place", "Valuable materials/equipment will be stored on site"]
    },
    {
        "sop_id": "9.3.100",
        "title": "Setup Laydown Area",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.2.080 - Prepare Material Handling Plan", "9.3.020 - Setup Storage Trailer", "9.3.030 - Setup Site Fencing and Access Control"],
        "purpose": "To establish a standardized process for setting up material laydown areas at the project site to support material handling and storage.",
        "scope_items": ["Material laydown area is required", "Material handling plan has been developed", "Site logistics have been coordinated with GC"]
    },
    {
        "sop_id": "9.3.110",
        "title": "Setup Equipment Staging",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.2.080 - Prepare Material Handling Plan", "9.3.100 - Setup Laydown Area"],
        "purpose": "To establish a standardized process for setting up equipment staging areas at the project site to support construction operations.",
        "scope_items": ["Equipment staging is required", "Laydown area has been established", "Equipment needs have been identified"]
    },
    {
        "sop_id": "9.3.120",
        "title": "Setup Parking",
        "department": "Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.2.080 - Prepare Material Handling Plan", "9.3.100 - Setup Laydown Area", "9.3.110 - Setup Equipment Staging"],
        "purpose": "To establish a standardized process for setting up worker parking areas at the project site to support workforce logistics.",
        "scope_items": ["Worker parking is required", "Site visit has identified parking options", "Workforce size has been determined"]
    },
    {
        "sop_id": "9.3.470",
        "title": "Develop Site-Specific Safety Plan",
        "department": "Safety / Field Operations",
        "related_sops": ["9.2.070 - Conduct Site Visit", "9.3.010 - Setup Office Trailer", "9.3.060 - Setup Sanitary Facilities", "Document Filing Standards (SOP # TBD)"],
        "purpose": "To establish a standardized process for developing a comprehensive site-specific safety plan that addresses hazards, requirements, and emergency procedures for the project.",
        "scope_items": ["Site visit has been completed", "Site hazards have been identified", "Safety requirements need to be documented", "Emergency procedures need to be established"]
    },
    {
        "sop_id": "9.4.010",
        "title": "Conduct Coordination Meetings",
        "department": "Project Management",
        "related_sops": ["Conduct Subcontractor Coordination Meetings (SOP # TBD)", "Conduct Owner/GC Coordination Meetings (SOP # TBD)", "Conduct Progress Meetings (SOP # TBD)", "Document Filing Standards (SOP # TBD)"],
        "purpose": "To establish a standardized process for conducting effective coordination meetings that support project communication and issue resolution.",
        "scope_items": ["Project requires regular coordination meetings", "Multiple parties need to be aligned", "Issues need to be tracked and resolved", "Communication needs to be documented"]
    }
]

# Add standard content for mobilization SOPs
def add_standard_mobilization_content(sop):
    """Add standard roles, requirements, procedure, and appendix for mobilization SOPs"""
    sop["roles"] = {
        "Project Manager": [
            "Coordinates setup requirements with GC/Owner",
            "Ensures resources are available for setup",
            "Monitors setup progress and completion"
        ],
        "Foreman / Field Supervisor": [
            "Oversees physical setup activities",
            "Ensures setup meets project requirements",
            "Reports completion and any issues"
        ],
        "General Superintendent": [
            "Validates setup requirements",
            "Supports resource allocation",
            "Reviews completion for compliance"
        ]
    }
    sop["requirements"] = [
        "Site visit documentation",
        "Project requirements and specifications",
        "Necessary permits and approvals",
        "Required materials and equipment"
    ]
    sop["procedure"] = [
        ("Review Requirements",
         "Review site visit findings related to this setup. Confirm requirements with GC/Owner. Verify permits and approvals are in place."),
        ("Plan Setup Activities",
         "Identify resources and materials needed. Coordinate with other setup activities. Schedule setup work."),
        ("Execute Setup",
         "Perform setup work per plan. Document setup completion. Address any issues encountered."),
        ("Verify and Document",
         "Verify setup meets requirements. Complete documentation. Communicate completion to project team.")
    ]
    sop["appendix"] = [
        "Setup Checklist",
        "Completion Documentation Form"
    ]
    return sop


def main():
    print("=" * 70)
    print("CREATING SOPs FROM GSL MASTER TEMPLATE")
    print("=" * 70)

    all_sops = ALL_SOPS + [add_standard_mobilization_content(sop) for sop in MOBILIZATION_SOPS]

    for sop_data in all_sops:
        filename = f"{sop_data['sop_id']} - {sop_data['title']}.docx"
        print(f"  Creating: {filename}")

        doc = create_sop_from_template(sop_data)
        doc.save(os.path.join(OUTPUT_FOLDER, filename))

    print("\n" + "=" * 70)
    print(f"CREATED {len(all_sops)} SOPs in {OUTPUT_FOLDER}")
    print("=" * 70)

    print("\nFiles created:")
    for f in sorted(os.listdir(OUTPUT_FOLDER)):
        if f.endswith('.docx'):
            print(f"  - {f}")


if __name__ == "__main__":
    main()
