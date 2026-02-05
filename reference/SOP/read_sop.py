import docx
import os
import glob

def read_sop(pattern):
    # Find file matching pattern
    folder = r"C:\Users\tewing\Desktop\Claude Projects\SOP_Revisions"
    files = [f for f in os.listdir(folder) if pattern in f and f.endswith('.docx')]

    if not files:
        print(f"No file found matching: {pattern}")
        return

    filepath = os.path.join(folder, files[0])
    print(f"=== Reading: {files[0]} ===\n")

    doc = docx.Document(filepath)

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            style = para.style.name if para.style else "None"
            print(f"[{i}] ({style}) {text[:150]}{'...' if len(text) > 150 else ''}")

    # Check for tables
    if doc.tables:
        print(f"\n=== Tables: {len(doc.tables)} ===")
        for t_idx, table in enumerate(doc.tables):
            print(f"\nTable {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols")
            for r_idx, row in enumerate(table.rows[:5]):  # First 5 rows
                cells = [cell.text[:30] for cell in row.cells]
                print(f"  Row {r_idx}: {cells}")

if __name__ == "__main__":
    import sys
    pattern = sys.argv[1] if len(sys.argv) > 1 else "9.2.060"
    read_sop(pattern)
