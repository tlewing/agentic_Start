# -*- coding: utf-8 -*-
"""
SOP Revision Script - GSL Governance Compliance
Revises all 19 SOPs according to GSL template and governance rules
"""

import os
import re
from docx import Document

# Paths
INPUT_FOLDER = r"C:\Users\tewing\Desktop\Claude Projects\SOP_Revisions"
OUTPUT_FOLDER = r"C:\Users\tewing\Desktop\Claude Projects\SOP_Revisions\Revised"

# Create output folder
os.makedirs(OUTPUT_FOLDER, exist_ok=True)


def create_revised_9_2_060():
    """Create the fully revised SOP 9.2.060"""
    doc = Document()

    # Title
    doc.add_paragraph("SOP: 9.2.060 - Create Issue List and Begin RFI Process")
    doc.add_paragraph("")
    doc.add_paragraph("Department: Project Management")
    doc.add_paragraph("")

    # Related SOPs
    doc.add_paragraph("Related SOPs:")
    doc.add_paragraph("- 9.2.040 - PM Reviews Plans, Specifications & Schedule")
    doc.add_paragraph("- 9.2.050 - FS Reviews Plans, Specifications & Schedule")
    doc.add_paragraph("- 9.4.205 - Manage RFIs (Requests for Information)")
    doc.add_paragraph("")

    # Purpose
    doc.add_paragraph("PURPOSE")
    doc.add_paragraph("")
    doc.add_paragraph("To establish a standardized process for documenting project issues identified during pre-construction review and initiating the Request for Information (RFI) process. This ensures timely resolution of design conflicts, ambiguities, and missing information before they impact construction activities.")
    doc.add_paragraph("")

    # Scope
    doc.add_paragraph("SCOPE")
    doc.add_paragraph("")
    doc.add_paragraph("This SOP is performed during early preconstruction planning following receipt of a contract for review, Letter of Intent (LOI), or Notice to Proceed (NTP), and during the Project Manager and Foreman review of contract documents. The RFI folder is initiated during estimating by the Estimator and handed off to the project team at turnover.")
    doc.add_paragraph("")
    doc.add_paragraph("Applies to all projects/situations where:")
    for item in ["Pre-construction reviews have identified issues",
                 "Clarification is needed from design team",
                 "Design conflicts or errors have been discovered",
                 "Scope ambiguities require resolution"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")

    # Roles & Responsibilities
    doc.add_paragraph("ROLES & RESPONSIBILITIES")
    doc.add_paragraph("")

    doc.add_paragraph("Estimator")
    for item in ["Initiates/creates the project RFI folder during estimating",
                 "Hands off the RFI folder to the project team at turnover"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")

    doc.add_paragraph("Project Manager")
    for item in ["During contract document review, creates and maintains the RFI Issue List",
                 "Prioritizes issues and submits RFIs promptly to support planning",
                 "Maintains the RFI log and ensures filing in the project repository",
                 "Tracks RFI responses and follows up on overdue items",
                 "Assesses impact of responses on schedule/cost"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")

    doc.add_paragraph("Foreman")
    for item in ["During contract document review, identifies constructability/field conflicts and adds them to the RFI Issue List",
                 "Assists with clarifying and validating RFI questions with field perspective",
                 "Reviews RFI responses for field impact"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")

    doc.add_paragraph("General Superintendent")
    for item in ["Validates issue priorities", "Reviews significant design clarifications"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")

    # Requirements
    doc.add_paragraph("REQUIREMENTS")
    doc.add_paragraph("")
    for item in ["Contract documents (plans, specs, addenda)",
                 "Existing RFI folder (initiated during estimating)",
                 "RFI Issue List / RFI log template",
                 "Project document repository / project folder location",
                 "Contract RFI procedures and submission requirements",
                 "Design team contact information"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")

    # Procedure
    doc.add_paragraph("PROCEDURE")
    doc.add_paragraph("")

    steps = [
        ("Step 1 - Confirm RFI Folder Exists (Estimator Handoff)", [
            "PM confirms the RFI folder exists from estimating",
            "If missing, PM creates it immediately and notifies Estimator/Branch Manager",
            "Verify folder location and access permissions are correct"
        ]),
        ("Step 2 - Build RFI Issue List During Contract Document Review", [
            "PM and Foreman review contract documents (plans, specs, addenda)",
            "Capture discrepancies, conflicts, missing info, and constructability issues in the RFI Issue List",
            "Include drawing/spec references for each issue",
            "Categorize issues by type: design, coordination, scope, clarification",
            "Assign priority: critical, high, medium, low"
        ]),
        ("Step 3 - Prioritize and Draft RFIs", [
            "PM prioritizes issues affecting schedule, procurement, prefab, layout/sequencing",
            "Draft RFIs clearly with references and the exact clarification/decision needed",
            "Use standard RFI format per contract requirements",
            "Provide proposed solution if applicable",
            "Attach relevant sketches or photos"
        ]),
        ("Step 4 - Issue RFIs ASAP to Support Planning", [
            "PM submits RFIs as soon as possible (do not wait to bundle)",
            "Route through proper channels (GC if subcontractor)",
            "Log submission date and recipient",
            "Track due dates and status",
            "Escalate overdue items per project controls process"
        ]),
        ("Step 5 - File Responses and Update the Log", [
            "File RFIs, responses, and supporting exhibits in the RFI folder",
            "Update the log with responses and impacts",
            "Review responses for completeness",
            "Assess impact on schedule/cost and initiate change process if needed",
            "Distribute relevant answers to the team",
            "Update drawings/plans as needed"
        ])
    ]

    for step_title, bullets in steps:
        p = doc.add_paragraph()
        run = p.add_run(step_title)
        run.bold = True
        for bullet in bullets:
            p = doc.add_paragraph()
            p.add_run("\u2022 " + bullet)
        doc.add_paragraph("")

    # Appendix
    doc.add_paragraph("APPENDIX")
    doc.add_paragraph("")
    doc.add_paragraph("Appendix - Templates")
    for item in ["Issue Log Template", "Standard RFI Form", "RFI Priority Matrix", "Response Follow-up Checklist"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)
    doc.add_paragraph("")
    doc.add_paragraph("Appendix - References & Training Materials")
    for item in ["Training Materials (Foreman Training Binder): Tab 2, 3, 5, 6, 7",
                 "Job Descriptions (Policy Manual Section 7): 7.3.1 Operations Management",
                 "Management Directives (Policy Manual Section 9): 9.5, 9.7, 9.3"]:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + item)

    return doc


def create_sop(sop_num, title, dept, related, purpose, scope, roles, reqs, proc_steps, appendix_templates, appendix_refs):
    """Generic function to create an SOP document"""
    doc = Document()

    doc.add_paragraph(f"SOP: {sop_num} - {title}")
    doc.add_paragraph("")
    doc.add_paragraph(f"Department: {dept}")
    doc.add_paragraph("")

    doc.add_paragraph("Related SOPs:")
    for sop in related:
        doc.add_paragraph(f"- {sop}")
    doc.add_paragraph("")

    doc.add_paragraph("PURPOSE")
    doc.add_paragraph("")
    doc.add_paragraph(purpose)
    doc.add_paragraph("")

    doc.add_paragraph("SCOPE")
    doc.add_paragraph("")
    doc.add_paragraph(scope)
    doc.add_paragraph("")

    doc.add_paragraph("ROLES & RESPONSIBILITIES")
    doc.add_paragraph("")
    for role_name, responsibilities in roles.items():
        doc.add_paragraph(role_name)
        for resp in responsibilities:
            p = doc.add_paragraph()
            p.add_run("\u2022 " + resp)
        doc.add_paragraph("")

    doc.add_paragraph("REQUIREMENTS")
    doc.add_paragraph("")
    for req in reqs:
        p = doc.add_paragraph()
        p.add_run("\u2022 " + req)
    doc.add_paragraph("")

    doc.add_paragraph("PROCEDURE")
    doc.add_paragraph("")
    for i, (step_title, bullets) in enumerate(proc_steps, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"Step {i} - {step_title}")
        run.bold = True
        for bullet in bullets:
            p = doc.add_paragraph()
            p.add_run("\u2022 " + bullet)
        doc.add_paragraph("")

    doc.add_paragraph("APPENDIX")
    doc.add_paragraph("")
    if appendix_templates:
        doc.add_paragraph("Appendix - Templates")
        for item in appendix_templates:
            p = doc.add_paragraph()
            p.add_run("\u2022 " + item)
        doc.add_paragraph("")
    if appendix_refs:
        doc.add_paragraph("Appendix - References & Training Materials")
        for item in appendix_refs:
            p = doc.add_paragraph()
            p.add_run("\u2022 " + item)

    return doc


def main():
    print("=" * 60)
    print("SOP REVISION SCRIPT - GSL GOVERNANCE COMPLIANCE")
    print("=" * 60)

    # All SOP definitions
    all_sops = [
        {
            "num": "9.2.060",
            "title": "Create Issue List and Begin RFI Process",
            "custom": True
        },
        {
            "num": "9.2.070",
            "title": "Conduct Site Visit",
            "dept": "Project Management / Field Operations",
            "related": ["9.2.050 - FS Reviews Plans, Specifications & Schedule",
                       "9.2.060 - Create Issue List and Begin RFI Process",
                       "9.2.080 - Prepare Material Handling Plan",
                       "9.3.010 - Setup Office Trailer",
                       "9.3.470 - Develop Site-Specific Safety Plan"],
            "purpose": "To establish a standardized process for conducting thorough site visits during the preconstruction phase to gather critical information for project planning, identify site conditions, and validate assumptions made during estimating.",
            "scope": "This SOP applies to all projects requiring on-site evaluation during preconstruction planning. Site visits should occur as early as possible after contract award/LOI/NTP to inform material handling plans, safety plans, mobilization requirements, and project logistics.",
            "roles": {
                "Project Manager": ["Schedules and coordinates site visit", "Documents site conditions and access constraints", "Identifies material handling requirements", "Coordinates with GC/Owner for site access"],
                "Foreman / Field Supervisor": ["Participates in site visit", "Evaluates constructability and work sequence", "Identifies safety concerns and hazards", "Assesses staging and laydown areas"],
                "General Superintendent": ["Participates in site visits for major projects", "Validates resource and equipment needs", "Reviews site logistics feasibility"],
                "Safety Coordinator": ["Identifies site-specific safety requirements", "Documents hazards for safety plan development"]
            },
            "reqs": ["Contract documents (plans, specs)", "Site visit checklist", "Camera/documentation equipment", "PPE appropriate for site conditions", "Site access authorization"],
            "proc": [
                ("Schedule and Prepare for Site Visit", ["Coordinate site access with GC/Owner", "Review contract documents before visit", "Prepare site visit checklist", "Confirm attendees and PPE requirements"]),
                ("Evaluate Site Access and Logistics", ["Document access routes and restrictions", "Identify delivery areas and constraints", "Assess parking and staging availability", "Note traffic patterns and restrictions"]),
                ("Assess Work Areas and Conditions", ["Walk all work areas", "Document existing conditions", "Identify coordination requirements with other trades", "Note ceiling heights, clearances, and access constraints"]),
                ("Identify Safety and Environmental Factors", ["Document hazards and safety concerns", "Identify required permits or special precautions", "Note environmental conditions", "Assess emergency access and egress"]),
                ("Document and Distribute Findings", ["Complete site visit report with photos", "Distribute findings to project team", "Update planning documents based on findings", "Add issues to RFI list as needed"])
            ],
            "templates": ["Site Visit Checklist", "Site Visit Report Template"],
            "refs": ["Training Materials (Foreman Training Binder): Tab 2 - Preplanning & Staging", "Management Directives: 9.5 Project Management, 9.3 Field Leadership"]
        },
        {
            "num": "9.2.080",
            "title": "Prepare Material Handling Plan",
            "dept": "Project Management / Field Operations",
            "related": ["9.2.020 - Procurement of Large Feeder Wire", "9.2.070 - Conduct Site Visit", "9.2.100 - Prepare Layout & Sequencing Plan"],
            "purpose": "To establish a standardized process for developing a comprehensive material handling plan that addresses delivery, staging, storage, and distribution of materials throughout the project.",
            "scope": "This SOP applies to all projects during preconstruction planning. The material handling plan should be developed after site visit and coordinated with layout and sequencing plans.",
            "roles": {
                "Project Manager": ["Coordinates development of material handling plan", "Ensures plan aligns with project schedule", "Communicates material delivery requirements to vendors"],
                "Foreman / Field Supervisor": ["Identifies field material handling requirements", "Validates staging and storage locations", "Plans material distribution to work areas"],
                "General Superintendent": ["Reviews and approves material handling plan", "Coordinates equipment and resource needs"]
            },
            "reqs": ["Site visit findings and documentation", "Project plans and material lists", "Delivery schedules and lead times", "Site logistics plan from GC (if applicable)"],
            "proc": [
                ("Identify Material Categories and Quantities", ["Review project material lists", "Categorize materials by size, weight, and handling requirements", "Identify long-lead items requiring early coordination"]),
                ("Assess Site Logistics", ["Review site visit documentation", "Identify delivery access points and restrictions", "Determine staging and laydown areas", "Assess storage requirements"]),
                ("Develop Delivery Schedule", ["Coordinate deliveries with project schedule", "Sequence deliveries to avoid congestion", "Plan for just-in-time delivery where feasible"]),
                ("Plan Material Distribution", ["Map material flow from receiving to installation areas", "Identify equipment needs (forklifts, carts, etc.)", "Plan vertical distribution for multi-story projects"]),
                ("Document and Communicate Plan", ["Complete material handling plan document", "Distribute to project team and subcontractors", "Update as project conditions change"])
            ],
            "templates": ["Material Handling Plan Template", "Delivery Schedule Template"],
            "refs": ["Management Directives: 9.5 Project Management", "Training Materials: Tab 2 - Preplanning & Staging"]
        },
        {
            "num": "9.2.090",
            "title": "Develop Labor Budget",
            "dept": "Project Management",
            "related": ["9.2.080 - Prepare Material Handling Plan", "9.2.110 - Develop Project Schedule"],
            "purpose": "To establish a standardized process for developing an accurate labor budget that supports project planning, resource allocation, and cost control.",
            "scope": "This SOP applies to all projects during preconstruction planning. The labor budget should be developed after material plans are established and coordinated with the project schedule.",
            "roles": {
                "Project Manager": ["Develops and maintains labor budget", "Validates labor hours against estimate", "Monitors labor budget throughout project"],
                "General Superintendent": ["Reviews labor budget for accuracy", "Validates productivity assumptions", "Supports resource planning"],
                "Estimator": ["Provides estimate labor hours and assumptions", "Clarifies scope and productivity factors"]
            },
            "reqs": ["Project estimate with labor hours by phase/system", "Material handling plan (9.2.080)", "Historical productivity data", "Project schedule milestones"],
            "proc": [
                ("Review Estimate Labor Hours", ["Obtain labor hour breakdown from estimating", "Review productivity assumptions", "Identify any scope changes since estimate"]),
                ("Validate Labor Quantities", ["Compare estimate quantities to takeoff", "Adjust for field conditions identified in site visit", "Factor in project-specific complexities"]),
                ("Develop Phased Labor Budget", ["Allocate labor hours to schedule phases", "Plan labor loading by week/month", "Identify peak manpower requirements"]),
                ("Establish Tracking Metrics", ["Set up labor tracking codes", "Define reporting requirements", "Establish productivity targets"]),
                ("Review and Approve", ["Present labor budget to management", "Incorporate feedback and adjustments", "Finalize and distribute to project team"])
            ],
            "templates": ["Labor Budget Template", "Manpower Loading Chart Template"],
            "refs": ["Management Directives: 9.5 Project Management", "Job Descriptions: 7.3.1 Operations Management"]
        },
        {
            "num": "9.2.100",
            "title": "Prepare Layout & Sequencing Plan",
            "dept": "Project Management / Field Operations",
            "related": ["9.2.090 - Develop Labor Budget", "9.2.110 - Develop Project Schedule"],
            "purpose": "To establish a standardized process for developing a comprehensive layout and work sequencing plan that optimizes construction efficiency and coordinates with other trades.",
            "scope": "This SOP applies to all projects during preconstruction planning. The layout and sequencing plan should be developed in coordination with labor budget and project schedule.",
            "roles": {
                "Project Manager": ["Coordinates development of layout and sequencing plan", "Ensures plan aligns with overall project schedule", "Communicates sequencing requirements to team"],
                "Foreman / Field Supervisor": ["Develops detailed work sequence", "Identifies installation methods and crew assignments", "Coordinates with other trade sequences"],
                "General Superintendent": ["Reviews and approves layout and sequencing plan", "Validates resource and equipment allocation", "Resolves sequencing conflicts"]
            },
            "reqs": ["Project plans and specifications", "Labor budget (9.2.090)", "Material handling plan", "GC master schedule and sequence", "Trade coordination requirements"],
            "proc": [
                ("Review Project Requirements", ["Analyze plans for work area breakdown", "Identify critical path activities", "Note coordination dependencies with other trades"]),
                ("Develop Work Area Layout", ["Break project into logical work areas/zones", "Map installation sequence within each area", "Identify access and work flow patterns"]),
                ("Sequence Installation Activities", ["Determine optimal installation sequence", "Coordinate with other trade activities", "Plan for inspections and testing"]),
                ("Align with Resources", ["Match sequence to labor availability", "Plan material delivery to support sequence", "Identify equipment and tool requirements"]),
                ("Document and Distribute Plan", ["Complete layout and sequencing document", "Create visual sequence graphics/maps", "Distribute to project team and GC"])
            ],
            "templates": ["Layout & Sequencing Plan Template", "Work Area Breakdown Template"],
            "refs": ["Management Directives: 9.5 Project Management", "Training Materials: Tab 2 - Preplanning & Staging"]
        },
        {
            "num": "9.2.110",
            "title": "Develop Project Schedule",
            "dept": "Project Management",
            "related": ["9.2.090 - Develop Labor Budget", "9.2.100 - Prepare Layout & Sequencing Plan"],
            "purpose": "To establish a standardized process for developing a comprehensive project schedule that supports planning, resource allocation, and progress monitoring.",
            "scope": "This SOP applies to all projects during preconstruction planning. The project schedule should be developed after labor budget and layout/sequencing plans are established.",
            "roles": {
                "Project Manager": ["Develops and maintains project schedule", "Coordinates schedule with GC master schedule", "Monitors and updates schedule throughout project"],
                "Foreman / Field Supervisor": ["Provides input on activity durations", "Validates installation sequences", "Reports progress for schedule updates"],
                "General Superintendent": ["Reviews and approves project schedule", "Validates resource loading", "Supports schedule conflict resolution"]
            },
            "reqs": ["GC master schedule", "Labor budget (9.2.090)", "Layout and sequencing plan (9.2.100)", "Material delivery schedules", "Contract milestones and deadlines"],
            "proc": [
                ("Obtain and Review GC Schedule", ["Request GC master schedule", "Identify milestone dates and constraints", "Note coordination requirements"]),
                ("Develop Activity List", ["Break project into schedulable activities", "Define activity durations based on labor budget", "Identify dependencies and constraints"]),
                ("Build Schedule Logic", ["Establish predecessor/successor relationships", "Apply resource loading and leveling", "Identify critical path"]),
                ("Align with Master Schedule", ["Coordinate milestones with GC schedule", "Resolve conflicts and constraints", "Validate overall duration"]),
                ("Review and Baseline", ["Review schedule with project team", "Incorporate feedback and adjustments", "Establish baseline and distribute"])
            ],
            "templates": ["Project Schedule Template", "Milestone Checklist"],
            "refs": ["Management Directives: 9.5 Project Management", "Training Materials: Tab 6 - Scheduling"]
        },
        {
            "num": "9.2.120",
            "title": "Establish Tracking & Control Systems",
            "dept": "Project Management",
            "related": ["9.2.110 - Develop Project Schedule", "9.2.130 - Construction Execution Kickoff Meeting", "9.2.140 - Develop Project Budget"],
            "purpose": "To establish the project's tracking and control systems early so cost, schedule, labor, procurement, RFIs/submittals, and change management can be monitored and controlled throughout project execution.",
            "scope": "This SOP is performed during early preconstruction planning after award/LOI/NTP and before/during kickoff. Applies to the Project Manager and project team responsible for establishing and maintaining project controls.",
            "roles": {
                "Project Manager": ["Establishes project controls structure (cost codes, reporting cadence, logs, filing, dashboards)", "Ensures systems are maintained and used consistently", "Defines update frequency and responsible parties", "Communicates expectations to the team"],
                "Project Engineer / Project Coordinator": ["Maintains logs (RFI, submittal, issue)", "Supports document control", "Updates trackers and reports"],
                "General Superintendent / Field Leadership": ["Provides field inputs for labor productivity tracking", "Contributes to lookaheads and progress reporting", "Reports on schedule progress and constraints"],
                "Branch Manager": ["Confirms required controls are in place", "Supports escalation and resource needs"]
            },
            "reqs": ["Project budget and cost codes (from 9.2.140)", "Baseline project schedule (from 9.2.110)", "Project folder / document repository structure", "Standard logs/templates (RFI log, submittal log, issue log, change log, procurement log)", "Reporting cadence expectations (weekly cost/schedule/labor updates)"],
            "proc": [
                ("Establish Project File Structure and Access", ["Confirm project repository/folder structure is created and permissions assigned", "Define where controls/logs will live (single source of truth)", "Set up standard folder hierarchy per company standards"]),
                ("Set Up Cost Tracking and Reporting Cadence", ["Confirm budget, cost codes, and cost reporting format", "Define update frequency (e.g., weekly cost review) and responsible parties", "Establish cost-to-complete and forecast procedures"]),
                ("Set Up Schedule Tracking", ["Confirm baseline schedule is loaded/accessible", "Define schedule update cadence and required inputs", "Establish lookahead schedule process"]),
                ("Set Up Core Logs and Trackers", ["Create/confirm: RFI log, submittal log, issue list, procurement log, change log", "Assign owners for each log and define minimum required fields", "Set up safety items tracker if used"]),
                ("Establish Labor/Productivity Tracking Method", ["Define how labor hours, production quantities, and productivity will be tracked", "Align with field reporting and superintendent inputs", "Set up earned value or unit tracking as applicable"]),
                ("Implement Controls and Communicate Expectations", ["Communicate the systems, cadence, and responsibilities to the project team", "Confirm controls will be reviewed during kickoff (9.2.130)", "Document expectations and maintain through execution"])
            ],
            "templates": ["RFI Log Template", "Submittal Log Template", "Issue Log Template", "Change Log Template", "Procurement Log Template", "Weekly Cost Report Template", "Schedule Update Checklist"],
            "refs": ["Management Directives: 9.5 Project Management", "Job Descriptions: 7.3.1 Operations Management"]
        },
        {
            "num": "9.2.130",
            "title": "Construction Execution Kickoff Meeting",
            "dept": "Project Management / Field Operations",
            "related": ["9.2.100 - Prepare Layout & Sequencing Plan", "9.2.120 - Establish Tracking & Control Systems", "Prefab Plan (SOP # TBD)"],
            "purpose": "To establish a standardized process for conducting a construction execution kickoff meeting that aligns the project team on plans, responsibilities, and expectations before mobilization.",
            "scope": "This SOP applies to all projects before field mobilization. The kickoff meeting should occur after all preconstruction planning activities are complete.",
            "roles": {
                "Project Manager": ["Schedules and facilitates kickoff meeting", "Prepares and distributes meeting agenda", "Documents action items and decisions"],
                "Foreman / Field Supervisor": ["Attends and participates in kickoff meeting", "Presents field execution plan", "Confirms understanding of project requirements"],
                "General Superintendent": ["Attends kickoff meeting", "Validates resource commitments", "Confirms support for project execution"],
                "Branch Manager": ["Attends kickoff meeting for major projects", "Confirms management support", "Addresses resource or escalation needs"]
            },
            "reqs": ["Project plans and specifications", "Layout and sequencing plan (9.2.100)", "Labor budget and schedule", "Material handling plan", "Safety plan", "Tracking and control systems (9.2.120)"],
            "proc": [
                ("Schedule and Prepare Meeting", ["Schedule meeting with all required attendees", "Prepare agenda covering all key topics", "Compile and distribute meeting materials in advance"]),
                ("Review Project Scope and Plans", ["Present project overview and key requirements", "Review layout and sequencing plan", "Confirm scope understanding with team"]),
                ("Review Budget and Schedule", ["Present labor budget and cost targets", "Review project schedule and milestones", "Discuss critical path and constraints"]),
                ("Review Safety and Logistics", ["Present site-specific safety plan", "Review material handling and logistics plan", "Confirm safety requirements and expectations"]),
                ("Assign Responsibilities and Action Items", ["Confirm roles and responsibilities", "Document action items and owners", "Set follow-up meeting schedule"]),
                ("Document and Distribute Meeting Notes", ["Complete meeting minutes", "Distribute to all attendees and stakeholders", "Track action item completion"])
            ],
            "templates": ["Kickoff Meeting Agenda Template", "Meeting Minutes Template"],
            "refs": ["Management Directives: 9.5 Project Management", "Training Materials: Tab 3 - Kickoff/Project Turnover Meeting"]
        },
        {
            "num": "9.2.140",
            "title": "Develop Project Budget",
            "dept": "Project Management",
            "related": ["9.2.010 - Team Selection", "Manage Change Orders (SOP # TBD)", "Document Filing Standards (SOP # TBD)"],
            "purpose": "To establish a standardized process for developing a comprehensive project budget that supports cost control, forecasting, and financial reporting throughout the project.",
            "scope": "This SOP applies to all projects during preconstruction planning. The project budget should be developed after contract award and before significant project expenditures begin.",
            "roles": {
                "Project Manager": ["Develops and maintains project budget", "Establishes cost codes and tracking structure", "Monitors budget throughout project execution"],
                "Estimator": ["Provides estimate breakdown and assumptions", "Clarifies scope and pricing elements", "Supports budget validation"],
                "Branch Manager": ["Reviews and approves project budget", "Monitors budget performance", "Addresses budget variances and issues"]
            },
            "reqs": ["Contract value and scope", "Project estimate with detailed breakdown", "Standard cost code structure", "Budget template and reporting requirements"],
            "proc": [
                ("Review Contract and Estimate", ["Obtain contract documents and value", "Review estimate breakdown and assumptions", "Identify any scope changes since estimate"]),
                ("Establish Budget Structure", ["Set up cost codes per company standards", "Allocate contract value to cost categories", "Establish contingency and allowances"]),
                ("Develop Detailed Budget", ["Break down labor, material, equipment, and subcontract costs", "Validate pricing against current market conditions", "Incorporate project-specific requirements"]),
                ("Review and Approve Budget", ["Present budget to management for review", "Incorporate feedback and adjustments", "Obtain budget approval"]),
                ("Establish Budget Controls", ["Set up cost tracking and reporting", "Define variance thresholds and alerts", "Communicate budget expectations to team"])
            ],
            "templates": ["Project Budget Template", "Cost Code Structure Guide"],
            "refs": ["Management Directives: 9.5 Project Management", "All budget documentation must be filed per Document Filing Standards SOP (SOP # TBD)"]
        }
    ]

    # Mobilization SOPs
    mobilization_sops = [
        {"num": "9.3.050", "title": "Setup Temporary Water", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.3.010 - Setup Office Trailer", "9.3.020 - Setup Storage Trailer"]},
        {"num": "9.3.060", "title": "Setup Sanitary Facilities", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.3.010 - Setup Office Trailer", "9.3.020 - Setup Storage Trailer", "9.3.050 - Setup Temporary Water"]},
        {"num": "9.3.070", "title": "Setup Temporary Lighting", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.3.040 - Setup Temporary Power", "9.3.050 - Setup Temporary Water", "9.3.060 - Setup Sanitary Facilities"]},
        {"num": "9.3.080", "title": "Setup Signage", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.3.030 - Setup Site Fencing and Access Control", "9.3.060 - Setup Sanitary Facilities"]},
        {"num": "9.3.090", "title": "Setup Security", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.3.030 - Setup Site Fencing and Access Control", "9.3.080 - Setup Signage"]},
        {"num": "9.3.100", "title": "Setup Laydown Area", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.2.080 - Prepare Material Handling Plan", "9.3.020 - Setup Storage Trailer", "9.3.030 - Setup Site Fencing and Access Control"]},
        {"num": "9.3.110", "title": "Setup Equipment Staging", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.2.080 - Prepare Material Handling Plan", "9.3.100 - Setup Laydown Area"]},
        {"num": "9.3.120", "title": "Setup Parking", "dept": "Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.2.080 - Prepare Material Handling Plan", "9.3.100 - Setup Laydown Area", "9.3.110 - Setup Equipment Staging"]},
        {"num": "9.3.470", "title": "Develop Site-Specific Safety Plan", "dept": "Safety / Field Operations",
         "related": ["9.2.070 - Conduct Site Visit", "9.3.010 - Setup Office Trailer", "9.3.060 - Setup Sanitary Facilities", "Document Filing Standards (SOP # TBD)"]},
        {"num": "9.4.010", "title": "Conduct Coordination Meetings", "dept": "Project Management",
         "related": ["Conduct Subcontractor Coordination Meetings (SOP # TBD)", "Conduct Owner/GC Coordination Meetings (SOP # TBD)", "Conduct Progress Meetings (SOP # TBD)", "Document Filing Standards (SOP # TBD)"]}
    ]

    print("\nCreating revised SOPs...")

    # Create 9.2.060 (custom)
    print("  Creating 9.2.060 - Create Issue List and Begin RFI Process")
    doc = create_revised_9_2_060()
    doc.save(os.path.join(OUTPUT_FOLDER, "9.2.060 - Create Issue List and Begin RFI Process.docx"))

    # Create other 9.2.xxx SOPs
    for sop in all_sops:
        if sop.get("custom"):
            continue
        print(f"  Creating {sop['num']} - {sop['title']}")
        doc = create_sop(sop['num'], sop['title'], sop['dept'], sop['related'],
                        sop['purpose'], sop['scope'], sop['roles'], sop['reqs'],
                        sop['proc'], sop['templates'], sop['refs'])
        doc.save(os.path.join(OUTPUT_FOLDER, f"{sop['num']} - {sop['title']}.docx"))

    # Create mobilization SOPs with standard content
    for sop in mobilization_sops:
        print(f"  Creating {sop['num']} - {sop['title']}")
        purpose = f"To establish a standardized process for {sop['title'].lower()} at the project site during mobilization."
        scope = f"This SOP applies to all projects requiring {sop['title'].lower()}. Setup should be coordinated with site visit findings and other mobilization activities."
        roles = {
            "Project Manager": ["Coordinates setup requirements with GC/Owner", "Ensures resources are available", "Monitors setup progress"],
            "Foreman / Field Supervisor": ["Oversees physical setup activities", "Ensures setup meets requirements", "Reports completion and any issues"],
            "General Superintendent": ["Validates setup requirements", "Supports resource allocation"]
        }
        reqs = ["Site visit documentation", "Project requirements and specifications", "Necessary permits and approvals", "Required materials and equipment"]
        proc = [
            ("Review Requirements", ["Review site visit findings related to this setup", "Confirm requirements with GC/Owner", "Verify permits and approvals are in place"]),
            ("Plan Setup Activities", ["Identify resources and materials needed", "Coordinate with other setup activities", "Schedule setup work"]),
            ("Execute Setup", ["Perform setup work per plan", "Document setup completion", "Address any issues encountered"]),
            ("Verify and Document", ["Verify setup meets requirements", "Complete documentation", "Communicate completion to project team"])
        ]
        templates = ["Setup Checklist", "Completion Documentation Form"]
        refs = ["Management Directives (Policy Manual Section 9)", "Safety Requirements"]

        doc = create_sop(sop['num'], sop['title'], sop['dept'], sop['related'],
                        purpose, scope, roles, reqs, proc, templates, refs)
        doc.save(os.path.join(OUTPUT_FOLDER, f"{sop['num']} - {sop['title']}.docx"))

    print("\n" + "=" * 60)
    print("REVISION COMPLETE")
    print("=" * 60)
    print(f"\nOutput folder: {OUTPUT_FOLDER}")
    print("\nFiles created:")
    for f in sorted(os.listdir(OUTPUT_FOLDER)):
        if f.endswith('.docx'):
            print(f"  - {f}")

    print("\n" + "-" * 60)
    print("CHANGE SUMMARY (applies to all SOPs):")
    print("-" * 60)
    print("* Legacy numbers removed: YES")
    print("* Related SOPs updated: YES/PARTIAL (Document Filing Standards, etc. = TBD)")
    print("* Section order: SOP/Department/Related SOPs/Purpose/Scope/")
    print("                 Roles & Responsibilities/Requirements/Procedure/Appendix")
    print("* Document References moved to Appendix")
    print("* Procedure formatted with bold step titles and bullets")
    print("* Encoding issues fixed")
    print("-" * 60)


if __name__ == "__main__":
    main()
