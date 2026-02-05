"""
Build SOP v3 - Full Word Document Generation

Follows all 20 formatting rules. Uses python-docx for document manipulation
and direct XML editing for SharePoint metadata.
"""

import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import re
import shutil
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, Twips
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

BASE_DIR = Path(r"C:\Users\tewing\Desktop\Holding\SOP")
SOURCE_DIR = BASE_DIR / "SharePoint-SOPs"
OUTPUT_DIR = BASE_DIR / "Revised-SOPs"
TEMPLATE_PATH = BASE_DIR / "Templates" / "GSL_SOP_Master_Template_SHAREPOINT_ENABLED (v1).docx"

OUTPUT_DIR.mkdir(exist_ok=True)


def extract_source_content(source_path):
    """Extract structured content from source SOP."""
    with zipfile.ZipFile(source_path, 'r') as zf:
        xml_content = zf.read('word/document.xml').decode('utf-8')

    texts = re.findall(r'<w:t[^>]*>([^<]+)</w:t>', xml_content)
    full_text = ' '.join(texts)
    # Decode XML entities
    full_text = full_text.replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')

    content = {
        'sop_id': '',
        'title': '',
        'department': '',
        'related_sops': [],
        'purpose': '',
        'scope_intro': '',
        'scope_items': [],
        'roles': [],
        'requirements': [],
        'procedure': [],
        'appendix_items': [],
    }

    # Extract from filename
    filename = source_path.name
    id_match = re.search(r'(9\.\d+\.\d+)', filename)
    if id_match:
        content['sop_id'] = id_match.group(1)

    title_match = re.search(r'9\.\d+\.\d+\s*[–-]\s*(.+?)(?:\.docx|$)', filename)
    if title_match:
        content['title'] = title_match.group(1).strip()

    # Department
    dept_match = re.search(r'Department:\s*([^R]+?)(?:Related|Purpose|$)', full_text)
    if dept_match:
        content['department'] = dept_match.group(1).strip()

    # Related SOPs
    related_section = re.search(r'Related SOPs:(.+?)Purpose', full_text)
    if related_section:
        related_text = related_section.group(1).strip()
        # Find all SOP references - title starts with letter, ends before next SOP number
        sop_refs = re.findall(r'(9\.\d+\.\d+)\s*[-–]\s*([A-Za-z][^0-9]*?)(?=\s*9\.\d|\s*$)', related_text)
        for sop_id, sop_title in sop_refs:
            # Clean up title (remove trailing special chars)
            title = re.sub(r'[\s–•]+$', '', sop_title.strip())
            content['related_sops'].append({
                'id': sop_id,
                'title': title
            })

    # Purpose
    purpose_match = re.search(r'Purpose\s+(.+?)(?:Scope|$)', full_text)
    if purpose_match:
        content['purpose'] = purpose_match.group(1).strip()

    # Scope
    scope_match = re.search(r'Scope\s+(.+?)(?:Roles|$)', full_text, re.DOTALL)
    if scope_match:
        scope_text = scope_match.group(1).strip()
        # Look for intro
        intro_match = re.match(r'(This SOP applies[^.]+\.)', scope_text)
        if intro_match:
            content['scope_intro'] = intro_match.group(1)
        else:
            applies_match = re.match(r'(Applies to[^.]+\.?)', scope_text)
            if applies_match:
                content['scope_intro'] = applies_match.group(1)

        # Extract scope items if present
        # Usually follows the intro as bullet points or conditions

    # Roles & Responsibilities
    roles_match = re.search(r'Roles & Responsibilities\s+(.+?)(?:Requirements|$)', full_text, re.DOTALL)
    if roles_match:
        roles_text = roles_match.group(1).strip()

        # Define role names to look for
        role_names = ['Branch Manager', 'General Superintendent', 'Project Manager', 'Foreman',
                      'Safety', 'Estimator', 'Purchasing', 'Admin', 'Lead Journeyman', 'Operations Manager']
        role_pattern = '|'.join(role_names)

        # Find all role headers with their positions
        role_headers = list(re.finditer(rf'({role_pattern})\s*(?:\([^)]*\))?\s*:', roles_text, re.IGNORECASE))

        for i, match in enumerate(role_headers):
            role_name = match.group(1).strip()
            start_pos = match.end()
            # End position is start of next role or end of text
            end_pos = role_headers[i + 1].start() if i + 1 < len(role_headers) else len(roles_text)
            resp_text = roles_text[start_pos:end_pos].strip()

            responsibilities = []
            # Split by sentence endings followed by capital letters
            for resp in re.split(r'(?<=[.!])\s+(?=[A-Z])', resp_text):
                resp = resp.strip()
                if resp and len(resp) > 5:
                    responsibilities.append(resp)

            if responsibilities:
                content['roles'].append({
                    'name': role_name,
                    'responsibilities': responsibilities
                })

    # Requirements
    req_match = re.search(r'Requirements\s+(.+?)(?:Procedure|$)', full_text, re.DOTALL)
    if req_match:
        req_text = req_match.group(1).strip()
        for item in re.split(r'[•–]\s*', req_text):
            item = item.strip()
            if item and len(item) > 3 and not item.startswith('['):
                content['requirements'].append(item)

    # Procedure
    proc_match = re.search(r'Procedure\s+(.+?)(?:Appendix|$)', full_text, re.DOTALL)
    if proc_match:
        proc_text = proc_match.group(1).strip()
        # Parse steps
        step_pattern = r'Step\s+(\d+):\s*([^–•]+?)(?:–|•)\s*(.+?)(?=Step\s+\d+:|$)'
        step_matches = re.findall(step_pattern, proc_text, re.DOTALL)

        for step_num, step_name, step_desc in step_matches:
            # Clean up description - combine bullet points
            desc_parts = re.split(r'[•–]\s*', step_desc)
            desc_clean = ' '.join([p.strip() for p in desc_parts if p.strip()])

            content['procedure'].append({
                'number': step_num,
                'name': step_name.strip(),
                'description': desc_clean
            })

    # Appendix
    app_match = re.search(r'Appendix\s+(.+?)$', full_text, re.DOTALL)
    if app_match:
        app_text = app_match.group(1).strip()
        # Extract items
        for item in re.findall(r'(?:–|•)\s*([^–•\n]+)', app_text):
            item = item.strip()
            if item and len(item) > 3 and not item.startswith('['):
                content['appendix_items'].append(item)

    return content


def add_keep_next(paragraph):
    """Add keepNext property to paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    keepNext = OxmlElement('w:keepNext')
    pPr.insert(0, keepNext)


def add_keep_lines(paragraph):
    """Add keepLines property to paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    keepLines = OxmlElement('w:keepLines')
    pPr.insert(0, keepLines)


def set_paragraph_indent(paragraph, left_twips):
    """Set left indent on paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), str(left_twips))
    pPr.append(ind)


def add_bullet_numbering(paragraph, num_id=37):
    """Add bullet numbering to paragraph (uses template's numId 37)."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    numId = OxmlElement('w:numId')
    numId.set(qn('w:val'), str(num_id))
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)


def add_horizontal_rule(doc):
    """Add a proper Word horizontal rule using lxml."""
    from lxml import etree

    p = doc.add_paragraph()
    add_keep_lines(p)

    # Create VML horizontal rule using raw XML
    hr_xml = '''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                        xmlns:v="urn:schemas-microsoft-com:vml"
                        xmlns:o="urn:schemas-microsoft-com:office:office">
        <v:rect style="width:0;height:1.5pt" o:hralign="center" o:hrstd="t" o:hr="t" fillcolor="#a0a0a0" stroked="f"/>
    </w:pict>'''

    run = p.add_run()
    pict_element = etree.fromstring(hr_xml.encode())
    run._r.append(pict_element)

    return p


def build_document(content, output_path):
    """Build Word document from content following all formatting rules."""

    # Start from template to preserve styles, numbering, etc.
    doc = Document(TEMPLATE_PATH)

    # Clear all paragraphs from template body
    for para in doc.paragraphs[:]:
        p = para._element
        p.getparent().remove(p)

    # Clear all tables from template body
    for table in doc.tables[:]:
        t = table._tbl
        t.getparent().remove(t)

    # === HEADER BLOCK ===

    # SOP Title (bold)
    p = doc.add_paragraph()
    add_keep_next(p)
    run = p.add_run(f"SOP: {content['sop_id']} – {content['title']}")
    run.bold = True

    # Department (bold label, normal value)
    p = doc.add_paragraph()
    add_keep_next(p)
    run1 = p.add_run("Department:")
    run1.bold = True
    run2 = p.add_run(f" {content['department'] or 'Project Management / Operations'}")

    # Related SOPs label
    p = doc.add_paragraph()
    add_keep_next(p)
    run = p.add_run("Related SOPs:")
    run.bold = True

    # Related SOP entries (plain text, not bullets)
    for related in content['related_sops']:
        p = doc.add_paragraph()
        add_keep_next(p)
        add_keep_lines(p)
        p.add_run(f"{related['id']} – {related['title']}")

    # Horizontal rule
    add_horizontal_rule(doc)

    # === PURPOSE ===
    p = doc.add_paragraph()
    add_keep_next(p)
    add_keep_lines(p)
    run = p.add_run("Purpose")
    run.bold = True

    p = doc.add_paragraph()
    add_keep_lines(p)
    p.add_run(content['purpose'] or "To establish standardized procedures for this activity.")

    doc.add_paragraph()  # Empty paragraph for spacing

    # Horizontal rule
    add_horizontal_rule(doc)

    # === SCOPE ===
    p = doc.add_paragraph()
    add_keep_next(p)
    add_keep_lines(p)
    run = p.add_run("Scope")
    run.bold = True

    if content['scope_intro']:
        p = doc.add_paragraph()
        add_keep_lines(p)
        p.add_run(content['scope_intro'])

    # Scope items as bullets
    for item in content['scope_items']:
        p = doc.add_paragraph(item, style='List Paragraph')
        add_keep_lines(p)
        add_bullet_numbering(p)

    doc.add_paragraph()

    # Horizontal rule
    add_horizontal_rule(doc)

    # === ROLES & RESPONSIBILITIES ===
    p = doc.add_paragraph()
    add_keep_next(p)
    run = p.add_run("Roles & Responsibilities")
    run.bold = True

    for role in content['roles']:
        # Role name (not bold)
        p = doc.add_paragraph()
        add_keep_next(p)
        p.add_run(role['name'])

        # Responsibilities as bullets
        for resp in role['responsibilities']:
            p = doc.add_paragraph(resp, style='List Paragraph')
            add_keep_lines(p)
            add_bullet_numbering(p)

        doc.add_paragraph()  # Spacing between roles

    # Horizontal rule
    add_horizontal_rule(doc)

    # === REQUIREMENTS ===
    p = doc.add_paragraph()
    add_keep_next(p)
    add_keep_lines(p)
    run = p.add_run("Requirements")
    run.bold = True

    for req in content['requirements']:
        p = doc.add_paragraph(req, style='List Paragraph')
        add_keep_lines(p)
        add_bullet_numbering(p)

    doc.add_paragraph()

    # Horizontal rule
    add_horizontal_rule(doc)

    # === PROCEDURE ===
    p = doc.add_paragraph()
    add_keep_next(p)
    add_keep_lines(p)
    run = p.add_run("Procedure")
    run.bold = True

    for step in content['procedure']:
        # Step header (bold, indented)
        p = doc.add_paragraph()
        add_keep_next(p)
        set_paragraph_indent(p, 360)
        run = p.add_run(f"Step {step['number']} – {step['name']}")
        run.bold = True

        # Step description (normal, indented)
        if step.get('description'):
            p = doc.add_paragraph()
            add_keep_lines(p)
            set_paragraph_indent(p, 360)
            p.add_run(step['description'])

        doc.add_paragraph()  # Spacing between steps

    # === APPENDIX ===
    p = doc.add_paragraph()
    add_keep_next(p)
    add_keep_lines(p)
    run = p.add_run("Appendix")
    run.bold = True

    for item in content['appendix_items']:
        p = doc.add_paragraph(item, style='List Paragraph')
        add_keep_lines(p)
        add_bullet_numbering(p)

    # Save document
    doc.save(output_path)

    return True


def determine_raci(content):
    """Determine RACI from content."""
    raci = {'responsible': [], 'accountable': [], 'consulted': [], 'informed': []}

    role_names = [r['name'] for r in content['roles']]
    title_lower = content['title'].lower()

    # Determine responsible
    if 'team selection' in title_lower:
        raci['responsible'] = ['Branch Manager']
    elif 'safety' in title_lower:
        raci['responsible'] = ['Safety']
    elif 'foreman' in title_lower or 'field' in title_lower:
        raci['responsible'] = ['Foreman']
    elif any('Branch Manager' in r for r in role_names):
        raci['responsible'] = ['Branch Manager']
    elif any('Project Manager' in r for r in role_names):
        raci['responsible'] = ['Project Manager']
    else:
        raci['responsible'] = ['Project Manager']

    # Determine accountable
    if raci['responsible'][0] in ['Foreman', 'Safety']:
        raci['accountable'] = ['General Superintendent']
    elif raci['responsible'][0] in ['Branch Manager', 'Project Manager']:
        raci['accountable'] = ['Operations Manager']
    else:
        raci['accountable'] = ['Project Manager']

    # Consulted = other roles
    for role in role_names:
        if role not in raci['responsible'] and role not in raci['accountable']:
            raci['consulted'].append(role)

    return raci


def update_metadata(output_path, content, raci):
    """Update SharePoint metadata in the document."""
    temp_path = output_path.with_suffix('.tmp')

    with zipfile.ZipFile(output_path, 'r') as zf_in:
        with zipfile.ZipFile(temp_path, 'w', zipfile.ZIP_DEFLATED) as zf_out:
            for item in zf_in.namelist():
                data = zf_in.read(item)

                if item == 'customXml/item4.xml':
                    # Update SharePoint metadata
                    xml_content = data.decode('utf-8')
                    root = ET.fromstring(xml_content)

                    sp_ns = 'b11c1c0a-1848-4118-b7cf-ce9450f86f68'
                    xsi_ns = 'http://www.w3.org/2001/XMLSchema-instance'

                    doc_mgmt = None
                    for child in root:
                        if 'documentManagement' in child.tag:
                            doc_mgmt = child
                            break

                    if doc_mgmt is not None:
                        values = {
                            'SOPID': content['sop_id'],
                            'SOPFileName': f"{content['sop_id']} – {content['title']}.docx",
                            'Description': content['title'],
                            'Department_x002f_Division': content['department'],
                            'Status': 'Draft',
                            'RACI_Responsible': ', '.join(raci['responsible']),
                            'RACI_Accountable': ', '.join(raci['accountable']),
                            'RACI_Consulted': ', '.join(raci['consulted']),
                            'RACI_Informed': ', '.join(raci['informed']),
                        }

                        for child in doc_mgmt:
                            local_name = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                            if local_name in values and values[local_name]:
                                nil_attr = f'{{{xsi_ns}}}nil'
                                if nil_attr in child.attrib:
                                    del child.attrib[nil_attr]
                                child.text = str(values[local_name])

                    data = ET.tostring(root, encoding='unicode', xml_declaration=True).encode('utf-8')

                elif item == 'docProps/core.xml':
                    # Update core properties
                    xml_content = data.decode('utf-8')
                    root = ET.fromstring(xml_content)

                    ns = {
                        'dc': 'http://purl.org/dc/elements/1.1/',
                        'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
                        'dcterms': 'http://purl.org/dc/terms/',
                    }

                    for el in root:
                        local = el.tag.split('}')[-1] if '}' in el.tag else el.tag
                        if local == 'title':
                            el.text = f"SOP: {content['sop_id']} – {content['title']}"
                        elif local == 'subject':
                            el.text = content['department']
                        elif local == 'keywords':
                            el.text = f"{content['sop_id']}, {content['department']}, Team Selection"
                        elif local == 'description':
                            el.text = content['purpose'][:500] if content['purpose'] else ''

                    data = ET.tostring(root, encoding='unicode', xml_declaration=True).encode('utf-8')

                zf_out.writestr(item, data)

    temp_path.replace(output_path)


def process_sop(source_path):
    """Process a single SOP."""
    output_path = OUTPUT_DIR / source_path.name

    # Extract content
    content = extract_source_content(source_path)

    if not content['sop_id']:
        return None, "No SOP ID found"

    # Determine RACI
    raci = determine_raci(content)

    # Build document
    build_document(content, output_path)

    # Update metadata
    update_metadata(output_path, content, raci)

    return {
        'sop_id': content['sop_id'],
        'title': content['title'],
        'roles_count': len(content['roles']),
        'steps_count': len(content['procedure']),
        'raci': raci
    }, None


def main():
    import sys

    print("=" * 60)
    print("BUILD SOP v3 - PROPER FORMATTING")
    print("=" * 60)

    # Check for --all flag
    process_all = '--all' in sys.argv

    if not process_all:
        # Test with single SOP
        test_file = SOURCE_DIR / "9.2.010 – Team Selection.docx"

        if not test_file.exists():
            print(f"Source not found: {test_file}")
            return

        print(f"\nProcessing: {test_file.name}")

        result, error = process_sop(test_file)

        if result:
            print(f"\n  [OK] Generated: {result['sop_id']} – {result['title']}")
            print(f"       Roles: {result['roles_count']}, Steps: {result['steps_count']}")
            print(f"       RACI R: {result['raci']['responsible']}")
            print(f"       RACI A: {result['raci']['accountable']}")
            print(f"       RACI C: {result['raci']['consulted']}")
        else:
            print(f"  [ERR] {error}")

        print(f"\nOutput: {OUTPUT_DIR / test_file.name}")
        print("\nRun with --all to process all SOPs")
    else:
        # Process all SOPs
        sop_files = sorted(SOURCE_DIR.glob("*.docx"))
        print(f"\nProcessing {len(sop_files)} SOPs...")

        success = 0
        errors = []

        for sop_file in sop_files:
            try:
                result, error = process_sop(sop_file)
                if result:
                    print(f"  [OK] {result['sop_id']} - R: {', '.join(result['raci']['responsible'])}")
                    success += 1
                else:
                    errors.append((sop_file.name, error))
                    print(f"  [SKIP] {sop_file.name}: {error}")
            except Exception as e:
                errors.append((sop_file.name, str(e)))
                print(f"  [ERR] {sop_file.name}: {e}")

        print(f"\n{'='*60}")
        print(f"COMPLETE")
        print(f"  Success: {success}")
        print(f"  Errors: {len(errors)}")
        print(f"  Output: {OUTPUT_DIR}")
        print(f"{'='*60}")


if __name__ == "__main__":
    main()
