"""
Rebuild SOPs on Template
Creates new SOPs based on the SharePoint-enabled template,
transferring content from original SOPs into the template structure.
"""
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import json
import re
import shutil
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches

BASE_DIR = Path(r"C:\Users\tewing\Desktop\Holding\SOP")
SOURCE_DIR = BASE_DIR / "SharePoint-SOPs"
OUTPUT_DIR = BASE_DIR / "Revised-SOPs"
TEMPLATE_PATH = BASE_DIR / "Templates" / "GSL_SOP_Master_Template_SHAREPOINT_ENABLED (v1).docx"
REFERENCE_FILE = BASE_DIR / "docs" / "policy_reference.json"

OUTPUT_DIR.mkdir(exist_ok=True)

# Namespace for SharePoint metadata
NAMESPACE = "b11c1c0a-1848-4118-b7cf-ce9450f86f68"
ET.register_namespace('p', 'http://schemas.microsoft.com/office/2006/metadata/properties')
ET.register_namespace('xsi', 'http://www.w3.org/2001/XMLSchema-instance')
ET.register_namespace('pc', 'http://schemas.microsoft.com/office/infopath/2007/PartnerControls')
ET.register_namespace('', NAMESPACE)


def load_reference():
    with open(REFERENCE_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)


def extract_sop_content(source_path):
    """Extract structured content from source SOP using python-docx."""
    doc = Document(source_path)

    content = {
        'sop_id': '',
        'title': '',
        'department': '',
        'related_sops': [],
        'purpose': '',
        'scope': [],
        'roles': [],
        'requirements': [],
        'procedure': [],
        'appendix': [],
        'full_text': ''
    }

    # Extract SOP ID and title from filename
    filename = source_path.name
    id_match = re.search(r'(9\.\d+\.?\d*)', filename)
    if id_match:
        content['sop_id'] = id_match.group(1)

    title_match = re.search(r'9\.\d+\.?\d*\s*[–-]\s*(.+?)(?:\.docx|$)', filename)
    if title_match:
        content['title'] = title_match.group(1).strip()

    # Parse document sections
    current_section = None
    section_keywords = {
        'purpose': ['purpose'],
        'scope': ['scope', 'applies to'],
        'roles': ['roles', 'responsibilities'],
        'requirements': ['requirements', 'required'],
        'procedure': ['procedure', 'step'],
        'appendix': ['appendix', 'references']
    }

    all_text = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        all_text.append(text)
        text_lower = text.lower()

        # Check for section headers
        for section, keywords in section_keywords.items():
            if any(kw in text_lower for kw in keywords) and len(text) < 50:
                current_section = section
                break

        # Extract department
        if 'department' in text_lower and ':' in text:
            content['department'] = text.split(':', 1)[1].strip()

        # Extract related SOPs
        if 'related sop' in text_lower or (current_section == 'related' and text.startswith('9.')):
            if '9.' in text:
                content['related_sops'].append(text)

        # Add content to current section
        if current_section and text:
            if current_section == 'purpose' and not any(kw in text_lower for kw in ['purpose']):
                content['purpose'] += text + ' '
            elif current_section == 'scope':
                content['scope'].append(text)
            elif current_section == 'roles':
                content['roles'].append(text)
            elif current_section == 'requirements':
                content['requirements'].append(text)
            elif current_section == 'procedure':
                content['procedure'].append(text)
            elif current_section == 'appendix':
                content['appendix'].append(text)

    content['full_text'] = '\n'.join(all_text)

    # Clean up
    content['purpose'] = content['purpose'].strip()
    content['scope'] = [s for s in content['scope'] if s and 'scope' not in s.lower()]
    content['roles'] = [r for r in content['roles'] if r and 'roles' not in r.lower() and 'responsibilities' not in r.lower()]
    content['requirements'] = [r for r in content['requirements'] if r and 'requirement' not in r.lower()]
    content['procedure'] = [p for p in content['procedure'] if p and 'procedure' not in p.lower()]
    content['appendix'] = [a for a in content['appendix'] if a and 'appendix' not in a.lower()]

    return content


def determine_raci(content, reference):
    """Determine RACI based on content."""
    raci = {'responsible': [], 'accountable': [], 'consulted': [], 'informed': []}

    text = content['full_text'].lower()
    title_lower = content['title'].lower()
    phase = ''

    if content['sop_id']:
        phase_match = re.match(r'9\.(\d)', content['sop_id'])
        if phase_match:
            phase = f"9.{phase_match.group(1)}"

    # Role detection
    role_patterns = {
        'Project Manager': r'project\s*manager|PM\b',
        'General Superintendent': r'general\s*superintendent|GS\b',
        'Foreman': r'foreman|field\s*supervisor',
        'Safety': r'safety\s*(rep|coordinator|officer)?',
        'Estimator': r'estimator|estimating',
        'Purchasing': r'purchas(ing|er)|buyer',
        'Prefab': r'prefab|pre-fab',
        'Admin': r'admin|administrative',
    }

    mentioned = []
    for role, pattern in role_patterns.items():
        if re.search(pattern, text, re.IGNORECASE):
            mentioned.append(role)

    # Title-based hints
    title_hints = {
        'safety': 'Safety', 'quality': 'General Superintendent',
        'schedule': 'Project Manager', 'budget': 'Project Manager',
        'cost': 'Project Manager', 'billing': 'Project Manager',
        'procurement': 'Purchasing', 'purchase': 'Purchasing',
        'material': 'Purchasing', 'vendor': 'Purchasing',
        'prefab': 'Prefab', 'field': 'Foreman', 'setup': 'Foreman',
        'coordination': 'Project Manager', 'meeting': 'Project Manager',
        'closeout': 'Project Manager', 'documentation': 'Admin',
    }

    phase_default = {
        '9.1': 'Estimator', '9.2': 'Project Manager', '9.3': 'Foreman',
        '9.4': 'Foreman', '9.5': 'General Superintendent', '9.6': 'Project Manager'
    }

    responsible = None
    for hint, role in title_hints.items():
        if hint in title_lower:
            responsible = role
            break

    if not responsible:
        responsible = phase_default.get(phase, 'Project Manager')

    raci['responsible'] = [responsible]

    if responsible in ['Foreman', 'Safety']:
        raci['accountable'] = ['General Superintendent']
    elif responsible in ['Purchasing', 'Admin', 'Prefab']:
        raci['accountable'] = ['Project Manager']
    elif responsible == 'General Superintendent':
        raci['accountable'] = ['Project Manager']
    elif responsible == 'Project Manager':
        raci['accountable'] = ['Operations Manager']
    else:
        raci['accountable'] = ['Project Manager']

    for role in mentioned:
        if role not in raci['responsible'] and role not in raci['accountable']:
            raci['consulted'].append(role)

    raci['consulted'] = raci['consulted'][:3]

    return raci


def populate_template(template_path, output_path, content, raci, reference):
    """Create new document from template with content inserted."""

    # Copy template
    shutil.copy2(template_path, output_path)

    # Open and modify
    doc = Document(output_path)

    # Get phase info
    phase = ''
    if content['sop_id']:
        phase_match = re.match(r'9\.(\d)', content['sop_id'])
        if phase_match:
            phase = f"9.{phase_match.group(1)}"

    phase_mapping = reference.get('sop_phase_mapping', {})
    phase_info = phase_mapping.get(phase, {})

    # Replace content in paragraphs
    for para in doc.paragraphs:
        text = para.text

        # Header replacements
        if '[SOP #]' in text:
            para.text = text.replace('[SOP #]', content['sop_id'])
        if '[SOP Title]' in text:
            para.text = para.text.replace('[SOP Title]', content['title'])
        if '[Department(s)]' in text:
            para.text = text.replace('[Department(s)]', content['department'] or 'Project Management / Operations')

        # Related SOPs
        if '[Related SOP Title]' in text:
            if content['related_sops']:
                para.text = '\n'.join(content['related_sops'][:5])
            else:
                para.text = ''

        # Purpose
        if '[Describe the purpose' in text:
            para.text = content['purpose'] if content['purpose'] else 'To provide standardized procedures for this activity.'

        # Scope
        if '[Condition / trigger' in text:
            if content['scope']:
                para.text = content['scope'][0] if content['scope'] else ''

        # Roles
        if '[Role / Title]' in text:
            if content['roles']:
                para.text = content['roles'][0] if content['roles'] else ''
        if '[Responsibility]' in text:
            para.text = ''

        # Requirements
        if '[Requirement' in text:
            if content['requirements']:
                para.text = content['requirements'][0] if content['requirements'] else ''

        # Procedure steps
        if '[Describe what happens' in text:
            if content['procedure']:
                para.text = content['procedure'][0] if content['procedure'] else ''
        if '[Step Name]' in text:
            para.text = text.replace('[Step Name]', '')

    doc.save(output_path)

    # Now update the metadata XML
    update_sharepoint_metadata(output_path, content, raci, phase_info)

    return True


def update_sharepoint_metadata(docx_path, content, raci, phase_info):
    """Update SharePoint metadata in the document."""

    metadata = {
        'Status': 'Draft',
        'DocumentType': 'SOP',
        'SOPFileName': docx_path.name,
        'Description': content['title'][:255],
        'SOPID': content['sop_id'],
        'Tags_x002f_Keywords': phase_info.get('name', '')[:100],
        'Department_x002f_Division': content['department'] or 'Project Management',
        'RACI_Responsible': ', '.join(raci['responsible']),
        'RACI_Accountable': ', '.join(raci['accountable']),
        'RACI_Consulted': ', '.join(raci['consulted']),
        'RACI_Informed': ', '.join(raci['informed']),
    }

    temp_path = docx_path.with_suffix('.tmp')

    with zipfile.ZipFile(docx_path, 'r') as zf_in:
        with zipfile.ZipFile(temp_path, 'w', zipfile.ZIP_DEFLATED) as zf_out:
            for item in zf_in.namelist():
                data = zf_in.read(item)

                if item == 'customXml/item4.xml':
                    # Update metadata
                    xml_content = data.decode('utf-8')
                    xml_content = update_metadata_xml(xml_content, metadata)
                    data = xml_content.encode('utf-8')

                zf_out.writestr(item, data)

    temp_path.replace(docx_path)


def update_metadata_xml(xml_content, metadata):
    """Update the metadata XML values."""
    root = ET.fromstring(xml_content)

    doc_mgmt = None
    for child in root:
        if 'documentManagement' in child.tag:
            doc_mgmt = child
            break

    if doc_mgmt is None:
        return xml_content

    for child in doc_mgmt:
        local_name = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if local_name in metadata and metadata[local_name]:
            nil_attr = '{http://www.w3.org/2001/XMLSchema-instance}nil'
            if nil_attr in child.attrib:
                del child.attrib[nil_attr]
            child.text = str(metadata[local_name])

    return ET.tostring(root, encoding='unicode', xml_declaration=True)


def main():
    print("=" * 60)
    print("REBUILD SOPs ON SHAREPOINT TEMPLATE")
    print("=" * 60)

    reference = load_reference()
    sop_files = list(SOURCE_DIR.glob("*.docx"))

    print(f"\nProcessing {len(sop_files)} SOPs...")
    print("Each SOP will be rebuilt on the SharePoint-enabled template.\n")

    success = 0
    errors = []

    for sop_file in sorted(sop_files):
        try:
            # Extract content from original
            content = extract_sop_content(sop_file)

            if not content['sop_id']:
                errors.append((sop_file.name, "No SOP ID found"))
                print(f"  [SKIP] {sop_file.name} - No SOP ID")
                continue

            # Determine RACI
            raci = determine_raci(content, reference)

            # Create output path
            output_path = OUTPUT_DIR / sop_file.name

            # Populate template with content
            populate_template(TEMPLATE_PATH, output_path, content, raci, reference)

            print(f"  [OK] {content['sop_id']} - {content['title'][:40]}")
            success += 1

        except Exception as e:
            errors.append((sop_file.name, str(e)))
            print(f"  [ERR] {sop_file.name}: {e}")

    print(f"\n{'='*60}")
    print(f"COMPLETE")
    print(f"  Success: {success}")
    print(f"  Errors: {len(errors)}")
    print(f"  Output: {OUTPUT_DIR}")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
