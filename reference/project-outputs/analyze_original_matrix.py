#!/usr/bin/env python3
"""
Analyze original Key_SOP_Matrix.xlsx structure
"""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from openpyxl import load_workbook

original_path = r"C:\Users\tewing\Desktop\Claude Projects\Key_SOP_Matrix.xlsx"

wb = load_workbook(original_path, read_only=True, data_only=True)

print(f"Sheets in original Key_SOP_Matrix.xlsx:")
print("=" * 60)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n### {sheet_name}")
    print(f"Size: {ws.max_row} rows x {ws.max_column} cols")

    # Get headers (row 1)
    headers = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        headers.append(val if val else f"(empty col {col})")

    print(f"Headers: {headers}")

    # Show sample data (rows 2-5)
    print("Sample rows:")
    for row in range(2, min(6, ws.max_row + 1)):
        row_data = []
        for col in range(1, min(ws.max_column + 1, 8)):  # First 7 columns
            val = ws.cell(row=row, column=col).value
            if val:
                val_str = str(val)[:40] + "..." if len(str(val)) > 40 else str(val)
                row_data.append(val_str)
            else:
                row_data.append("(empty)")
        print(f"  Row {row}: {row_data}")

wb.close()
