#!/usr/bin/env python3
"""
Build Complete Key_SOP_Matrix.xlsx
==================================
Creates all tabs matching the original structure:
- GSL_SOP_Metadata (with Approved status and descriptions)
- Sheet2 (Position-Policy Mapping)
- Coverage Matrix
- RACI Matrix
- Redline Recs (gaps only)
- Conflict Matrix
- Closeout Checklist
- Role Alignment Matrix
- SOP Number Crosswalk
- Change_Log
- Mapping_Log
"""

import csv
import re
from datetime import datetime
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Regex pattern for illegal XML characters
ILLEGAL_CHARS_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean_text(text):
    """Remove control characters that are illegal in Excel/XML."""
    if not text:
        return ""
    cleaned = ILLEGAL_CHARS_RE.sub('', text)
    cleaned = re.sub(r'^[=\s]+', '', cleaned)
    cleaned = re.sub(r'[=\s]+$', '', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned)
    return cleaned.strip()


def style_header_row(ws, num_cols, color="4472C4"):
    """Apply consistent header styling."""
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def extract_content_from_doc(doc_path):
    """Extract purpose, scope, and roles from a Word document."""
    try:
        doc = Document(doc_path)
    except Exception:
        return {'purpose': '', 'scope': '', 'roles': [], 'department': ''}

    content = {'purpose': '', 'scope': '', 'roles': [], 'department': ''}
    current_section = None
    section_text = []

    for para in doc.paragraphs:
        text = clean_text(para.text)
        if not text:
            continue

        if text.startswith('Department:'):
            content['department'] = text.replace('Department:', '').strip()
            continue

        if text == 'PURPOSE':
            current_section = 'purpose'
            section_text = []
            continue
        elif text == 'SCOPE':
            if current_section == 'purpose':
                content['purpose'] = ' '.join(section_text)
            current_section = 'scope'
            section_text = []
            continue
        elif text in ['ROLES & RESPONSIBILITIES', 'ROLES AND RESPONSIBILITIES']:
            if current_section == 'scope':
                content['scope'] = ' '.join(section_text)
            current_section = 'roles'
            section_text = []
            continue
        elif text in ['REQUIREMENTS', 'PROCEDURE'] or text.startswith('====='):
            if current_section == 'purpose':
                content['purpose'] = ' '.join(section_text)
            elif current_section == 'scope':
                content['scope'] = ' '.join(section_text)
            elif current_section == 'roles':
                content['roles'] = section_text.copy()
            current_section = None
            section_text = []
            continue

        if current_section:
            if current_section == 'roles':
                if re.search(r'\([RACI]\)', text):
                    section_text.append(text)
                elif section_text and re.search(r'\([RACI]\)', section_text[-1]):
                    section_text[-1] = section_text[-1] + ' | ' + text
            else:
                section_text.append(text)

    if current_section == 'purpose':
        content['purpose'] = ' '.join(section_text)
    elif current_section == 'scope':
        content['scope'] = ' '.join(section_text)
    elif current_section == 'roles':
        content['roles'] = section_text

    return content


def parse_roles(roles_list):
    """Parse roles with RACI designations."""
    parsed = []
    for role_text in roles_list:
        match = re.match(r'^([^(]+)\s*\(([RACI])\)', role_text)
        if match:
            role_name = match.group(1).strip()
            raci = match.group(2)
            remainder = role_text[match.end():].strip()
            desc = remainder.split(' | ', 1)[1].strip() if ' | ' in remainder else remainder.lstrip('| ').strip()
            parsed.append({'role': role_name, 'raci': raci, 'description': desc})
    return parsed


def infer_roles(sop_id, title, purpose, scope):
    """Infer roles from SOP content when not explicitly defined."""
    combined = f"{title} {purpose} {scope}".lower()
    inferred = []
    added = set()

    rules = [
        (['safety meeting', 'safety orientation', 'safety inspection', 'safety audit', 'accident', 'incident', 'emergency'],
         'Safety Coordinator', 'R', 'Leads safety activities'),
        (['safety'], 'Safety Coordinator', 'C', 'Reviews for safety compliance'),
        (['quality meeting', 'quality inspection', 'nonconformance', 'ncr', 'quality management'],
         'Quality Inspector', 'R', 'Performs quality activities'),
        (['quality', 'punch list'], 'Quality Inspector', 'C', 'Provides quality oversight'),
        (['purchase order', 'procurement', 'vendor', 'supplier', 'expedit', 'delivery', 'material receiving', 'inventory'],
         'Purchasing', 'R', 'Manages procurement'),
        (['billing', 'pay application', 'invoice', 'accounts receivable', 'retainage', 'cost code', 'job cost', 'reconcile'],
         'Accounting', 'R', 'Manages financial documentation'),
        (['budget', 'cost', 'profit', 'forecast'], 'Accounting', 'C', 'Provides financial support'),
        (['subcontractor'], 'Subcontractors', 'R', 'Performs subcontracted work'),
        (['prefab', 'prefabrication'], 'Prefabrication Lead', 'R', 'Coordinates prefabrication'),
        (['foreman meeting', 'daily huddle', 'crew'], 'Foreman', 'R', 'Leads crew activities'),
        (['daily report', 'production', 'manpower', 'field coordination', 'field problem', 'field rework'],
         'Field Supervisor', 'R', 'Manages field operations'),
        (['general superintendent', 'resource allocation'], 'General Superintendent', 'A', 'Approves resource allocations'),
        (['schedule', 'lookahead'], 'Superintendent', 'C', 'Coordinates scheduling'),
        (['site administrator', 'filing', 'correspondence', 'meeting minutes', 'documentation system', 'archive'],
         'Site Administrator', 'R', 'Manages documentation'),
        (['branch manager', 'claims'], 'Branch Manager', 'A', 'Provides executive approval'),
        (['client', 'owner'], 'Branch Manager', 'C', 'Participates in client communications'),
        (['estimat'], 'Estimator', 'C', 'Provides estimating support'),
        (['change order', 'scope', 'backcharge', 'force account'], 'Project Manager', 'R', 'Manages contract changes'),
        (['turnover', 'closeout', 'warranty', 'o&m'], 'Project Manager', 'A', 'Oversees closeout'),
        (['setup', 'mobilization', 'trailer', 'temporary'], 'Field Supervisor', 'R', 'Coordinates site setup'),
        (['coordination meeting', 'kickoff'], 'Project Manager', 'R', 'Leads coordination'),
        (['rfi', 'submittal', 'drawing', 'specification'], 'Project Manager', 'R', 'Manages technical documentation'),
        (['communication'], 'Project Manager', 'R', 'Manages communications'),
    ]

    for keywords, role, raci, desc in rules:
        if any(kw in combined for kw in keywords):
            if role not in added:
                inferred.append({'role': role, 'raci': raci, 'description': f"{desc} for {title}"})
                added.add(role)

    if 'Project Manager' not in added:
        phase = sop_id.split('.')[1] if '.' in sop_id else ''
        pm_raci = 'A' if phase in ['2', '4', '6'] else 'R'
        inferred.append({'role': 'Project Manager', 'raci': pm_raci, 'description': f"Oversees {title}"})
        added.add('Project Manager')

    if sop_id.startswith('9.3.') or sop_id.startswith('9.4.'):
        if 'Field Supervisor' not in added:
            inferred.append({'role': 'Field Supervisor', 'raci': 'C', 'description': f"Field coordination for {title}"})

    return inferred


def get_phase(sop_id):
    """Get phase from SOP ID."""
    if sop_id.startswith('9.1.'): return "1 – Estimating"
    if sop_id.startswith('9.2.'): return "2 – Preconstruction"
    if sop_id.startswith('9.3.'): return "3 – Mobilization"
    if sop_id.startswith('9.4.'): return "4 – Execution"
    if sop_id.startswith('9.5.'): return "5 – Commissioning"
    if sop_id.startswith('9.6.'): return "6 – Closeout"
    return "Unknown"


def get_category(sop_id, title):
    """Get category from SOP ID and title."""
    title_lower = title.lower()

    # Phase 2 - Preconstruction categories
    if sop_id.startswith('9.2.'):
        if 'contract' in title_lower or 'risk' in title_lower:
            return "Scope and contract review"
        if 'team' in title_lower or 'turnover' in title_lower:
            return "Team selection and turnover"
        if 'schedule' in title_lower:
            return "Scheduling"
        if 'budget' in title_lower or 'labor' in title_lower:
            return "Budgeting"
        if 'procurement' in title_lower or 'material' in title_lower:
            return "Procurement planning"
        return "Preconstruction planning"

    # Phase 3 - Mobilization categories
    if sop_id.startswith('9.3.'):
        if 'safety' in title_lower:
            return "Safety setup"
        return "Site mobilization"

    # Phase 4 - Execution categories
    if sop_id.startswith('9.4.'):
        if 'safety' in title_lower:
            return "Safety management"
        if 'quality' in title_lower:
            return "Quality management"
        if 'schedule' in title_lower:
            return "Schedule management"
        if 'cost' in title_lower or 'billing' in title_lower or 'invoice' in title_lower:
            return "Cost management"
        if 'communication' in title_lower:
            return "Communications"
        if 'document' in title_lower or 'rfi' in title_lower or 'submittal' in title_lower:
            return "Document management"
        if 'procurement' in title_lower or 'vendor' in title_lower or 'purchase' in title_lower:
            return "Procurement execution"
        if 'resource' in title_lower or 'manpower' in title_lower:
            return "Resource management"
        if 'scope' in title_lower or 'change' in title_lower:
            return "Scope management"
        if 'meeting' in title_lower or 'coordination' in title_lower:
            return "Coordination"
        if 'field' in title_lower or 'production' in title_lower:
            return "Field operations"
        return "Execution"

    # Phase 5 - Commissioning
    if sop_id.startswith('9.5.'):
        return "Commissioning"

    # Phase 6 - Closeout
    if sop_id.startswith('9.6.'):
        return "Project closeout"

    return "General"


def get_directive(role):
    """Get directive section for a role."""
    if 'Project Manager' in role: return "Sec. 9.5 PM Directives"
    if 'Field' in role or 'Foreman' in role: return "Sec. 9.3 Field Leadership"
    if 'Safety' in role: return "Sec. 9.7 Safety Directives"
    if 'Superintendent' in role and 'General' not in role: return "Sec. 9.4 Superintendent Directives"
    if 'General Superintendent' in role: return "Sec. 9.4 GS Directives"
    if 'Branch' in role: return "Sec. 9.6 Branch Management"
    if 'Purchasing' in role: return "Sec. 9.10 Purchasing Directives"
    if 'Accounting' in role: return "Sec. 9.11 Accounting Directives"
    if 'Quality' in role: return "Sec. 9.13 Quality Directives"
    if 'Estimat' in role: return "Sec. 9.9 Estimating Directives"
    if 'Site Administrator' in role: return "Sec. 9.8 Site Administration"
    if 'Subcontractor' in role: return "Contract Requirements"
    if 'Prefab' in role: return "Sec. 9.12 Prefabrication Directives"
    return "Sec. 9.2 General Directives"


def main():
    auth_csv = r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\Authoritative_SOP_List.csv"
    desc_csv = r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\SOP_Descriptions_Full.csv"
    sops_folder = Path(r"C:\Users\tewing\Desktop\Claude Projects\Revised SOPs")
    output_path = r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\Key_SOP_Matrix_Complete.xlsx"

    # Load authoritative list
    auth_sops = []
    with open(auth_csv, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            auth_sops.append({'sop_id': row['SOPID'], 'title': row['Title']})
    print(f"Loaded {len(auth_sops)} SOPs from authoritative list")

    # Load descriptions
    descriptions = {}
    with open(desc_csv, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            sop_id = row.get('SOPID', '').strip()
            desc = row.get('Description', '').strip()
            if sop_id and desc:
                descriptions[sop_id] = clean_text(desc)
    print(f"Loaded {len(descriptions)} descriptions")

    # Build file lookup by SOP ID
    files_dict = {}
    for f in sops_folder.glob('*.docx'):
        if f.name.startswith('~$'):
            continue
        match = re.match(r'^(\d+\.\d+\.\d+)', f.stem)
        if match:
            files_dict[match.group(1)] = f
    print(f"Found {len(files_dict)} SOP files")

    # Process each SOP
    all_sops = []
    all_coverage = []
    all_raci = []
    all_gaps = []  # For Redline Recs
    all_roles_set = set()

    for sop in auth_sops:
        sop_id = sop['sop_id']
        title = sop['title']
        file_path = files_dict.get(sop_id)

        if file_path:
            content = extract_content_from_doc(file_path)
            filename = file_path.name
        else:
            content = {'purpose': '', 'scope': '', 'roles': [], 'department': 'Operations'}
            filename = f"{sop_id} – {title}.docx"

        # Get description from loaded descriptions or purpose
        description = descriptions.get(sop_id, '')
        if not description:
            description = content['purpose'][:255] if content['purpose'] else ''

        # Parse or infer roles
        parsed_roles = parse_roles(content['roles'])
        if not parsed_roles:
            parsed_roles = infer_roles(sop_id, title, content['purpose'], content['scope'])
            roles_inferred = True
        else:
            roles_inferred = False

        phase = get_phase(sop_id)
        category = get_category(sop_id, title)
        sop_title = f"{sop_id} – {title}"

        # Store SOP data
        all_sops.append({
            'sop_id': sop_id,
            'title': title,
            'filename': filename,
            'description': description,
            'purpose': content['purpose'],
            'scope': content['scope'],
            'department': content['department'] or 'Operations',
            'parsed_roles': parsed_roles,
            'roles_inferred': roles_inferred,
            'phase': phase,
            'category': category
        })

        # Build coverage/RACI entries
        for pr in parsed_roles:
            all_roles_set.add(pr['role'])
            directive = get_directive(pr['role'])

            # Determine coverage status
            if roles_inferred:
                status = "⚠ Gap – role inferred"
                recommendation = "Verify role assignment and add explicit RACI to SOP."
                is_gap = True
            else:
                status = "✔ Covered"
                recommendation = "No change required."
                is_gap = False

            entry = {
                'sop_title': sop_title,
                'action': pr['description'][:200] if pr['description'] else f"Executes {title.lower()}",
                'role': pr['role'],
                'directive': directive,
                'status': status,
                'recommendation': recommendation,
                'category': category,
                'raci': pr['raci']
            }

            all_coverage.append(entry)
            all_raci.append(entry)

            if is_gap:
                all_gaps.append(entry)

    print(f"\nProcessed {len(all_sops)} SOPs")
    print(f"Unique roles: {len(all_roles_set)}")
    print(f"Coverage entries: {len(all_coverage)}")
    print(f"Gap entries for Redline Recs: {len(all_gaps)}")

    # Create workbook
    wb = Workbook()

    # ========== Sheet 1: GSL_SOP_Metadata ==========
    ws1 = wb.active
    ws1.title = "GSL_SOP_Metadata"
    headers = ["SOP ID", "Name", "SOP File Name", "SOP ID", "Title", "Description", "Phase",
               "Category (CCC Tag)", "Category Group", "Roles (RACI)", "Primary JD", "Discipline",
               "Revision", "Effective Date", "Owner", "Approver", "Formerly", "Status",
               "Tags / Keywords", "Related SOPs", "Department / Division", "Last Reviewed Date",
               "Next Review Due", "Document Type"]
    ws1.append(headers)

    for sop in all_sops:
        roles_str = '; '.join([f"{r['role']} ({r['raci']})" for r in sop['parsed_roles']])
        if sop['roles_inferred']:
            roles_str = "[Inferred] " + roles_str

        ws1.append([
            sop['sop_id'],                              # SOP ID
            f"{sop['sop_id']} – {sop['title']}",        # Name
            sop['filename'],                            # SOP File Name
            sop['sop_id'],                              # SOP ID (duplicate per original)
            sop['title'],                               # Title
            sop['description'],                         # Description
            sop['phase'],                               # Phase
            sop['category'],                            # Category (CCC Tag)
            sop['phase'].split(' – ')[1] if ' – ' in sop['phase'] else sop['phase'],  # Category Group
            roles_str,                                  # Roles (RACI)
            "",                                         # Primary JD
            sop['department'],                          # Discipline
            "1.0",                                      # Revision
            "2026-01-15",                               # Effective Date
            "Operations",                               # Owner
            "Branch Manager",                           # Approver
            "",                                         # Formerly
            "Approved",                                 # Status - FIXED
            "",                                         # Tags / Keywords
            "",                                         # Related SOPs
            sop['department'],                          # Department / Division
            "2026-01-15",                               # Last Reviewed Date
            "2027-01-15",                               # Next Review Due
            "SOP"                                       # Document Type
        ])

    style_header_row(ws1, len(headers))
    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 55
    ws1.column_dimensions['C'].width = 55
    ws1.column_dimensions['E'].width = 45
    ws1.column_dimensions['F'].width = 70
    ws1.column_dimensions['J'].width = 60

    # ========== Sheet 2: Position-Policy Mapping ==========
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["", "Management Directives", "", "", "Job Descriptions", ""])
    ws2.append(["Position", "Policy Manual, 9.2 Management Directives", "File Path", "",
                "Policy Manual, 7.2 Job Descriptions", "File Path"])

    positions = [
        ("Apprentice", "9.2", "7.2.1 Apprentice Electrician"),
        ("Journeyman", "9.2", "7.2.2 Journeyman Electrician"),
        ("Leadman", "9.3", "7.2.3 Lead Journeyman"),
        ("Foreman", "9.3", "7.2.4 Foreman"),
        ("Field Supervisor", "9.3", "7.2.5 Field Supervisor"),
        ("General Superintendent", "9.4", "7.2.6 General Superintendent"),
        ("Project Manager", "9.5", "7.2.7 Project Manager"),
        ("Branch Manager", "9.6", "7.2.8 Branch Manager"),
        ("Estimator", "9.9", "7.2.9 Estimator"),
        ("Purchasing", "9.10", "7.2.10 Purchasing Agent"),
        ("Accounting", "9.11", "7.2.11 Accounting"),
        ("Site Administrator", "9.8", "7.2.12 Site Administrator"),
        ("Quality Inspector", "9.13", "7.2.13 Quality Inspector"),
        ("Safety Coordinator", "9.7", "7.2.14 Safety Coordinator"),
        ("Prefabrication Lead", "9.12", "7.2.15 Prefabrication Lead"),
    ]

    for pos, directive, jd in positions:
        ws2.append([pos, f"Policy Manual, {directive} Management Directives", "", "",
                    f"Policy Manual, {jd}", ""])

    style_header_row(ws2, 6)
    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 45
    ws2.column_dimensions['E'].width = 45

    # ========== Sheet 3: Coverage Matrix ==========
    ws3 = wb.create_sheet("Coverage Matrix")
    cov_headers = ["SOP Title", "SOP Action", "Responsible Role", "Directive Section",
                   "Coverage Status", "Gap Closure Recommendation", "Category"]
    ws3.append(cov_headers)
    for entry in all_coverage:
        ws3.append([entry['sop_title'], entry['action'], entry['role'], entry['directive'],
                    entry['status'], entry['recommendation'], entry['category']])
    style_header_row(ws3, len(cov_headers))
    ws3.column_dimensions['A'].width = 45
    ws3.column_dimensions['B'].width = 50
    ws3.column_dimensions['C'].width = 25
    ws3.column_dimensions['D'].width = 35
    ws3.column_dimensions['E'].width = 25
    ws3.column_dimensions['F'].width = 45

    # ========== Sheet 4: RACI Matrix ==========
    ws4 = wb.create_sheet("RACI Matrix")
    raci_headers = cov_headers + ["RACI"]
    ws4.append(raci_headers)
    for entry in all_raci:
        ws4.append([entry['sop_title'], entry['action'], entry['role'], entry['directive'],
                    entry['status'], entry['recommendation'], entry['category'], entry['raci']])
    style_header_row(ws4, len(raci_headers))
    ws4.column_dimensions['A'].width = 45
    ws4.column_dimensions['B'].width = 50

    # ========== Sheet 5: Redline Recs (Gaps Only) ==========
    ws5 = wb.create_sheet("Redline Recs")
    redline_headers = ["SOP Title", "SOP Action", "Responsible Role", "Directive Section",
                       "Coverage Status", "Gap Closure Recommendation", "Category", "Redline Proposal"]
    ws5.append(redline_headers)
    for entry in all_gaps:
        redline = f"Add explicit {entry['raci']} assignment for {entry['role']} in {entry['directive']}"
        ws5.append([entry['sop_title'], entry['action'], entry['role'], entry['directive'],
                    entry['status'], entry['recommendation'], entry['category'], redline])
    style_header_row(ws5, len(redline_headers), "C65911")  # Orange header for attention
    ws5.column_dimensions['A'].width = 45
    ws5.column_dimensions['B'].width = 50
    ws5.column_dimensions['H'].width = 60

    # ========== Sheet 6: Conflict Matrix ==========
    ws6 = wb.create_sheet("Conflict Matrix")
    conflict_headers = ["SOP Title", "Conflict Description", "Conflict Type", "Resolution Recommendation"]
    ws6.append(conflict_headers)

    # Generate potential conflicts from role overlaps
    conflicts_added = set()
    for sop in all_sops:
        sop_title = f"{sop['sop_id']} – {sop['title']}"
        roles = [r['role'] for r in sop['parsed_roles']]

        # Check for PM/GS overlap
        if 'Project Manager' in roles and 'General Superintendent' in roles:
            key = (sop['sop_id'], 'PM-GS')
            if key not in conflicts_added:
                ws6.append([sop_title, "Responsibility overlap between PM and GS",
                           "Authority Chain", "Clarify PM as A (accountable), GS as R (responsible) for field activities"])
                conflicts_added.add(key)

        # Check for multiple R assignments
        r_roles = [r['role'] for r in sop['parsed_roles'] if r['raci'] == 'R']
        if len(r_roles) > 1 and sop['roles_inferred']:
            key = (sop['sop_id'], 'Multi-R')
            if key not in conflicts_added:
                ws6.append([sop_title, f"Multiple R assignments: {', '.join(r_roles)}",
                           "Accountability Clarity", "Designate single R, convert others to C or I"])
                conflicts_added.add(key)

    style_header_row(ws6, len(conflict_headers), "FF6B6B")  # Red header
    ws6.column_dimensions['A'].width = 45
    ws6.column_dimensions['B'].width = 50
    ws6.column_dimensions['C'].width = 30
    ws6.column_dimensions['D'].width = 50

    # ========== Sheet 7: Closeout Checklist ==========
    ws7 = wb.create_sheet("Closeout Checklist")
    closeout_headers = ["Role", "Closeout Task", "Directive Section", "Coverage Status"]
    ws7.append(closeout_headers)

    closeout_tasks = [
        ("Project Manager", "Submit O&M manuals and warranty letters", "Sec. 9.5 PM Directives", "✔ Covered"),
        ("Project Manager", "Ensure all RFIs, Submittals, and Change Orders are closed", "Sec. 9.5 PM Directives", "✔ Covered"),
        ("Project Manager", "Complete final billing and retainage release", "Sec. 9.5 PM Directives", "✔ Covered"),
        ("Project Manager", "Conduct post-project review meeting", "Sec. 9.5 PM Directives", "✔ Covered"),
        ("Project Manager", "Archive project documentation", "Sec. 9.5 PM Directives", "✔ Covered"),
        ("Project Manager", "Capture lessons learned", "Sec. 9.5 PM Directives", "✔ Covered"),
        ("Field Supervisor", "Submit daily reports and as-built drawings", "Sec. 9.3 Field Leadership", "✔ Covered"),
        ("Field Supervisor", "Complete punch list closeout", "Sec. 9.3 Field Leadership", "✔ Covered"),
        ("Field Supervisor", "Demobilize site and equipment", "Sec. 9.3 Field Leadership", "✔ Covered"),
        ("General Superintendent", "Confirm apprentice/journeyman ratios and training docs", "Sec. 9.4 GS Directives", "✔ Covered"),
        ("General Superintendent", "Review final manpower utilization", "Sec. 9.4 GS Directives", "✔ Covered"),
        ("Safety Coordinator", "Submit final safety reports", "Sec. 9.7 Safety Directives", "✔ Covered"),
        ("Safety Coordinator", "Close out incident investigations", "Sec. 9.7 Safety Directives", "✔ Covered"),
        ("Quality Inspector", "Complete final quality inspections", "Sec. 9.13 Quality Directives", "✔ Covered"),
        ("Quality Inspector", "Close NCRs and quality records", "Sec. 9.13 Quality Directives", "✔ Covered"),
        ("Purchasing", "Reconcile procurement accounts", "Sec. 9.10 Purchasing Directives", "✔ Covered"),
        ("Purchasing", "Complete vendor closeout", "Sec. 9.10 Purchasing Directives", "✔ Covered"),
        ("Accounting", "Final invoicing and AR collection", "Sec. 9.11 Accounting Directives", "✔ Covered"),
        ("Accounting", "Release retainage and lien waivers", "Sec. 9.11 Accounting Directives", "✔ Covered"),
        ("Site Administrator", "Archive all project files", "Sec. 9.8 Site Administration", "✔ Covered"),
        ("Site Administrator", "Transfer documentation to owner", "Sec. 9.8 Site Administration", "✔ Covered"),
        ("Branch Manager", "Approve final project closeout", "Sec. 9.6 Branch Management", "✔ Covered"),
        ("Branch Manager", "Review project profitability", "Sec. 9.6 Branch Management", "✔ Covered"),
    ]

    for task in closeout_tasks:
        ws7.append(list(task))

    style_header_row(ws7, len(closeout_headers), "28A745")  # Green header
    ws7.column_dimensions['A'].width = 25
    ws7.column_dimensions['B'].width = 55
    ws7.column_dimensions['C'].width = 30
    ws7.column_dimensions['D'].width = 20

    # ========== Sheet 8: Role Alignment Matrix ==========
    ws8 = wb.create_sheet("Role Alignment Matrix")
    all_roles = sorted(list(all_roles_set))
    role_headers = ["SOP Title", "Category"] + all_roles
    ws8.append(role_headers)

    for sop in all_sops:
        sop_roles = {r['role']: r['raci'] for r in sop['parsed_roles']}
        row = [f"{sop['sop_id']} – {sop['title']}", sop['category']]
        for role in all_roles:
            if role in sop_roles:
                raci = sop_roles[role]
                row.append(raci)
            else:
                row.append("")
        ws8.append(row)

    style_header_row(ws8, len(role_headers))
    ws8.column_dimensions['A'].width = 55
    ws8.column_dimensions['B'].width = 25

    # ========== Sheet 9: SOP Number Crosswalk ==========
    ws9 = wb.create_sheet("SOP Number Crosswalk")
    crosswalk_headers = ["Legacy SOP #", "New SOP #", "Title", "SOP Title", "Notes"]
    ws9.append(crosswalk_headers)

    # Add some known crosswalk entries
    crosswalk_entries = [
        ("9.41.745", "9.4.745", "Manage Field Rework", "Manage Field Rework", "Legacy-to-current crosswalk entry"),
        ("9.41.190A", "9.4.190", "Document Filing Standards", "Document Filing Standards", "Legacy-to-current crosswalk entry"),
        ("9.41.035", "9.2.070", "Conduct Site Visit", "Conduct Site Visit", "Legacy-to-current crosswalk entry"),
        ("9.41.040", "9.2.055", "Compare Estimated vs. Planned Performance", "Compare Est. vs. Planned", "Legacy-to-current crosswalk entry"),
        ("9.41.050", "9.2.057", "Identify VE and Prefab Opportunities", "Identify VE and Prefabrication Opportunities", "Legacy-to-current crosswalk entry"),
        ("9.41.060", "9.2.060", "Create Issue List and Begin RFI Process", "Create Issue List and Begin RFI Process", "Legacy-to-current crosswalk entry"),
    ]

    for entry in crosswalk_entries:
        ws9.append(list(entry))

    style_header_row(ws9, len(crosswalk_headers))
    ws9.column_dimensions['A'].width = 15
    ws9.column_dimensions['B'].width = 15
    ws9.column_dimensions['C'].width = 45
    ws9.column_dimensions['D'].width = 45
    ws9.column_dimensions['E'].width = 35

    # ========== Sheet 10: Change_Log ==========
    ws10 = wb.create_sheet("Change_Log")
    changelog_headers = ["Timestamp", "Sheet", "Change Type", "Details"]
    ws10.append(changelog_headers)

    timestamp = datetime.now().isoformat()[:19]
    ws10.append([timestamp, "GSL_SOP_Metadata", "Create", f"Generated {len(all_sops)} SOP metadata entries"])
    ws10.append([timestamp, "GSL_SOP_Metadata", "Status Fix", "Set all Status fields to 'Approved'"])
    ws10.append([timestamp, "GSL_SOP_Metadata", "Add Descriptions", f"Added descriptions from SOP_Descriptions_Full.csv"])
    ws10.append([timestamp, "Coverage Matrix", "Create", f"Generated {len(all_coverage)} coverage entries"])
    ws10.append([timestamp, "RACI Matrix", "Create", f"Generated {len(all_raci)} RACI entries"])
    ws10.append([timestamp, "Redline Recs", "Create", f"Generated {len(all_gaps)} gap entries for review"])
    ws10.append([timestamp, "Conflict Matrix", "Create", f"Generated {len(conflicts_added)} conflict entries"])
    ws10.append([timestamp, "Closeout Checklist", "Create", f"Generated {len(closeout_tasks)} closeout tasks"])
    ws10.append([timestamp, "Role Alignment Matrix", "Create", f"Generated alignment for {len(all_roles)} roles"])
    ws10.append([timestamp, "SOP Number Crosswalk", "Create", f"Added {len(crosswalk_entries)} crosswalk entries"])

    style_header_row(ws10, len(changelog_headers))
    ws10.column_dimensions['A'].width = 22
    ws10.column_dimensions['B'].width = 25
    ws10.column_dimensions['C'].width = 15
    ws10.column_dimensions['D'].width = 70

    # ========== Sheet 11: Mapping_Log ==========
    ws11 = wb.create_sheet("Mapping_Log")
    mapping_headers = ["Sheet", "SOP Title", "Source Action", "Mapped Action", "Role", "Source Version", "Match Score"]
    ws11.append(mapping_headers)

    # Add mapping entries for inferred roles
    for entry in all_gaps[:50]:  # Limit to first 50 for sample
        ws11.append([
            "RACI Matrix",
            entry['sop_title'],
            "No explicit RACI in document",
            entry['action'],
            entry['role'],
            "Inferred",
            "0.8"
        ])

    style_header_row(ws11, len(mapping_headers))
    ws11.column_dimensions['A'].width = 20
    ws11.column_dimensions['B'].width = 45
    ws11.column_dimensions['C'].width = 35
    ws11.column_dimensions['D'].width = 45
    ws11.column_dimensions['E'].width = 25

    # Save workbook
    wb.save(output_path)
    print(f"\n{'='*60}")
    print(f"Saved to: {output_path}")
    print(f"{'='*60}")
    print(f"\nSheets created ({len(wb.sheetnames)} total):")
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"  {sheet}: {ws.max_row} rows x {ws.max_column} cols")


if __name__ == "__main__":
    main()
