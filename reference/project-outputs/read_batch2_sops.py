"""
read_batch2_sops.py
Reads 10 .docx files from the Revised SOPs folder (matching specific prefixes),
extracts all non-empty paragraphs (index, style, text) and all table contents,
and saves the full output to read_batch2_sops.txt.
"""

import os
import sys
from docx import Document

SOP_DIR = r"C:\Users\tewing\Desktop\Claude Projects\Revised SOPs"
OUTPUT_FILE = r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\read_batch2_sops.txt"

PREFIXES = [
    "9.4.010",
    "9.3.470",
    "9.3.120",
    "9.3.110",
    "9.3.100",
    "9.3.090",
    "9.3.080",
    "9.3.070",
    "9.3.060",
    "9.3.050",
]


def process_document(filepath, out):
    """Read a single .docx and write paragraph + table data to the output handle."""
    doc = Document(filepath)

    # --- Paragraphs ---
    out.write("  PARAGRAPHS (non-empty)\n")
    out.write("  " + "-" * 76 + "\n")
    para_count = 0
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            style_name = para.style.name if para.style else "(no style)"
            out.write(f"    [{idx:>4}]  Style: {style_name}\n")
            out.write(f"           Text:  {text}\n\n")
            para_count += 1
    if para_count == 0:
        out.write("    (no non-empty paragraphs found)\n\n")
    else:
        out.write(f"    --- {para_count} non-empty paragraphs ---\n\n")

    # --- Tables ---
    out.write("  TABLES\n")
    out.write("  " + "-" * 76 + "\n")
    if not doc.tables:
        out.write("    (no tables found)\n\n")
    else:
        for t_idx, table in enumerate(doc.tables):
            out.write(f"    Table {t_idx + 1}  ({len(table.rows)} rows x {len(table.columns)} cols)\n")
            for r_idx, row in enumerate(table.rows):
                cells_text = []
                for c_idx, cell in enumerate(row.cells):
                    cell_text = cell.text.strip().replace("\n", " | ")
                    cells_text.append(cell_text)
                out.write(f"      Row {r_idx:>3}: {cells_text}\n")
            out.write("\n")
        out.write(f"    --- {len(doc.tables)} table(s) ---\n\n")


def main():
    # Discover matching files
    all_files = sorted(os.listdir(SOP_DIR))
    matched = []
    for fname in all_files:
        if not fname.lower().endswith(".docx"):
            continue
        for prefix in PREFIXES:
            if fname.startswith(prefix):
                matched.append(fname)
                break

    print(f"Found {len(matched)} files matching the {len(PREFIXES)} prefixes.\n")

    with open(OUTPUT_FILE, "w", encoding="utf-8") as out:
        out.write("=" * 80 + "\n")
        out.write("  BATCH 2 SOP DOCUMENT READER OUTPUT\n")
        out.write(f"  Source: {SOP_DIR}\n")
        out.write(f"  Files processed: {len(matched)}\n")
        out.write("=" * 80 + "\n\n")

        for i, fname in enumerate(matched, 1):
            filepath = os.path.join(SOP_DIR, fname)
            header = f"FILE {i}/{len(matched)}: {fname}"
            out.write("=" * 80 + "\n")
            out.write(f"  {header}\n")
            out.write("=" * 80 + "\n\n")
            print(f"  [{i}/{len(matched)}] Processing: {fname}")

            try:
                process_document(filepath, out)
            except Exception as e:
                msg = f"  ERROR processing {fname}: {e}"
                out.write(msg + "\n\n")
                print(msg)

        out.write("=" * 80 + "\n")
        out.write("  END OF REPORT\n")
        out.write("=" * 80 + "\n")

    print(f"\nDone. Output saved to:\n  {OUTPUT_FILE}")
    print(f"  File size: {os.path.getsize(OUTPUT_FILE):,} bytes")


if __name__ == "__main__":
    main()
