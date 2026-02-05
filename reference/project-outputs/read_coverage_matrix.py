import openpyxl
import os
import sys
from collections import Counter

# Force UTF-8 output
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# Paths
INPUT_FILE = r"C:\Users\tewing\Desktop\Claude Projects\Key_SOP_Matrix_RACI_Updated.xlsx"
OUTPUT_FILE = r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\Coverage_Matrix_Complete.txt"

def truncate(val, maxlen=80):
    """Convert value to string, replace None/empty with '(empty)', truncate to maxlen."""
    if val is None:
        return "(empty)"
    s = str(val).strip()
    if s == "" or s.lower() == "none":
        return "(empty)"
    if len(s) > maxlen:
        return s[:maxlen-3] + "..."
    return s

def main():
    # Load workbook
    wb = openpyxl.load_workbook(INPUT_FILE, read_only=True, data_only=True)
    
    # Check sheet exists
    sheet_name = "Coverage Matrix"
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")
        return
    
    ws = wb[sheet_name]
    
    # Read all rows
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    
    if not rows:
        print("ERROR: No data found in sheet.")
        return
    
    # First row is header
    header_row = rows[0]
    data_rows = rows[1:]
    
    # Define our output columns (A=0 through K=10)
    col_labels = [
        "SOP Title",           # A = 0
        "SOP Action",          # B = 1
        "Responsible Role",    # C = 2
        "RACI-R",              # D = 3
        "RACI-A",              # E = 4
        "RACI-C",              # F = 5
        "RACI-I",              # G = 6
        "Directive Section",   # H = 7
        "Coverage Status",     # I = 8
        "Gap Recommendation",  # J = 9
        "Category",            # K = 10
    ]
    
    # Build output lines
    lines = []
    
    # Title
    lines.append("=" * 120)
    lines.append("COVERAGE MATRIX - COMPLETE DATA DUMP")
    lines.append(f"Source: {INPUT_FILE}")
    lines.append(f"Sheet: {sheet_name}")
    lines.append(f"Total data rows: {len(data_rows)}")
    lines.append(f"Excel headers: {[str(h) for h in header_row]}")
    lines.append("=" * 120)
    lines.append("")
    
    # Header line
    header_parts = ["Row#"] + col_labels
    header_line = " | ".join(header_parts)
    lines.append(header_line)
    lines.append("-" * len(header_line))
    
    # Tracking stats
    raci_r_count = 0
    raci_a_count = 0
    raci_c_count = 0
    raci_i_count = 0
    coverage_counter = Counter()
    category_counter = Counter()
    
    for idx, row in enumerate(data_rows, start=1):
        # Safely get values (row may have fewer columns)
        def get_col(col_idx):
            if col_idx < len(row):
                return row[col_idx]
            return None
        
        vals = [get_col(i) for i in range(11)]  # cols A-K (0-10)
        
        # Stats - check if cell has real data
        def has_data(v):
            if v is None:
                return False
            s = str(v).strip()
            return s != "" and s.lower() != "none"
        
        if has_data(vals[3]):
            raci_r_count += 1
        if has_data(vals[4]):
            raci_a_count += 1
        if has_data(vals[5]):
            raci_c_count += 1
        if has_data(vals[6]):
            raci_i_count += 1
        
        coverage_val = truncate(vals[8])
        category_val = truncate(vals[10])
        coverage_counter[coverage_val] += 1
        category_counter[category_val] += 1
        
        # Build line
        truncated = [truncate(v) for v in vals]
        parts = [f"{idx:<4}"] + truncated
        line = " | ".join(parts)
        lines.append(line)
    
    # Summary stats
    lines.append("")
    lines.append("=" * 120)
    lines.append("SUMMARY STATISTICS")
    lines.append("=" * 120)
    lines.append(f"Total data rows: {len(data_rows)}")
    lines.append("")
    lines.append("RACI Column Population:")
    lines.append(f"  RACI-R (col D): {raci_r_count} / {len(data_rows)} rows have data ({raci_r_count*100/len(data_rows):.1f}%)")
    lines.append(f"  RACI-A (col E): {raci_a_count} / {len(data_rows)} rows have data ({raci_a_count*100/len(data_rows):.1f}%)")
    lines.append(f"  RACI-C (col F): {raci_c_count} / {len(data_rows)} rows have data ({raci_c_count*100/len(data_rows):.1f}%)")
    lines.append(f"  RACI-I (col G): {raci_i_count} / {len(data_rows)} rows have data ({raci_i_count*100/len(data_rows):.1f}%)")
    lines.append("")
    lines.append("Coverage Status Distribution:")
    for status, count in sorted(coverage_counter.items(), key=lambda x: -x[1]):
        lines.append(f"  {status}: {count} ({count*100/len(data_rows):.1f}%)")
    lines.append("")
    lines.append("Category Distribution:")
    for cat, count in sorted(category_counter.items(), key=lambda x: -x[1]):
        lines.append(f"  {cat}: {count} ({count*100/len(data_rows):.1f}%)")
    lines.append("")
    
    wb.close()
    
    # Join all lines
    output = "\n".join(lines)
    
    # Save to file first (UTF-8)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        f.write(output)
    
    # Print to console
    print(output)
    print(f"\nOutput saved to: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
