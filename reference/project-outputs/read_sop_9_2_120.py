"""
Read SOP 9.2.120 .docx from Revised SOPs folder using python-docx.
Print filename, all non-empty paragraphs (index, style, text), and all tables.
Save output to read_9.2.120.txt.
"""

import glob
import sys
from docx import Document

SEARCH_DIR = r"C:\Users\tewing\Desktop\Claude Projects\Revised SOPs"
OUTPUT_FILE = r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\read_9.2.120.txt"

def main():
    pattern = SEARCH_DIR + chr(92) + "9.2.120*"
    matches = glob.glob(pattern)
    if not matches:
        print(f"ERROR: No file matching pattern {pattern}")
        sys.exit(1)

    filepath = matches[0]
    lines = []

    def out(text=""):
        print(text)
        lines.append(text)

    out("=" * 90)
    out(f"FILE: {filepath}")
    out("=" * 90)

    doc = Document(filepath)

    out("")
    out("-" * 90)
    out("PARAGRAPHS (non-empty)")
    out("-" * 90)
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            style_name = para.style.name if para.style else "(no style)"
            out(f"  [{idx:>3}]  Style: {style_name:<30s}  Text: {text}")

    out("")
    out("-" * 90)
    out(f"TABLES  (total: {len(doc.tables)})")
    out("-" * 90)
    for t_idx, table in enumerate(doc.tables):
        num_rows = len(table.rows)
        num_cols = len(table.columns)
        out(f"" + chr(10) + f"  TABLE {t_idx}  ({num_rows} rows x {num_cols} cols)")
        out("  " + "-" * 80)
        for r_idx, row in enumerate(table.rows):
            cells_text = []
            for c_idx, cell in enumerate(row.cells):
                cells_text.append(cell.text.strip())
            out(f"    Row {r_idx:>3}: {' | '.join(cells_text)}")
        out("  " + "-" * 80)

    out("")
    out("=" * 90)
    out("END OF DOCUMENT")
    out("=" * 90)

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        f.write(chr(10).join(lines) + chr(10))

    print(f"" + chr(10) + f">>> Output saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
