#!/usr/bin/env python3
"""Verify Key_SOP_Matrix_Complete.xlsx"""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\tewing\Desktop\Claude Projects\Project Outputs\Key_SOP_Matrix_Complete.xlsx", read_only=True)

print("=" * 70)
print("Key_SOP_Matrix_Complete.xlsx Verification")
print("=" * 70)

# Check GSL_SOP_Metadata
ws = wb["GSL_SOP_Metadata"]
print(f"\n### GSL_SOP_Metadata")
print(f"Rows: {ws.max_row - 1} SOPs")

# Find Status column
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
status_col = headers.index("Status") + 1 if "Status" in headers else None

if status_col:
    # Count status values
    status_counts = {}
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row=row, column=status_col).value
        status_counts[val] = status_counts.get(val, 0) + 1
    print(f"Status column values: {status_counts}")

# Check descriptions
desc_col = headers.index("Description") + 1 if "Description" in headers else None
if desc_col:
    with_desc = 0
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row=row, column=desc_col).value
        if val and len(str(val)) > 5:
            with_desc += 1
    print(f"SOPs with descriptions: {with_desc}/{ws.max_row - 1}")

# Show sample rows
print("\nSample data (first 3 rows):")
for row in range(2, min(5, ws.max_row + 1)):
    sop_id = ws.cell(row=row, column=1).value
    status = ws.cell(row=row, column=status_col).value if status_col else "N/A"
    desc = ws.cell(row=row, column=desc_col).value if desc_col else ""
    desc_preview = (desc[:50] + "...") if desc and len(desc) > 50 else desc
    print(f"  {sop_id}: Status={status}, Desc='{desc_preview}'")

# List all sheets
print(f"\n### All Sheets ({len(wb.sheetnames)}):")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"  - {name}: {ws.max_row} rows x {ws.max_column} cols")

wb.close()
