"""
Build Complete Course to Skills Mapping

Chain: Course → Training Tab → SOPs → Roles → Skills

Sources:
- TRAINING_SOP_MAPPING.md: Tab → SOPs
- JD_SOP_MAPPING.md: SOP → Roles
- MD_SOP_MAPPING.md: MD → SOPs → Roles
- TARGET_SKILL_RULES.md: Role → Skills
"""

import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# =============================================================================
# MAPPING 1: Training Tab → SOPs (from TRAINING_SOP_MAPPING.md)
# =============================================================================

TAB_TO_SOPS = {
    "Tab 1": {
        "name": "Tool Control",
        "sops": ["9.4.565", "9.4.700"],
        "skill_levels": "L1-L8",
        "roles": ["FE", "FL", "GS"]
    },
    "Tab 2": {
        "name": "Preplanning & Staging",
        "sops": ["9.2.070", "9.2.080", "9.2.100", "9.3.100", "9.3.110", "9.4.570", "9.4.640", "9.4.645", "9.4.705", "9.4.710"],
        "skill_levels": "L5-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 3": {
        "name": "Kickoff/Project Turnover Meeting",
        "sops": ["9.2.010", "9.2.015", "9.2.130"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 4": {
        "name": "Staging & Job Familiarization",
        "sops": ["9.2.070", "9.3.010", "9.3.020", "9.3.100", "9.3.110", "9.4.640"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 5": {
        "name": "Purchasing Buy out",
        "sops": ["9.2.020", "9.2.160", "9.4.605", "9.4.610", "9.4.615", "9.4.620"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM", "PUR"]
    },
    "Tab 6": {
        "name": "Scheduling",
        "sops": ["9.2.100", "9.2.110", "9.4.315", "9.4.320", "9.4.325", "9.4.330", "9.4.335", "9.4.340", "9.4.345", "9.4.350"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 7": {
        "name": "Job Plans",
        "sops": ["9.2.090", "9.2.100", "9.4.705", "9.4.710"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 8": {
        "name": "Change Order Process",
        "sops": ["9.4.355", "9.4.360", "9.4.365", "9.4.370", "9.4.375", "9.4.380", "9.4.385"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS", "PM", "CM"]
    },
    "Tab 9": {
        "name": "Submittals",
        "sops": ["9.4.200", "9.4.605", "9.4.620"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 10": {
        "name": "Job Setup/Schedule of Values",
        "sops": ["9.2.120", "9.2.140", "9.4.440", "9.4.445"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "PM"]
    },
    "Tab 11": {
        "name": "Estimating Take-off Procedures",
        "sops": ["9.2.020", "9.2.055", "9.2.090", "9.4.360", "9.4.365"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "EST"]
    },
    "Tab 12": {
        "name": "Look Ahead",
        "sops": ["9.4.335", "9.4.675", "9.4.680", "9.4.710"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 14": {
        "name": "Daily Reports",
        "sops": ["9.4.215", "9.4.490", "9.4.690", "9.4.695", "9.4.700"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 15": {
        "name": "Integrating Schedule and Daily Reports",
        "sops": ["9.4.215", "9.4.320", "9.4.325"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 16": {
        "name": "Weekly Summary Letters",
        "sops": ["9.4.270", "9.4.275", "9.4.285"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 17": {
        "name": "Creating Budgets",
        "sops": ["9.2.090", "9.2.140", "9.4.400", "9.4.405", "9.4.435"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 18": {
        "name": "Job Close Out",
        "sops": ["9.6.005", "9.6.010", "9.6.015", "9.6.020", "9.6.030", "9.6.040", "9.6.050", "9.6.060", "9.6.070", "9.6.085"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 19": {
        "name": "Manpower Projections",
        "sops": ["9.2.090", "9.4.340", "9.4.560", "9.4.690", "9.4.595"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 20": {
        "name": "Job Cost Projections",
        "sops": ["9.4.400", "9.4.405", "9.4.425", "9.4.430", "9.4.435"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 22": {
        "name": "Document Management",
        "sops": ["9.4.190A", "9.4.195", "9.4.200", "9.4.205", "9.4.210", "9.4.220", "9.4.225", "9.4.230", "9.4.235", "9.4.240"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 38": {
        "name": "Testing",
        "sops": ["9.4.530", "9.4.525", "9.5.550"],
        "skill_levels": "L5-L8",
        "roles": ["FE", "FL", "GS"]
    },
    "Tab 40": {
        "name": "Safety (General)",
        "sops": ["9.3.470", "9.4.475", "9.4.480", "9.4.485", "9.4.490", "9.4.495", "9.4.500", "9.4.505", "9.4.510"],
        "skill_levels": "L1-L8",
        "roles": ["FE", "FL", "GS", "PM", "BM", "EST", "CM", "PUR"]
    },
    "Tab 41": {
        "name": "Material Management & Control",
        "sops": ["9.2.080", "9.4.570", "9.4.625", "9.4.630", "9.4.631", "9.4.635", "9.4.640", "9.4.645", "9.4.650"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 42": {
        "name": "Prefab",
        "sops": ["9.2.057", "9.4.050", "9.4.585"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 49": {
        "name": "Bluebeam",
        "sops": ["9.2.100", "9.4.220", "9.6.030"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS"]
    },
    "Tab 51": {
        "name": "Job Structure & Org Chart",
        "sops": ["9.2.010", "9.2.015", "9.4.560"],
        "skill_levels": "L7-L8",
        "roles": ["FL", "GS", "PM"]
    },
    "Tab 57": {
        "name": "ProCore",
        "sops": ["9.2.120", "9.4.195", "9.4.200", "9.4.205", "9.4.210", "9.4.215", "9.4.320"],
        "skill_levels": "L6-L8",
        "roles": ["FL", "GS", "PM"]
    },
}

# Safety-specific training tabs
SAFETY_TABS = {
    "Safety - LOTO": {
        "name": "Lockout Tagout",
        "sops": ["9.4.480", "9.4.510"],
        "roles": ["FE", "FL", "GS"],
        "skill_area": "Electrical Safety"
    },
    "Safety - Arc Flash": {
        "name": "Arc Flash / NFPA 70E",
        "sops": ["9.4.480", "9.4.510"],
        "roles": ["FE", "FL", "GS"],
        "skill_area": "Electrical Safety"
    },
    "Safety - Fall Protection": {
        "name": "Fall Protection",
        "sops": ["9.4.480", "9.4.510"],
        "roles": ["FE", "FL", "GS"],
        "skill_area": "Fall Protection"
    },
    "Safety - Scaffolding": {
        "name": "Scaffolding Safety",
        "sops": ["9.4.480", "9.4.510"],
        "roles": ["FE", "FL", "GS"],
        "skill_area": "Fall Protection"
    },
    "Safety - HazCom": {
        "name": "Hazard Communication",
        "sops": ["9.4.480", "9.4.510"],
        "roles": ["FE", "FL", "GS"],
        "skill_area": "Chemical Safety"
    },
    "Safety - Confined Space": {
        "name": "Confined Space",
        "sops": ["9.4.480", "9.4.510"],
        "roles": ["FE", "FL", "GS"],
        "skill_area": "Confined Space"
    },
    "Safety - Coordinator": {
        "name": "Safety Coordinator Training",
        "sops": ["9.4.475", "9.4.480", "9.4.485", "9.4.490", "9.4.495", "9.4.500", "9.4.505", "9.4.510"],
        "roles": ["FL", "GS"],
        "skill_area": "Safety Management"
    },
    "Safety - Orientation": {
        "name": "New Hire Safety Orientation",
        "sops": ["9.4.475"],
        "roles": ["FE", "FL", "GS", "PM", "BM", "EST", "CM", "PUR"],
        "skill_area": "Safety Awareness"
    },
}

# =============================================================================
# MAPPING 2: Role → Required Skills (from TARGET_SKILL_RULES.md)
# =============================================================================

ROLE_SKILLS = {
    "FE": {  # Field Employee
        "name": "Field Employee (Journeyman/Apprentice)",
        "md": "MD 9.2",
        "skills": {
            "Safety": "Basic",
            "Lean Construction": "Awareness",
            "Communication & Coaching": "Awareness",
            "Emotional Intelligence": "Awareness",
            "Project Planning & Productivity": "Awareness",
            "Documentation & Compliance": "Awareness",
        }
    },
    "FL": {  # Field Leadership (Foreman)
        "name": "Field Leadership (Foreman)",
        "md": "MD 9.3",
        "skills": {
            "Safety": "Proficient",
            "Lean Construction": "Basic",
            "Leadership & Field Management": "Proficient",
            "Communication & Coaching": "Basic",
            "Emotional Intelligence": "Basic",
            "Performance Management": "Basic",
            "Construction Software": "Basic",
            "Project Planning & Productivity": "Basic",
            "Documentation & Compliance": "Basic",
            "Professional Development": "Awareness",
        }
    },
    "GS": {  # General Superintendent
        "name": "General Superintendent",
        "md": "MD 9.4",
        "skills": {
            "Safety": "Proficient",
            "Lean Construction": "Proficient",
            "Leadership & Field Management": "Expert",
            "Communication & Coaching": "Proficient",
            "Emotional Intelligence": "Proficient",
            "Performance Management": "Proficient",
            "Construction Software": "Basic",
            "Project Planning & Productivity": "Proficient",
            "Documentation & Compliance": "Proficient",
            "Professional Development": "Basic",
        }
    },
    "PM": {  # Project Manager
        "name": "Project Manager",
        "md": "MD 9.5",
        "skills": {
            "Safety": "Basic",
            "Lean Construction": "Basic",
            "Leadership & Field Management": "Basic",
            "Communication & Coaching": "Basic",
            "Emotional Intelligence": "Basic",
            "Performance Management": "Basic",
            "Construction Software": "Proficient",
            "Project Planning & Productivity": "Proficient",
            "Documentation & Compliance": "Proficient",
            "Professional Development": "Basic",
        }
    },
    "BM": {  # Branch Manager
        "name": "Branch Manager",
        "md": "MD 9.6",
        "skills": {
            "Safety": "Basic",
            "Lean Construction": "Proficient",
            "Leadership & Field Management": "Proficient",
            "Communication & Coaching": "Proficient",
            "Emotional Intelligence": "Proficient",
            "Performance Management": "Proficient",
            "Construction Software": "Basic",
            "Project Planning & Productivity": "Proficient",
            "Documentation & Compliance": "Basic",
            "Professional Development": "Proficient",
        }
    },
    "EST": {  # Estimator
        "name": "Estimator",
        "md": "MD 9.7",
        "skills": {
            "Safety": "Awareness",
            "Lean Construction": "Awareness",
            "Communication & Coaching": "Awareness",
            "Emotional Intelligence": "Awareness",
            "Construction Software": "Proficient",
            "Project Planning & Productivity": "Basic",
            "Documentation & Compliance": "Awareness",
            "Professional Development": "Awareness",
        }
    },
    "CM": {  # Contract Manager
        "name": "Contract Manager",
        "md": "MD 9.14",
        "skills": {
            "Safety": "Awareness",
            "Communication & Coaching": "Basic",
            "Emotional Intelligence": "Awareness",
            "Construction Software": "Basic",
            "Project Planning & Productivity": "Basic",
            "Documentation & Compliance": "Proficient",
            "Professional Development": "Awareness",
        }
    },
    "PUR": {  # Purchasing
        "name": "Purchasing",
        "md": "MD 9.16",
        "skills": {
            "Safety": "Awareness",
            "Communication & Coaching": "Awareness",
            "Emotional Intelligence": "Awareness",
            "Construction Software": "Basic",
            "Project Planning & Productivity": "Awareness",
            "Documentation & Compliance": "Basic",
            "Professional Development": "Awareness",
        }
    },
}

# =============================================================================
# MAPPING 3: Skill Level Set → Individual Skills
# =============================================================================

SKILL_LEVEL_SETS = {
    "Safety": [
        "Safety Awareness & Culture",
        "Accident Prevention",
        "Hazard Recognition",
        "Pre-Task Safety Planning",
        "Safety Inspections",
        "Incident Investigation",
        "Injury Management",
        "Emergency Response",
        "NFPA 70E Fundamentals",
        "Arc Flash Awareness",
        "Arc Flash Risk Assessment",
        "Arc Flash PPE Selection",
        "Energized Work Practices",
        "Electrical Hazard Recognition",
        "LOTO Fundamentals",
        "Simple LOTO Procedures",
        "Complex LOTO Procedures",
        "Abnormal Lock Removal",
        "De-Energization Verification",
        "Return to Service",
        "Fall Hazard Recognition",
        "Personal Fall Arrest Systems",
        "Ladder Safety",
        "Scaffold Fundamentals",
        "Scaffold Access & Egress",
        "Fall Protection on Scaffolds",
        "HazCom Program Fundamentals",
        "GHS System Understanding",
        "SDS Use & Access",
        "Chemical Hazard Identification",
        "Confined Space Program",
    ],
    "Lean Construction": [
        "Lean Fundamentals",
        "Lean Application",
        "Eight Wastes Recognition",
        "Eliminate Waste",
        "Continuous Improvement",
        "Pull Planning",
    ],
    "Leadership & Field Management": [
        "Foreman Role & Responsibilities",
        "Field Supervision",
        "Trade Coordination",
        "Jobsite Efficiency",
        "Quality Control",
        "Manpower Projection",
        "Job Closeout",
        "Leadership Fundamentals",
        "Values-Driven Leadership",
        "Accountability & Ownership",
        "Team Building",
        "Team Management",
        "Delegation & Empowerment",
        "Employee Motivation",
        "Team Morale",
        "Extreme Ownership",
    ],
    "Communication & Coaching": [
        "Effective Communication",
        "Building Relationships",
        "Mentoring & Coaching",
        "Conflict Resolution",
        "Setting Expectations",
    ],
    "Emotional Intelligence": [
        "EQ Fundamentals",
        "Self-Awareness",
        "Self-Regulation",
        "Empathy",
        "Social Skills",
        "EQ Leadership",
    ],
    "Performance Management": [
        "Performance Management Fundamentals",
        "Conducting Evaluations",
        "Performance Reviews",
        "Corrective Counseling",
    ],
    "Construction Software": [
        "Accubid Fundamentals",
        "ViewPoint Change Orders",
        "ProCore Fundamentals",
        "Job Plan Fundamentals",
        "Vista",
    ],
    "Project Planning & Productivity": [
        "Project Planning",
        "Scheduling Techniques",
        "Productivity Improvement",
        "Time Management",
        "Construction Project Lifecycle",
    ],
    "Documentation & Compliance": [
        "Documentation Practices",
        "Documentation & Reporting",
        "Regulatory Inspections",
    ],
    "Professional Development": [
        "Presentation Skills",
        "Strategic Thinking",
        "Industry Leadership",
    ],
}

# =============================================================================
# MAPPING 4: Course → Training Tab/Category
# =============================================================================

def get_course_mapping(title):
    """Map course title to Training Tab, roles, SOPs, and skills"""
    title_lower = title.lower()

    result = {
        "training_tab": None,
        "tab_name": None,
        "sops": [],
        "roles": [],
        "skill_level_sets": [],
        "skills": []
    }

    # === SAFETY COURSES ===

    # Lockout Tagout
    if "lockout" in title_lower or "tagout" in title_lower or "loto" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Lockout Tagout"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["LOTO Fundamentals"]
        if "simple" in title_lower:
            result["skills"].append("Simple LOTO Procedures")
        elif "complex" in title_lower:
            result["skills"].append("Complex LOTO Procedures")
        elif "abnormal" in title_lower:
            result["skills"].append("Abnormal Lock Removal")
        elif "return" in title_lower or "service" in title_lower:
            result["skills"].append("Return to Service")
        elif "verif" in title_lower or "de-energiz" in title_lower:
            result["skills"].extend(["De-Energization Verification", "Electrical Hazard Recognition"])
        elif "hazardous energy" in title_lower:
            result["skills"].extend(["Simple LOTO Procedures", "Electrical Hazard Recognition"])
        else:
            result["skills"].extend(["Simple LOTO Procedures", "De-Energization Verification"])
        return result

    # Arc Flash / NFPA 70E
    if "arc flash" in title_lower or "nfpa 70" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Arc Flash / NFPA 70E"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["NFPA 70E Fundamentals", "Arc Flash Awareness"]
        if "ppe" in title_lower:
            result["skills"].append("Arc Flash PPE Selection")
        if "risk" in title_lower or "assessment" in title_lower:
            result["skills"].append("Arc Flash Risk Assessment")
        if "boundar" in title_lower:
            result["skills"].append("Arc Flash Awareness")
        return result

    # Energized Work
    if "energized" in title_lower and "work" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Energized Work Practices"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Energized Work Practices", "Electrical Hazard Recognition"]
        return result

    # Fall Protection
    if "fall protection" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Fall Protection"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Fall Hazard Recognition", "Personal Fall Arrest Systems"]
        return result

    # Scaffolding
    if "scaffold" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Scaffolding Safety"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Scaffold Fundamentals", "Fall Protection on Scaffolds"]
        if "access" in title_lower:
            result["skills"].append("Scaffold Access & Egress")
        return result

    # Ladder
    if "ladder" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Ladder Safety"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Ladder Safety", "Fall Hazard Recognition"]
        return result

    # Hazard Communication / GHS / SDS
    if "hazard communication" in title_lower or "hazcom" in title_lower or "ghs" in title_lower or "sds" in title_lower or "safety data sheet" in title_lower or "container label" in title_lower or "chemical hazard" in title_lower or "pictogram" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Hazard Communication"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["HazCom Program Fundamentals", "GHS System Understanding"]
        if "sds" in title_lower or "safety data" in title_lower:
            result["skills"].append("SDS Use & Access")
        if "chemical" in title_lower:
            result["skills"].append("Chemical Hazard Identification")
        return result

    # Confined Space
    if "confined space" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Confined Space"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Confined Space Program"]
        return result

    # Safety Coordinator Series
    if "safety coordinator" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Safety Coordinator"
        result["sops"] = ["9.4.475", "9.4.480", "9.4.485", "9.4.490", "9.4.495", "9.4.500", "9.4.505", "9.4.510"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Safety Awareness & Culture"]
        if "orientation" in title_lower:
            result["skills"].append("Safety Awareness & Culture")
        if "task training" in title_lower:
            result["skills"].append("Pre-Task Safety Planning")
        if "meeting" in title_lower:
            result["skills"].append("Safety Awareness & Culture")
        if "planning" in title_lower or "hazard" in title_lower:
            result["skills"].extend(["Hazard Recognition", "Pre-Task Safety Planning"])
        if "inspection" in title_lower:
            result["skills"].append("Safety Inspections")
        if "disciplin" in title_lower:
            result["skill_level_sets"].append("Performance Management")
            result["skills"].append("Corrective Counseling")
        if "accident" in title_lower or "incident" in title_lower:
            result["skills"].append("Incident Investigation")
        if "injury" in title_lower or "case management" in title_lower:
            result["skills"].append("Injury Management")
        if "osha" in title_lower or "msha" in title_lower:
            result["skills"].append("Regulatory Inspections")
            result["skill_level_sets"].append("Documentation & Compliance")
        return result

    # Safety Orientation / New Hire
    if "new hire" in title_lower or ("safety" in title_lower and "orientation" in title_lower):
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "New Hire Safety Orientation"
        result["sops"] = ["9.4.475"]
        result["roles"] = ["FE", "FL", "GS", "PM", "BM", "EST", "CM", "PUR"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Safety Awareness & Culture", "Hazard Recognition", "Accident Prevention"]
        return result

    # OSHA Focus Four
    if "focus four" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "OSHA Focus Four"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Hazard Recognition", "Fall Hazard Recognition", "Electrical Hazard Recognition"]
        return result

    # Safety Management
    if "safety management" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Safety Management"
        result["sops"] = ["9.4.480", "9.4.485", "9.4.510"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Safety Awareness & Culture", "Safety Inspections", "Hazard Recognition"]
        return result

    # T.R.A.C.K.
    if "track" in title_lower and ("hazard" in title_lower or "planning" in title_lower):
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Hazard Planning (T.R.A.C.K.)"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Hazard Recognition", "Pre-Task Safety Planning"]
        return result

    # Slips/Trips/Falls
    if "slip" in title_lower or "trip" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Slip/Trip/Fall Prevention"
        result["sops"] = ["9.4.480", "9.4.510"]
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Safety"]
        result["skills"] = ["Hazard Recognition", "Accident Prevention"]
        return result

    # === LEAN CONSTRUCTION ===

    if "lean" in title_lower:
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Lean Construction"
        result["sops"] = ["9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Lean Construction"]
        if "what is lean" in title_lower or "introduction" in title_lower:
            result["skills"] = ["Lean Fundamentals"]
        elif "waste" in title_lower:
            result["skills"] = ["Eight Wastes Recognition", "Eliminate Waste"]
        elif "pull planning" in title_lower:
            result["skills"] = ["Pull Planning"]
            result["skill_level_sets"].append("Project Planning & Productivity")
            result["skills"].append("Scheduling Techniques")
        elif "continuous improvement" in title_lower:
            result["skills"] = ["Continuous Improvement"]
        elif "productivity" in title_lower:
            result["skills"] = ["Lean Application"]
            result["skill_level_sets"].append("Project Planning & Productivity")
            result["skills"].append("Productivity Improvement")
        else:
            result["skills"] = ["Lean Fundamentals", "Lean Application"]
        return result

    # Eight Wastes
    if "eight waste" in title_lower or ("waste" in title_lower and "construction" in title_lower):
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Lean Construction - Waste"
        result["sops"] = ["9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Lean Construction"]
        result["skills"] = ["Eight Wastes Recognition", "Eliminate Waste"]
        return result

    # Pull Planning / Sticky Note
    if "pull planning" in title_lower or "sticky note" in title_lower:
        result["training_tab"] = "Tab 6 - Scheduling"
        result["tab_name"] = "Pull Planning"
        result["sops"] = ["9.4.315", "9.4.335", "9.4.345"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Lean Construction", "Project Planning & Productivity"]
        result["skills"] = ["Pull Planning", "Scheduling Techniques"]
        return result

    # === VALUES-DRIVEN LEADERSHIP / COACH K ===

    if "values" in title_lower and "driven" in title_lower or "coach k" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Values-Driven Leadership"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Values-Driven Leadership", "Leadership Fundamentals"]
        if "team" in title_lower:
            result["skills"].append("Team Building")
        if "own" in title_lower:
            result["skills"].append("Accountability & Ownership")
        if "read" in title_lower or "people" in title_lower:
            result["skill_level_sets"].append("Emotional Intelligence")
            result["skills"].append("Empathy")
        if "communication" in title_lower:
            result["skill_level_sets"].append("Communication & Coaching")
            result["skills"].append("Effective Communication")
        if "emotion" in title_lower:
            result["skill_level_sets"].append("Emotional Intelligence")
            result["skills"].append("EQ Leadership")
        if "recruit" in title_lower or "retain" in title_lower:
            result["skills"].append("Team Building")
        if "performance" in title_lower:
            result["skill_level_sets"].append("Performance Management")
            result["skills"].append("Performance Management Fundamentals")
        if "feedback" in title_lower:
            result["skill_level_sets"].append("Communication & Coaching")
            result["skills"].append("Setting Expectations")
        if "develop" in title_lower and "leader" in title_lower:
            result["skills"].append("Leadership Fundamentals")
        if "motivation" in title_lower:
            result["skills"].append("Employee Motivation")
        if "bounce back" in title_lower:
            result["skill_level_sets"].append("Emotional Intelligence")
            result["skills"].append("Self-Regulation")
        return result

    # === EXTREME OWNERSHIP ===

    if "extreme ownership" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Extreme Ownership"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Extreme Ownership", "Accountability & Ownership", "Leadership Fundamentals"]
        return result

    # === LEADERSHIP LAWS ===

    if "leadership law" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Leadership Laws"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Leadership Fundamentals"]
        if "cover and move" in title_lower:
            result["skills"].append("Team Building")
        if "simple" in title_lower:
            result["skill_level_sets"].append("Communication & Coaching")
            result["skills"].append("Effective Communication")
        if "prioritize" in title_lower or "execute" in title_lower:
            result["skill_level_sets"].append("Project Planning & Productivity")
            result["skills"].append("Project Planning")
        if "decentralize" in title_lower:
            result["skills"].append("Delegation & Empowerment")
        return result

    # === FIELD LEADERSHIP INTRODUCTION ===

    if "field leadership" in title_lower or "introduction to field" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Field Leadership Introduction"
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Field Supervision", "Foreman Role & Responsibilities"]

        # Determine specific SOPs based on section
        if "section 1" in title_lower or "part 1" in title_lower:
            result["sops"] = ["9.2.010", "9.4.560"]
            if "skills" in title_lower and "traits" in title_lower:
                result["skills"].append("Foreman Role & Responsibilities")
        elif "section 2" in title_lower or "part 2" in title_lower:
            if "anatomy" in title_lower or "project" in title_lower:
                result["sops"] = ["9.2.015", "9.2.130"]
                result["skill_level_sets"].append("Project Planning & Productivity")
                result["skills"].append("Construction Project Lifecycle")
            elif "document" in title_lower:
                result["sops"] = ["9.4.195", "9.4.215"]
                result["skill_level_sets"].append("Documentation & Compliance")
                result["skills"].append("Documentation Practices")
            elif "efficiency" in title_lower:
                result["sops"] = ["9.4.705", "9.4.710"]
                result["skills"].append("Jobsite Efficiency")
            elif "coordinat" in title_lower:
                result["sops"] = ["9.4.345"]
                result["skills"].append("Trade Coordination")
            elif "manpower" in title_lower:
                result["sops"] = ["9.4.560", "9.4.690"]
                result["skills"].append("Manpower Projection")
            elif "job plan" in title_lower:
                result["sops"] = ["9.4.705", "9.4.710"]
                result["skill_level_sets"].append("Construction Software")
                result["skills"].append("Job Plan Fundamentals")
            elif "closeout" in title_lower:
                result["sops"] = ["9.6.005", "9.6.010"]
                result["skills"].append("Job Closeout")
            else:
                result["sops"] = ["9.2.010", "9.4.560"]
        elif "section 3" in title_lower or "part 3" in title_lower:
            result["sops"] = ["9.2.010", "9.4.560"]
            if "leadership" in title_lower:
                result["skills"].append("Leadership Fundamentals")
            if "communication" in title_lower:
                result["skill_level_sets"].append("Communication & Coaching")
                result["skills"].append("Effective Communication")
        elif "section 4" in title_lower or "part 4" in title_lower or "policy" in title_lower:
            result["sops"] = ["9.4.475"]
            result["skill_level_sets"].append("Documentation & Compliance")
            if "harassment" in title_lower or "discrimination" in title_lower:
                result["skills"].append("Documentation & Reporting")
            elif "fmla" in title_lower:
                result["skills"].append("Documentation & Reporting")
            elif "substance" in title_lower:
                result["skill_level_sets"].append("Safety")
                result["skills"].append("Safety Awareness & Culture")
            elif "eap" in title_lower or "mental health" in title_lower:
                result["skills"].append("Team Morale")
            else:
                result["skills"].append("Documentation & Reporting")
        else:
            result["sops"] = ["9.2.010", "9.4.560"]
        return result

    # === FOREMAN ROLES ===

    if "foreman" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Foreman Training"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Foreman Role & Responsibilities", "Field Supervision"]
        return result

    # === TEAM MANAGEMENT / MORALE ===

    if "team management" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Team Management"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Team Management"]
        return result

    if "morale" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Team Morale"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Team Morale", "Employee Motivation"]
        return result

    if "building" in title_lower and "team" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Team Building"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Team Building", "Team Management"]
        return result

    if "motivat" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Employee Motivation"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Employee Motivation", "Team Morale"]
        return result

    # === CRITICAL LEADERSHIP ===

    if "critical leadership" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Critical Leadership Training"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Leadership Fundamentals", "Accountability & Ownership"]
        return result

    # === SUB-CONTRACTORS ===

    if "sub-contractor" in title_lower or "subcontractor" in title_lower:
        result["training_tab"] = "Tab 5 - Purchasing"
        result["tab_name"] = "Subcontractor Management"
        result["sops"] = ["9.4.610", "9.4.615"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Trade Coordination"]
        return result

    # === EMOTIONAL INTELLIGENCE ===

    if "emotional intelligence" in title_lower or " eq" in title_lower or "eq " in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Emotional Intelligence"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Emotional Intelligence"]
        result["skills"] = ["EQ Fundamentals"]
        if "leading" in title_lower:
            result["skills"].append("EQ Leadership")
        return result

    if "self-awareness" in title_lower or "self awareness" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Emotional Intelligence"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Emotional Intelligence"]
        result["skills"] = ["Self-Awareness", "EQ Fundamentals"]
        return result

    if "self-regulation" in title_lower or "self regulation" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Emotional Intelligence"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Emotional Intelligence"]
        result["skills"] = ["Self-Regulation", "EQ Fundamentals"]
        return result

    if "empathy" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Emotional Intelligence"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Emotional Intelligence"]
        result["skills"] = ["Empathy", "EQ Fundamentals"]
        return result

    if "social skills" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Emotional Intelligence"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Emotional Intelligence"]
        result["skills"] = ["Social Skills", "EQ Fundamentals"]
        return result

    # === PERFORMANCE MANAGEMENT ===

    if "performance review" in title_lower or "performance evaluation" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Performance Reviews"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Performance Management"]
        result["skills"] = ["Performance Reviews", "Conducting Evaluations"]
        return result

    if "performance management" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Performance Management"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Performance Management"]
        result["skills"] = ["Performance Management Fundamentals"]
        return result

    if "corrective counseling" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Corrective Counseling"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Performance Management"]
        result["skills"] = ["Corrective Counseling"]
        return result

    # === COMMUNICATION ===

    if "communication" in title_lower:
        result["training_tab"] = "Tab 16 - Communications"
        result["tab_name"] = "Effective Communication"
        result["sops"] = ["9.4.270", "9.4.275", "9.4.285"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Communication & Coaching"]
        result["skills"] = ["Effective Communication"]
        if "feedback" in title_lower:
            result["skills"].append("Setting Expectations")
        return result

    # === PROJECT PLANNING / PRODUCTIVITY ===

    if "project planning" in title_lower or "planning and execution" in title_lower:
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Project Planning"
        result["sops"] = ["9.2.100", "9.2.110"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Project Planning & Productivity"]
        result["skills"] = ["Project Planning", "Construction Project Lifecycle"]
        return result

    if "time management" in title_lower:
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Time Management"
        result["sops"] = ["9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Project Planning & Productivity"]
        result["skills"] = ["Time Management", "Productivity Improvement"]
        return result

    if "time wasters" in title_lower:
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Time Management"
        result["sops"] = ["9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Project Planning & Productivity"]
        result["skills"] = ["Time Management", "Productivity Improvement"]
        return result

    if "productivity" in title_lower:
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Productivity"
        result["sops"] = ["9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Project Planning & Productivity"]
        result["skills"] = ["Productivity Improvement"]
        return result

    # === JOB PLANS ===

    if "job plan" in title_lower:
        result["training_tab"] = "Tab 7 - Job Plans"
        result["tab_name"] = "Job Plans"
        result["sops"] = ["9.2.090", "9.2.100", "9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Construction Software", "Project Planning & Productivity"]
        result["skills"] = ["Job Plan Fundamentals", "Project Planning"]
        return result

    # === CONSTRUCTION SOFTWARE ===

    if "accubid" in title_lower:
        result["training_tab"] = "Tab 11 - Estimating"
        result["tab_name"] = "Accubid Software"
        result["sops"] = ["9.2.055", "9.2.090"]
        result["roles"] = ["EST", "PM"]
        result["skill_level_sets"] = ["Construction Software"]
        result["skills"] = ["Accubid Fundamentals"]
        return result

    if "viewpoint" in title_lower or "vista" in title_lower:
        result["training_tab"] = "Tab 8 - Change Orders"
        result["tab_name"] = "ViewPoint Software"
        result["sops"] = ["9.4.360", "9.4.365"]
        result["roles"] = ["PM", "CM", "EST"]
        result["skill_level_sets"] = ["Construction Software"]
        result["skills"] = ["ViewPoint Change Orders", "Vista"]
        return result

    if "procore" in title_lower:
        result["training_tab"] = "Tab 57 - ProCore"
        result["tab_name"] = "ProCore Software"
        result["sops"] = ["9.2.120", "9.4.195", "9.4.200", "9.4.205", "9.4.210", "9.4.215", "9.4.320"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Construction Software"]
        result["skills"] = ["ProCore Fundamentals"]
        return result

    # === DOCUMENTATION ===

    if "documentation" in title_lower or "reporting" in title_lower:
        result["training_tab"] = "Tab 22 - Document Management"
        result["tab_name"] = "Documentation"
        result["sops"] = ["9.4.190A", "9.4.195", "9.4.215"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Documentation & Compliance"]
        result["skills"] = ["Documentation Practices", "Documentation & Reporting"]
        return result

    if "document management" in title_lower:
        result["training_tab"] = "Tab 22 - Document Management"
        result["tab_name"] = "Document Management"
        result["sops"] = ["9.4.190A", "9.4.195"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Documentation & Compliance"]
        result["skills"] = ["Documentation Practices"]
        return result

    # === PROFESSIONAL DEVELOPMENT / PRESENTING ===

    if "presenting" in title_lower or "presentation" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Presentation Skills"
        result["sops"] = ["9.4.270", "9.4.345"]
        result["roles"] = ["FL", "GS", "PM", "BM"]
        result["skill_level_sets"] = ["Professional Development"]
        result["skills"] = ["Presentation Skills"]
        return result

    # === FIELD OPERATIONS ===

    if "jobsite efficiency" in title_lower:
        result["training_tab"] = "Tab 2 - Preplanning"
        result["tab_name"] = "Jobsite Efficiency"
        result["sops"] = ["9.4.705", "9.4.710"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Jobsite Efficiency"]
        return result

    if "coordinat" in title_lower and "trade" in title_lower:
        result["training_tab"] = "Tab 6 - Scheduling"
        result["tab_name"] = "Trade Coordination"
        result["sops"] = ["9.4.345"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Trade Coordination"]
        return result

    if "quality control" in title_lower or "quality" in title_lower:
        result["training_tab"] = "Tab 38 - Testing"
        result["tab_name"] = "Quality Control"
        result["sops"] = ["9.4.525", "9.4.530"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Quality Control"]
        return result

    if "manpower projection" in title_lower:
        result["training_tab"] = "Tab 19 - Manpower"
        result["tab_name"] = "Manpower Projections"
        result["sops"] = ["9.2.090", "9.4.340", "9.4.560", "9.4.690"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Manpower Projection"]
        return result

    if "closeout" in title_lower or "close out" in title_lower:
        result["training_tab"] = "Tab 18 - Job Closeout"
        result["tab_name"] = "Job Closeout"
        result["sops"] = ["9.6.005", "9.6.010", "9.6.015", "9.6.020"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Job Closeout"]
        return result

    if "tool" in title_lower and "technology" in title_lower:
        result["training_tab"] = "Tab 1 - Tool Control"
        result["tab_name"] = "Tools and Technology"
        result["sops"] = ["9.4.565", "9.4.700"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management", "Construction Software"]
        result["skills"] = ["Field Supervision"]
        return result

    # === WHAT MAKES A LEADER ===

    if "what makes a leader" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Leadership Fundamentals"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Leadership Fundamentals"]
        return result

    # === HUMAN INTERACTION ===

    if "human interaction" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "Leadership & Relationships"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management", "Communication & Coaching"]
        result["skills"] = ["Leadership Fundamentals", "Building Relationships"]
        return result

    # === NEW MANAGER ===

    if "new manager" in title_lower or "manager 101" in title_lower:
        result["training_tab"] = "Tab 51 - Leadership"
        result["tab_name"] = "New Manager Training"
        result["sops"] = ["9.2.010", "9.4.560"]
        result["roles"] = ["FL", "GS", "PM"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Leadership Fundamentals", "Team Management"]
        return result

    # === INSPECTIONS ===

    if "inspection" in title_lower and "workplace" in title_lower:
        result["training_tab"] = "Tab 40 - Safety"
        result["tab_name"] = "Workplace Inspections"
        result["sops"] = ["9.4.480"]
        result["roles"] = ["FL", "GS"]
        result["skill_level_sets"] = ["Safety", "Documentation & Compliance"]
        result["skills"] = ["Safety Inspections", "Pre-Task Safety Planning"]
        return result

    # === DEFAULT FALLBACK ===

    # If no match found, return minimal defaults
    if not result["training_tab"]:
        result["training_tab"] = "General Training"
        result["tab_name"] = "General"
        result["sops"] = []
        result["roles"] = ["FE", "FL", "GS"]
        result["skill_level_sets"] = ["Leadership & Field Management"]
        result["skills"] = ["Field Supervision"]

    return result


def get_role_names(role_codes):
    """Convert role codes to full names"""
    names = []
    for code in role_codes:
        if code in ROLE_SKILLS:
            names.append(ROLE_SKILLS[code]["name"])
    return names


def process_catalog(input_csv, output_xlsx):
    """Process catalog and build complete Course → Skills mapping"""

    # Read existing catalog
    courses = []
    with open(input_csv, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get('title'):
                courses.append(row)

    print(f"Processing {len(courses)} courses...")

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Course Skills Mapping"

    # Headers
    headers = [
        "Course ID",
        "Title",
        "Training Tab",
        "Tab Name",
        "SOPs",
        "Roles (Codes)",
        "Roles (Names)",
        "Management Directives",
        "Skill Level Sets",
        "Skills",
        "Proficiency Level",
        "Filename"
    ]

    # Style headers
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Process courses
    for idx, course in enumerate(courses, 2):
        title = course.get('title', '')

        # Get complete mapping
        mapping = get_course_mapping(title)

        # Get Management Directives for roles
        mds = []
        for role_code in mapping["roles"]:
            if role_code in ROLE_SKILLS:
                mds.append(ROLE_SKILLS[role_code]["md"])

        # Get role full names
        role_names = get_role_names(mapping["roles"])

        # Remove duplicates from skills while preserving order
        seen = set()
        unique_skills = []
        for skill in mapping["skills"]:
            if skill not in seen:
                seen.add(skill)
                unique_skills.append(skill)

        # Remove duplicates from skill level sets
        unique_sls = list(dict.fromkeys(mapping["skill_level_sets"]))

        # Write row
        row_data = [
            course.get('course_id', ''),
            title,
            mapping["training_tab"],
            mapping["tab_name"],
            "; ".join(mapping["sops"]),
            "; ".join(mapping["roles"]),
            "; ".join(role_names),
            "; ".join(list(dict.fromkeys(mds))),
            "; ".join(unique_sls),
            "; ".join(unique_skills),
            course.get('proficiency_level', ''),
            course.get('filename', '')
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    # Column widths
    column_widths = [15, 50, 20, 25, 40, 15, 50, 25, 40, 60, 12, 40]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    # Freeze header
    ws.freeze_panes = 'A2'

    # Save
    wb.save(output_xlsx)
    print(f"Created: {output_xlsx}")

    # Statistics
    all_skills = set()
    all_sops = set()
    all_roles = set()
    all_sls = set()

    for course in courses:
        mapping = get_course_mapping(course.get('title', ''))
        all_skills.update(mapping["skills"])
        all_sops.update(mapping["sops"])
        all_roles.update(mapping["roles"])
        all_sls.update(mapping["skill_level_sets"])

    print(f"\nStatistics:")
    print(f"  Total Courses: {len(courses)}")
    print(f"  Unique Skills: {len(all_skills)}")
    print(f"  Unique SOPs: {len(all_sops)}")
    print(f"  Unique Roles: {len(all_roles)}")
    print(f"  Skill Level Sets: {len(all_sls)}")

    return len(courses)


if __name__ == "__main__":
    input_file = r"C:\Users\tewing\Documents\Projects\GSL-Operations-Framework\deliverables\scorm-catalog\SCORM_COURSE_CATALOG.csv"
    output_file = r"C:\Users\tewing\Desktop\Claude Projects\SOPs For Review\SCORM_Course_Skills_Mapping.xlsx"

    process_catalog(input_file, output_file)
