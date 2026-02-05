"""
Create Learn365 Skills Matrix Excel
Matching structure of 'THIS SOP Matrix.xlsx'
6 Sheets: Coverage Matrix, RACI Matrix, Redline Recs, Conflict Matrix, Closeout Checklist, Role Alignment Matrix
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from datetime import datetime

# Create workbook
wb = openpyxl.Workbook()

# Styles
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

def auto_width(ws):
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[column].width = adjusted_width

# ============================================================
# Sheet 1: Coverage Matrix
# Columns: Skill Title, Skill Action, Responsible Role, Directive Section, Coverage Status, Gap Closure Recommendation, Category
# ============================================================
ws1 = wb.active
ws1.title = "Coverage Matrix"

headers = ["Skill Title", "Skill Action", "Responsible Role", "Directive Section", "Coverage Status", "Gap Closure Recommendation", "Category"]
ws1.append(headers)
style_header(ws1)

COVERAGE_DATA = [
    # Safety Skills
    ("Safety Awareness & Culture", "Understand safety culture principles", "All Employees", "MD 9.2, 9.3 (Safety Requirements)", "✔ Covered", "No change required.", "Safety"),
    ("Safety Awareness & Culture", "Complete safety orientation", "Field Employee", "MD 9.2 (New Hire Requirements)", "✔ Covered", "No change required.", "Safety"),
    ("Safety Awareness & Culture", "Conduct safety meetings", "Foreman", "MD 9.3 (Field Leadership)", "✔ Covered", "No change required.", "Safety"),
    ("Accident Prevention", "Identify accident causes", "All Field Roles", "MD 9.3 (Safety Management)", "✔ Covered", "No change required.", "Safety"),
    ("Accident Prevention", "Implement prevention measures", "Foreman", "MD 9.3 (Field Leadership)", "✔ Covered", "No change required.", "Safety"),
    ("Hazard Recognition", "Identify workplace hazards", "All Field Roles", "MD 9.3 (Safety Management)", "✔ Covered", "No change required.", "Safety"),
    ("Hazard Recognition", "Document hazards in JHA", "Foreman", "MD 9.3 (Pre-task Planning)", "✔ Covered", "No change required.", "Safety"),
    ("Pre-Task Safety Planning", "Complete JHA before work", "Foreman", "MD 9.3 (Pre-task Planning)", "✔ Covered", "No change required.", "Safety"),
    ("Pre-Task Safety Planning", "Review hazards with crew", "Foreman", "MD 9.3 (Field Leadership)", "✔ Covered", "No change required.", "Safety"),
    ("Safety Inspections", "Conduct daily inspections", "Foreman", "MD 9.3 (Safety Inspections)", "✔ Covered", "No change required.", "Safety"),
    ("Safety Inspections", "Document inspection findings", "Foreman", "MD 9.3 (Documentation)", "✔ Covered", "No change required.", "Safety"),
    ("Incident Investigation", "Investigate safety incidents", "Safety Coordinator", "MD 9.3 (Incident Response)", "✔ Covered", "No change required.", "Safety"),
    ("Incident Investigation", "Report incidents to management", "Foreman", "MD 9.3 (Reporting)", "✔ Covered", "No change required.", "Safety"),
    ("First Aid/CPR/AED", "Maintain first aid certification", "All Field Roles", "MD 9.2 (Certifications)", "✔ Covered", "No change required.", "Safety"),
    ("Emergency Response", "Know emergency procedures", "All Employees", "MD 9.2, 9.3 (Emergency Response)", "✔ Covered", "No change required.", "Safety"),
    ("Emergency Response", "Lead emergency response", "Foreman", "MD 9.3 (Field Leadership)", "✔ Covered", "No change required.", "Safety"),
    ("NFPA 70E Fundamentals", "Understand electrical safety standards", "Electrical Workers", "MD 9.3 (Electrical Safety)", "✔ Covered", "No change required.", "Safety"),
    ("NFPA 70E Fundamentals", "Apply NFPA 70E requirements", "Journeyman", "MD 9.3 (Qualified Person)", "✔ Covered", "No change required.", "Safety"),
    ("Arc Flash Awareness", "Recognize arc flash hazards", "Electrical Workers", "MD 9.3 (Arc Flash)", "✔ Covered", "No change required.", "Safety"),
    ("Arc Flash Awareness", "Select appropriate PPE", "Journeyman", "MD 9.3 (PPE Requirements)", "✔ Covered", "No change required.", "Safety"),
    ("Arc Flash Risk Assessment", "Perform incident energy analysis", "Safety Coordinator", "MD 9.3 (Risk Assessment)", "✔ Covered", "No change required.", "Safety"),
    ("LOTO Fundamentals", "Understand LOTO requirements", "Electrical Workers", "MD 9.3 (LOTO)", "✔ Covered", "No change required.", "Safety"),
    ("LOTO Fundamentals", "Execute LOTO procedures", "Journeyman", "MD 9.3 (Energy Control)", "✔ Covered", "No change required.", "Safety"),
    ("Simple LOTO Procedures", "Perform simple lockout", "Journeyman", "MD 9.3 (LOTO)", "✔ Covered", "No change required.", "Safety"),
    ("Complex LOTO Procedures", "Manage complex lockout", "Foreman", "MD 9.3 (Complex LOTO)", "✔ Covered", "No change required.", "Safety"),
    ("Fall Hazard Recognition", "Identify fall hazards", "All Field Roles", "MD 9.3 (Fall Protection)", "✔ Covered", "No change required.", "Safety"),
    ("Fall Hazard Recognition", "Select fall protection method", "Foreman", "MD 9.3 (Fall Protection)", "✔ Covered", "No change required.", "Safety"),
    ("Personal Fall Arrest Systems", "Inspect PFAS equipment", "All Field Roles", "MD 9.3 (Equipment Inspection)", "✔ Covered", "No change required.", "Safety"),
    ("Personal Fall Arrest Systems", "Use PFAS correctly", "All Field Roles", "MD 9.3 (Fall Protection)", "✔ Covered", "No change required.", "Safety"),
    ("Ladder Safety", "Select appropriate ladder", "All Field Roles", "MD 9.3 (Ladder Safety)", "✔ Covered", "No change required.", "Safety"),
    ("Ladder Safety", "Inspect ladder before use", "All Field Roles", "MD 9.3 (Equipment Inspection)", "✔ Covered", "No change required.", "Safety"),
    ("Scaffold Fundamentals", "Understand scaffold requirements", "Field Roles", "MD 9.3 (Scaffolding)", "✔ Covered", "No change required.", "Safety"),
    ("Scaffold Fundamentals", "Inspect scaffold daily", "Foreman", "MD 9.3 (Scaffold Inspection)", "✔ Covered", "No change required.", "Safety"),
    ("HazCom Program Fundamentals", "Understand HazCom requirements", "All Employees", "MD 9.2 (HazCom)", "✔ Covered", "No change required.", "Safety"),
    ("HazCom Program Fundamentals", "Access and read SDS", "All Employees", "MD 9.2 (SDS Access)", "✔ Covered", "No change required.", "Safety"),
    ("Confined Space Program", "Recognize confined spaces", "Field Roles", "MD 9.3 (Confined Space)", "✔ Covered", "No change required.", "Safety"),
    ("Confined Space Program", "Follow permit procedures", "Foreman", "MD 9.3 (Permit Required)", "✔ Covered", "No change required.", "Safety"),

    # Lean Construction Skills
    ("Lean Fundamentals", "Understand lean principles", "FL, GS, PM, BM", "MD 9.5 (Lean Construction)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Lean Fundamentals", "Apply lean thinking to work", "Foreman", "MD 9.3 (Continuous Improvement)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Lean History", "Understand lean origins", "FL, GS, PM, BM", "MD 9.5 (Training)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Eight Wastes Recognition", "Identify waste in processes", "FL, GS, PM, BM", "MD 9.5 (Waste Elimination)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Eight Wastes Recognition", "Eliminate identified waste", "Foreman", "MD 9.3 (Efficiency)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Eliminate Waste", "Implement waste reduction", "FL, GS, PM", "MD 9.5 (Productivity)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Focus on Flow", "Optimize work flow", "FL, GS, PM", "MD 9.5 (Scheduling)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Continuous Improvement", "Lead CI initiatives", "GS, PM, BM", "MD 9.5 (Continuous Improvement)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Continuous Improvement", "Participate in CI activities", "Foreman", "MD 9.3 (Improvement)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Pull Planning", "Facilitate pull planning sessions", "GS, PM", "MD 9.5 (Pull Planning)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Pull Planning", "Participate in pull planning", "Foreman", "MD 9.3 (Planning)", "✔ Covered", "No change required.", "Lean Construction"),
    ("Field Team Empowerment", "Empower crew decisions", "Foreman", "MD 9.3 (Empowerment)", "✔ Covered", "No change required.", "Lean Construction"),

    # Leadership & Field Management Skills
    ("Foreman Role & Responsibilities", "Understand foreman duties", "Foreman", "MD 9.3 (Foreman Role)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Foreman Role & Responsibilities", "Execute foreman responsibilities", "Foreman", "MD 9.3 (Responsibilities)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Field Supervision", "Supervise crew effectively", "FL, GS", "MD 9.3, 9.4 (Supervision)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Field Supervision", "Ensure quality workmanship", "Foreman", "MD 9.3 (Quality)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Crew Planning", "Plan daily crew assignments", "Foreman", "MD 9.3 (Planning)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Crew Planning", "Coordinate with other trades", "Foreman", "MD 9.3 (Coordination)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Trade Coordination", "Coordinate with GC and trades", "FL, GS", "MD 9.3, 9.4 (Coordination)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Jobsite Efficiency", "Maximize crew productivity", "Foreman", "MD 9.3 (Productivity)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Quality Control", "Inspect work quality", "FL, GS", "MD 9.3, 9.4 (Quality)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Tool & Equipment Management", "Manage tools and equipment", "Foreman", "MD 9.3 (Equipment)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Manpower Projection", "Project manpower needs", "GS, PM", "MD 9.4, 9.5 (Manpower)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Job Closeout", "Complete job closeout tasks", "FL, GS, PM", "MD 9.3, 9.4, 9.5 (Closeout)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Leadership Fundamentals", "Apply leadership principles", "FL, GS, PM, BM", "MD 9.3-9.6 (Leadership)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Values-Driven Leadership", "Lead with company values", "FL, GS, PM, BM", "MD 9.3-9.6 (Values)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Accountability & Ownership", "Take ownership of outcomes", "FL, GS, PM, BM", "MD 9.3-9.6 (Accountability)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Team Building", "Build effective teams", "FL, GS, PM, BM", "MD 9.3-9.6 (Team Development)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Team Management", "Manage team performance", "FL, GS, PM, BM", "MD 9.3-9.6 (Management)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Delegation & Empowerment", "Delegate effectively", "FL, GS, PM, BM", "MD 9.3-9.6 (Delegation)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Employee Motivation", "Motivate team members", "FL, GS, PM, BM", "MD 9.3-9.6 (Motivation)", "✔ Covered", "No change required.", "Leadership & Field Management"),
    ("Extreme Ownership", "Take extreme ownership", "FL, GS, PM, BM", "MD 9.3-9.6 (Ownership)", "✔ Covered", "No change required.", "Leadership & Field Management"),

    # Communication & Coaching Skills
    ("Effective Communication", "Communicate clearly", "All Roles", "MD 9.2-9.6 (Communication)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Effective Communication", "Provide constructive feedback", "FL, GS, PM, BM", "MD 9.3-9.6 (Feedback)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Building Relationships", "Build professional relationships", "All Roles", "MD 9.2-9.6 (Relationships)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Conducting One-on-Ones", "Conduct effective one-on-ones", "FL, GS, PM, BM", "MD 9.3-9.6 (Employee Development)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Mentoring & Coaching", "Mentor and coach employees", "FL, GS, PM, BM", "MD 9.3-9.6 (Development)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Conflict Resolution", "Resolve workplace conflicts", "FL, GS, PM, BM", "MD 9.3-9.6 (Conflict)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Handling Resistance", "Address resistance to change", "FL, GS, PM, BM", "MD 9.3-9.6 (Change Management)", "✔ Covered", "No change required.", "Communication & Coaching"),
    ("Setting Expectations", "Set clear expectations", "FL, GS, PM, BM", "MD 9.3-9.6 (Expectations)", "✔ Covered", "No change required.", "Communication & Coaching"),

    # Emotional Intelligence Skills
    ("EQ Fundamentals", "Understand emotional intelligence", "All Roles", "MD 9.2-9.6 (EQ)", "✔ Covered", "No change required.", "Emotional Intelligence"),
    ("Self-Awareness", "Develop self-awareness", "All Roles", "MD 9.2-9.6 (Self-Development)", "✔ Covered", "No change required.", "Emotional Intelligence"),
    ("Self-Regulation", "Manage emotions effectively", "All Roles", "MD 9.2-9.6 (Self-Management)", "✔ Covered", "No change required.", "Emotional Intelligence"),
    ("Empathy", "Demonstrate empathy", "All Roles", "MD 9.2-9.6 (Relationships)", "✔ Covered", "No change required.", "Emotional Intelligence"),
    ("Social Skills", "Apply social skills", "All Roles", "MD 9.2-9.6 (Communication)", "✔ Covered", "No change required.", "Emotional Intelligence"),
    ("EQ Leadership", "Lead with EQ", "FL, GS, PM, BM", "MD 9.3-9.6 (Leadership)", "✔ Covered", "No change required.", "Emotional Intelligence"),
    ("Trust Building", "Build trust with team", "FL, GS, PM, BM", "MD 9.3-9.6 (Trust)", "✔ Covered", "No change required.", "Emotional Intelligence"),

    # Performance Management Skills
    ("Performance Management Fundamentals", "Understand PM principles", "FL, GS, PM, BM", "MD 9.3-9.6 (Performance)", "✔ Covered", "No change required.", "Performance Management"),
    ("Conducting Evaluations", "Conduct employee evaluations", "FL, GS, PM, BM", "MD 9.3-9.6 (Evaluations)", "✔ Covered", "No change required.", "Performance Management"),
    ("Performance Reviews", "Deliver performance reviews", "FL, GS, PM, BM", "MD 9.3-9.6 (Reviews)", "✔ Covered", "No change required.", "Performance Management"),
    ("Corrective Counseling", "Conduct corrective counseling", "FL, GS, PM, BM", "MD 9.3-9.6 (Discipline)", "✔ Covered", "No change required.", "Performance Management"),
    ("Escalation & Termination", "Handle escalation/termination", "GS, PM, BM", "MD 9.4-9.6 (Termination)", "✔ Covered", "No change required.", "Performance Management"),
    ("Career Development", "Support career development", "FL, GS, PM, BM", "MD 9.3-9.6 (Development)", "✔ Covered", "No change required.", "Performance Management"),
    ("Identifying Strengths", "Identify employee strengths", "FL, GS, PM, BM", "MD 9.3-9.6 (Development)", "✔ Covered", "No change required.", "Performance Management"),

    # Construction Software Skills
    ("Accubid Fundamentals", "Use Accubid for estimating", "Estimator", "MD 9.7 (Estimating Tools)", "✔ Covered", "No change required.", "Construction Software"),
    ("Accubid Data Export", "Export data from Accubid", "Estimator, PM", "MD 9.7, 9.5 (Data Management)", "✔ Covered", "No change required.", "Construction Software"),
    ("ViewPoint Change Orders", "Process change orders", "PM, CM", "MD 9.5, 9.14 (Change Orders)", "✔ Covered", "No change required.", "Construction Software"),
    ("ProCore Fundamentals", "Use ProCore basics", "PM, GS", "MD 9.4, 9.5 (Project Software)", "✔ Covered", "No change required.", "Construction Software"),
    ("ProCore Job Setup", "Set up jobs in ProCore", "PM", "MD 9.5 (Job Setup)", "✔ Covered", "No change required.", "Construction Software"),
    ("Job Plan Fundamentals", "Understand job plan theory", "FL, GS, PM", "MD 9.3-9.5 (Job Planning)", "✔ Covered", "No change required.", "Construction Software"),
    ("Job Plan Setup", "Set up job plans", "PM, GS", "MD 9.4, 9.5 (Planning)", "✔ Covered", "No change required.", "Construction Software"),
    ("Job Plan Maintenance", "Maintain job plans", "FL, GS, PM", "MD 9.3-9.5 (Plan Updates)", "✔ Covered", "No change required.", "Construction Software"),
    ("Excel Skills", "Use Excel effectively", "Office Roles", "MD 9.5-9.7 (Software)", "✔ Covered", "No change required.", "Construction Software"),
    ("Vista", "Use Vista system", "PM, Accounting", "MD 9.5 (Financial Systems)", "✔ Covered", "No change required.", "Construction Software"),
    ("Teams", "Use Microsoft Teams", "All Office", "MD 9.2-9.6 (Communication)", "✔ Covered", "No change required.", "Construction Software"),
    ("Outlook", "Use Outlook effectively", "All Office", "MD 9.2-9.6 (Communication)", "✔ Covered", "No change required.", "Construction Software"),

    # Project Planning & Productivity Skills
    ("Project Planning", "Plan project execution", "GS, PM", "MD 9.4, 9.5 (Planning)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Job Plan Analysis", "Analyze job plan data", "FL, GS, PM", "MD 9.3-9.5 (Analysis)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Scheduling Techniques", "Apply scheduling techniques", "GS, PM", "MD 9.4, 9.5 (Scheduling)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Productivity Analysis", "Analyze productivity metrics", "FL, GS, PM", "MD 9.3-9.5 (Productivity)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Productivity Improvement", "Improve productivity", "FL, GS, PM", "MD 9.3-9.5 (Improvement)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Time Management", "Manage time effectively", "All Roles", "MD 9.2-9.6 (Efficiency)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Construction Project Lifecycle", "Understand project phases", "FL, GS, PM", "MD 9.3-9.5 (Project Phases)", "✔ Covered", "No change required.", "Project Planning & Productivity"),
    ("Process Improvement", "Improve work processes", "FL, GS, PM, BM", "MD 9.3-9.6 (CI)", "✔ Covered", "No change required.", "Project Planning & Productivity"),

    # Documentation & Compliance Skills
    ("Documentation Practices", "Follow documentation standards", "All Roles", "MD 9.2-9.6 (Documentation)", "✔ Covered", "No change required.", "Documentation & Compliance"),
    ("Documentation & Reporting", "Complete required reports", "FL, GS, PM", "MD 9.3-9.5 (Reporting)", "✔ Covered", "No change required.", "Documentation & Compliance"),
    ("Regulatory Inspections", "Manage regulatory inspections", "GS, PM", "MD 9.4, 9.5 (Inspections)", "✔ Covered", "No change required.", "Documentation & Compliance"),
    ("Documentation & Compliance", "Ensure compliance documentation", "PM, CM", "MD 9.5, 9.14 (Compliance)", "✔ Covered", "No change required.", "Documentation & Compliance"),

    # Professional Development Skills
    ("Presentation Skills", "Deliver effective presentations", "GS, PM, BM", "MD 9.4-9.6 (Presentations)", "✔ Covered", "No change required.", "Professional Development"),
    ("Strategic Thinking", "Think strategically", "GS, PM, BM", "MD 9.4-9.6 (Strategy)", "✔ Covered", "No change required.", "Professional Development"),
    ("Industry Leadership", "Lead in industry", "BM", "MD 9.6 (Industry)", "✔ Covered", "No change required.", "Professional Development"),
]

for row in COVERAGE_DATA:
    ws1.append(row)

auto_width(ws1)

# ============================================================
# Sheet 2: RACI Matrix
# Same as Coverage Matrix + RACI column
# ============================================================
ws2 = wb.create_sheet("RACI Matrix")

headers = ["Skill Title", "Skill Action", "Responsible Role", "Directive Section", "Coverage Status", "Gap Closure Recommendation", "Category", "RACI"]
ws2.append(headers)
style_header(ws2)

# Add RACI assignments
for row in COVERAGE_DATA:
    # Determine RACI based on role
    role = row[2]
    if "All" in role:
        raci = "R"
    elif "PM" in role or "Project Manager" in role:
        raci = "A"
    elif "Foreman" in role or "FL" in role:
        raci = "R"
    elif "GS" in role or "Superintendent" in role:
        raci = "A"
    elif "BM" in role:
        raci = "A"
    else:
        raci = "R"

    ws2.append(list(row) + [raci])

auto_width(ws2)

# ============================================================
# Sheet 3: Redline Recs
# Coverage Matrix + Redline Proposal for gaps
# ============================================================
ws3 = wb.create_sheet("Redline Recs")

headers = ["Skill Title", "Skill Action", "Responsible Role", "Directive Section", "Coverage Status", "Gap Closure Recommendation", "Category", "Redline Proposal"]
ws3.append(headers)
style_header(ws3)

REDLINE_DATA = [
    ("BIM Coordination", "Coordinate using BIM models", "PM, GS", "MD 9.5 (BIM)", "⚠ Gap – Limited training", "Add BIM basics training module", "Construction Software", "Create 'BIM for Field Leaders' course covering model navigation, clash detection review, and coordination workflows."),
    ("Commissioning Procedures", "Execute commissioning tasks", "GS, PM", "MD 9.5 (Commissioning)", "⚠ Gap – Minimal training", "Expand commissioning training", "Project Planning & Productivity", "Develop commissioning training covering startup procedures, testing protocols, and documentation requirements."),
    ("Large Feeder Wire Management", "Manage large feeder installation", "Foreman", "MD 9.3 (Wire Management)", "✔ Covered (Gap Closed)", "Training module created", "Safety", "Tab 41 training module created covering procurement, handling, installation, and safety."),
    ("Prefabrication Operations", "Manage prefab operations", "GS, PM", "MD 9.4, 9.5 (Prefab)", "✔ Covered (Gap Closed)", "Training module expanded", "Project Planning & Productivity", "Tab 42 training module expanded with comprehensive prefab procedures."),
    ("Contract Review", "Review contract requirements", "PM, CM", "MD 9.5, 9.14 (Contracts)", "✔ Covered (Gap Closed)", "Training module created", "Documentation & Compliance", "Tab 2 training module created covering contract review and risk assessment."),
    ("RFI Management", "Manage RFI process", "PM, GS", "MD 9.5 (RFIs)", "✔ Covered (Gap Closed)", "Training module created", "Documentation & Compliance", "Tab 22 training module created covering complete RFI lifecycle."),
    ("Safety Coordinator Certification", "Obtain safety coordinator cert", "Safety Coordinator", "MD 9.3 (Certifications)", "⚠ Gap – Modules 1.0 & 2.0 missing", "Obtain missing SCORM modules", "Safety", "Request Safety Coordinator modules 1.0 and 2.0 from training vendor."),
    ("Coach K Leadership Series", "Complete Coach K training", "FL, GS, PM, BM", "MD 9.3-9.6 (Leadership)", "⚠ Gap – Modules 2-6 missing", "Obtain missing SCORM modules", "Leadership & Field Management", "Request Coach K modules 2-6 from training vendor or identify replacement content."),
    ("Scale Set Levels", "Configure proficiency levels", "LMS Admin", "N/A (System Config)", "⚠ Gap – API limitation", "Configure in Learn365 UI", "System Configuration", "Configure GSL Proficiency Scale levels (Awareness, Basic, Intermediate, Advanced, Expert) in Learn365 admin UI."),
    ("Target Skills", "Create role target skills", "LMS Admin", "N/A (System Config)", "⚠ Gap – Requires UI config", "Configure in Learn365 UI", "System Configuration", "Create Target Skills for each role (FE, FL, GS, PM, BM, EST, CM, PUR) with skill requirements and user field conditions."),
]

for row in REDLINE_DATA:
    ws3.append(row)

auto_width(ws3)

# ============================================================
# Sheet 4: Conflict Matrix
# ============================================================
ws4 = wb.create_sheet("Conflict Matrix")

headers = ["Skill Title", "Conflict Description", "Conflict Type", "Resolution Recommendation"]
ws4.append(headers)
style_header(ws4)

CONFLICT_DATA = [
    ("Skills vs Competencies", "Learn365 has separate Skills API and Competencies API with different purposes", "Architecture / System Design", "Use Skills API for role requirements (Target Skills); Use Competencies API for course completion rewards. Document this separation for LMS admins."),
    ("Duplicate Skill Names", "67 skills already existed with same titles when creating new framework", "Data Duplication", "Migrated existing skills to new categories rather than creating duplicates. Updated scaleSetId to GSL Proficiency Scale."),
    ("Scale Set Levels Not Persisting", "API creates scale sets but levels don't persist properly", "API Limitation", "Configure scale set levels manually in Learn365 admin UI. Document as known limitation."),
    ("Old Categories Cannot Be Deleted", "100 old skill categories are linked to courses and cannot be deleted", "Data Dependency", "Created 10 new clean categories alongside old ones. Old categories will coexist until courses are remapped."),
    ("SCORM Course Duplicates", "71 duplicate SCORM courses identified in LMS", "Content Duplication", "Documented in DUPLICATE_COURSES.csv. Recommend removal after confirming no active enrollments."),
    ("Skill Level Terminology", "Old framework used varied proficiency terms; new framework standardizes", "Terminology Inconsistency", "Standardized to 5-level GSL Proficiency Scale: Awareness (1), Basic (2), Intermediate (3), Advanced (4), Expert (5)."),
    ("Role Name Variations", "Different documents use different role names (e.g., FL vs Foreman)", "Naming Inconsistency", "Standardized role codes: FE (Field Employee), FL (Foreman), GS (General Superintendent), PM (Project Manager), BM (Branch Manager), EST (Estimator), CM (Contract Manager), PUR (Purchasing)."),
    ("Target Skill Rules vs Training Plans", "Confusion between Target Skills (requirements) and Training Plans (assignments)", "Conceptual Confusion", "Document clearly: Target Skills define what skills a role needs; Training Plans define what courses to assign. They work together but serve different purposes."),
]

for row in CONFLICT_DATA:
    ws4.append(row)

auto_width(ws4)

# ============================================================
# Sheet 5: Closeout Checklist (Role Training Requirements)
# ============================================================
ws5 = wb.create_sheet("Closeout Checklist")

headers = ["Role", "Required Training Task", "Directive Section", "Coverage Status"]
ws5.append(headers)
style_header(ws5)

CLOSEOUT_DATA = [
    ("Field Employee (FE)", "Complete New Hire Safety Orientation", "MD 9.2 (New Hire)", "✔ Covered"),
    ("Field Employee (FE)", "Complete What Is Emotional Intelligence course", "MD 9.2 (EQ)", "✔ Covered"),
    ("Field Employee (FE)", "Complete Lean Fundamentals (Awareness)", "MD 9.2 (Lean)", "✔ Covered"),
    ("Field Employee (FE)", "Complete Communication Basics", "MD 9.2 (Communication)", "✔ Covered"),
    ("Field Leadership (FL)", "Complete Introduction to Field Leadership (1-4)", "MD 9.3 (Leadership)", "✔ Covered"),
    ("Field Leadership (FL)", "Complete Roles and Responsibilities of a Foreman", "MD 9.3 (Foreman)", "✔ Covered"),
    ("Field Leadership (FL)", "Complete LOTO Series", "MD 9.3 (LOTO)", "✔ Covered"),
    ("Field Leadership (FL)", "Complete Fall Protection Training", "MD 9.3 (Fall Protection)", "✔ Covered"),
    ("Field Leadership (FL)", "Complete Values Driven Leadership (1-10)", "MD 9.3 (Leadership)", "✔ Covered"),
    ("Field Leadership (FL)", "Complete Corrective Counseling Training", "MD 9.3 (Performance)", "✔ Covered"),
    ("General Superintendent (GS)", "Complete All Foreman Requirements", "MD 9.4 (Prerequisites)", "✔ Covered"),
    ("General Superintendent (GS)", "Complete Safety Coordinator Series (3.0-11.0)", "MD 9.4 (Safety)", "⚠ Partial - Modules 1.0, 2.0 missing"),
    ("General Superintendent (GS)", "Complete Critical Leadership Series", "MD 9.4 (Leadership)", "✔ Covered"),
    ("General Superintendent (GS)", "Complete Full Values Driven Leadership (1-17)", "MD 9.4 (Leadership)", "✔ Covered"),
    ("General Superintendent (GS)", "Complete All EQ Courses", "MD 9.4 (EQ)", "✔ Covered"),
    ("Project Manager (PM)", "Complete ViewPoint Training", "MD 9.5 (Software)", "✔ Covered"),
    ("Project Manager (PM)", "Complete Accubid Overview", "MD 9.5 (Estimating)", "✔ Covered"),
    ("Project Manager (PM)", "Complete Job Plan Theory", "MD 9.5 (Planning)", "✔ Covered"),
    ("Project Manager (PM)", "Complete Values Driven Leadership (1-10)", "MD 9.5 (Leadership)", "✔ Covered"),
    ("Branch Manager (BM)", "Complete Full Lean Series", "MD 9.6 (Lean)", "✔ Covered"),
    ("Branch Manager (BM)", "Complete All Values Driven Leadership", "MD 9.6 (Leadership)", "✔ Covered"),
    ("Branch Manager (BM)", "Complete Art of Presenting Series", "MD 9.6 (Presentations)", "✔ Covered"),
    ("Estimator (EST)", "Complete Full Accubid Training", "MD 9.7 (Estimating)", "✔ Covered"),
    ("Estimator (EST)", "Complete ViewPoint Overview", "MD 9.7 (Software)", "✔ Covered"),
    ("Contract Manager (CM)", "Complete ViewPoint Change Orders", "MD 9.14 (Change Orders)", "✔ Covered"),
    ("Contract Manager (CM)", "Complete Documentation Training", "MD 9.14 (Documentation)", "✔ Covered"),
    ("Purchasing (PUR)", "Complete ViewPoint Basics", "MD 9.16 (Software)", "✔ Covered"),
    ("Purchasing (PUR)", "Complete Documentation Training", "MD 9.16 (Documentation)", "✔ Covered"),
]

for row in CLOSEOUT_DATA:
    ws5.append(row)

auto_width(ws5)

# ============================================================
# Sheet 6: Role Alignment Matrix
# Skills x Roles with X marks
# ============================================================
ws6 = wb.create_sheet("Role Alignment Matrix")

# Headers: Skill Title, Category, then role columns
role_headers = ["Skill Title", "Category", "Apprentice", "Journeyman", "Lead Journeyman", "Foreman",
                "General Superintendent", "Project Manager", "Branch Manager", "Estimator",
                "Contract Manager", "Purchasing", "Safety Coordinator", "Accounting", "All Employees"]
ws6.append(role_headers)
style_header(ws6)

# Skills with role assignments (X = required, blank = not required)
ALIGNMENT_DATA = [
    # Safety
    ("Safety Awareness & Culture", "Safety", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X"),
    ("Accident Prevention", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Hazard Recognition", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Pre-Task Safety Planning", "Safety", "", "", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Safety Inspections", "Safety", "", "", "", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Incident Investigation", "Safety", "", "", "", "X", "X", "", "", "", "", "", "X", "", ""),
    ("First Aid/CPR/AED", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Emergency Response", "Safety", "X", "X", "X", "X", "X", "X", "X", "", "", "", "X", "", "X"),
    ("NFPA 70E Fundamentals", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Arc Flash Awareness", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("LOTO Fundamentals", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Fall Hazard Recognition", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "X", "", ""),
    ("Ladder Safety", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "", "", ""),
    ("Scaffold Fundamentals", "Safety", "X", "X", "X", "X", "X", "", "", "", "", "", "", "", ""),
    ("HazCom Program Fundamentals", "Safety", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X", "X"),

    # Lean Construction
    ("Lean Fundamentals", "Lean Construction", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Eight Wastes Recognition", "Lean Construction", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Continuous Improvement", "Lean Construction", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Pull Planning", "Lean Construction", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Field Team Empowerment", "Lean Construction", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),

    # Leadership & Field Management
    ("Foreman Role & Responsibilities", "Leadership & Field Management", "", "", "", "X", "X", "", "", "", "", "", "", "", ""),
    ("Field Supervision", "Leadership & Field Management", "", "", "", "X", "X", "", "", "", "", "", "", "", ""),
    ("Crew Planning", "Leadership & Field Management", "", "", "", "X", "X", "", "", "", "", "", "", "", ""),
    ("Trade Coordination", "Leadership & Field Management", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),
    ("Quality Control", "Leadership & Field Management", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),
    ("Leadership Fundamentals", "Leadership & Field Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Values-Driven Leadership", "Leadership & Field Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Accountability & Ownership", "Leadership & Field Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Team Building", "Leadership & Field Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Extreme Ownership", "Leadership & Field Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),

    # Communication & Coaching
    ("Effective Communication", "Communication & Coaching", "", "", "", "X", "X", "X", "X", "", "X", "", "", "", ""),
    ("Building Relationships", "Communication & Coaching", "", "", "", "X", "X", "X", "X", "", "X", "", "", "", ""),
    ("Conflict Resolution", "Communication & Coaching", "", "", "", "X", "X", "X", "X", "", "X", "", "", "", ""),
    ("Mentoring & Coaching", "Communication & Coaching", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),

    # Emotional Intelligence
    ("EQ Fundamentals", "Emotional Intelligence", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Self-Awareness", "Emotional Intelligence", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Empathy", "Emotional Intelligence", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("EQ Leadership", "Emotional Intelligence", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),

    # Performance Management
    ("Performance Management Fundamentals", "Performance Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Conducting Evaluations", "Performance Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Performance Reviews", "Performance Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Corrective Counseling", "Performance Management", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),

    # Construction Software
    ("Accubid Fundamentals", "Construction Software", "", "", "", "", "", "X", "", "X", "", "", "", "", ""),
    ("ViewPoint Change Orders", "Construction Software", "", "", "", "", "", "X", "", "", "X", "", "", "", ""),
    ("ProCore Fundamentals", "Construction Software", "", "", "", "", "X", "X", "", "", "", "", "", "", ""),
    ("Job Plan Fundamentals", "Construction Software", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),
    ("Excel Skills", "Construction Software", "", "", "", "", "", "X", "X", "X", "X", "X", "", "X", ""),

    # Project Planning & Productivity
    ("Project Planning", "Project Planning & Productivity", "", "", "", "X", "X", "X", "X", "", "", "", "", "", ""),
    ("Scheduling Techniques", "Project Planning & Productivity", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),
    ("Productivity Analysis", "Project Planning & Productivity", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),
    ("Time Management", "Project Planning & Productivity", "", "", "", "X", "X", "X", "X", "X", "X", "X", "", "", ""),

    # Documentation & Compliance
    ("Documentation Practices", "Documentation & Compliance", "", "", "", "X", "X", "X", "", "", "X", "X", "", "", ""),
    ("Regulatory Inspections", "Documentation & Compliance", "", "", "", "X", "X", "X", "", "", "", "", "", "", ""),

    # Professional Development
    ("Presentation Skills", "Professional Development", "", "", "", "", "X", "X", "X", "", "", "", "", "", ""),
    ("Strategic Thinking", "Professional Development", "", "", "", "", "X", "X", "X", "", "", "", "", "", ""),
]

for row in ALIGNMENT_DATA:
    ws6.append(row)

auto_width(ws6)

# ============================================================
# Save workbook
# ============================================================
output_path = r"C:\Users\tewing\Desktop\Claude Projects\SOPs For Review\Learn365_Skills_Matrix.xlsx"
wb.save(output_path)
print(f"Created: {output_path}")
print(f"Sheets: {wb.sheetnames}")
print(f"Row counts:")
for sheet in wb.sheetnames:
    print(f"  {sheet}: {wb[sheet].max_row} rows")
