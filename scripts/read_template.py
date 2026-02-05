import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

doc = docx.Document(r"C:\Users\tewing\Desktop\Claude Projects\GSL_SOP_Template.docx")

print("=" * 70)
print("GSL SOP MASTER TEMPLATE STRUCTURE")
print("=" * 70)

# Show document styles
print("\n--- STYLES IN USE ---")
styles_used = set()
for para in doc.paragraphs:
    if para.style:
        styles_used.add(para.style.name)
for style in sorted(styles_used):
    print(f"  - {style}")

print("\n--- PARAGRAPH CONTENT ---")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text or para.style.name != "Normal":
        style_name = para.style.name if para.style else "None"
        # Check for bold
        is_bold = any(run.bold for run in para.runs if run.text.strip())
        bold_indicator = "[BOLD]" if is_bold else ""

        # Check alignment
        alignment = para.alignment.name if para.alignment else "LEFT"

        print(f"[{i:3}] ({style_name:20}) {bold_indicator:7} {text[:80]}{'...' if len(text) > 80 else ''}")

# Check for tables
print("\n--- TABLES ---")
for t_idx, table in enumerate(doc.tables):
    print(f"\nTable {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols")
    for r_idx, row in enumerate(table.rows):
        row_text = [cell.text.strip()[:25] for cell in row.cells]
        print(f"  Row {r_idx}: {row_text}")

# Check document properties
print("\n--- DOCUMENT SECTIONS ---")
for i, section in enumerate(doc.sections):
    print(f"Section {i}: margin_top={section.top_margin.inches:.2f}in, margin_bottom={section.bottom_margin.inches:.2f}in")
    print(f"  header: {section.header.paragraphs[0].text if section.header.paragraphs else 'None'}")
