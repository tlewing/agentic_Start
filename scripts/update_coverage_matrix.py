# -*- coding: utf-8 -*-
"""
Update Coverage Matrix from Revised SOPs
Extracts action items from revised SOPs and updates the Excel Coverage Matrix
"""

import os
import re
from docx import Document
import openpyxl

# Paths
REVISED_FOLDER = r"C:\Users\tewing\Desktop\Claude Projects\SOP_Revisions\Revised"
EXCEL_FILE = r"C:\Users\tewing\Desktop\Claude Projects\Key_SOP_Matrix_RACI_Updated.xlsx"

def extract_sop_data(filepath):
    """Extract SOP data from a revised document"""
    doc = Document(filepath)

    sop_data = {
        "sop_id": "",
        "title": "",
        "department": "",
        "actions": []
    }

    current_section = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Extract SOP ID and Title from the first line
        if text.startswith("SOP:"):
            match = re.match(r"SOP:\s*(\d+\.\d+\.\d+)\s*[-\u2013]\s*(.+)", text)
            if match:
                sop_data["sop_id"] = match.group(1)
                sop_data["title"] = match.group(2).strip()
            else:
                # Try simpler pattern
                parts = text.replace("SOP:", "").strip().split(" - ", 1)
                if len(parts) >= 2:
                    sop_data["sop_id"] = parts[0].strip()
                    sop_data["title"] = parts[1].strip()

        # Extract Department
        elif text.startswith("Department:"):
            sop_data["department"] = text.replace("Department:", "").strip()

        # Track section headers (Bold text that matches known sections)
        elif text in ["Purpose", "Scope", "Roles & Responsibilities", "Requirements", "Procedure", "Appendix"]:
            current_section = text

        # Extract procedure steps - look for "Step X - " patterns
        elif current_section == "Procedure":
            step_match = re.match(r"Step\s*(\d+)\s*[-\u2013]\s*(.+)", text)
            if step_match:
                step_name = step_match.group(2).strip()
                sop_data["actions"].append(step_name)

    return sop_data


def determine_responsible_role(action_text, sop_id):
    """Determine the responsible role for an action based on text analysis"""
    action_lower = action_text.lower()

    # Estimator indicators
    if any(kw in action_lower for kw in ["estimator", "estimate", "estimating"]):
        return "Estimator"

    # Safety indicators
    if any(kw in action_lower for kw in ["safety", "hazard", "emergency", "ppe"]):
        return "Safety Coordinator"

    # Foreman indicators
    if any(kw in action_lower for kw in ["foreman", "field", "constructability", "installation", "execute", "physical setup", "crew", "work area", "staging"]):
        return "Foreman / Field Supervisor"

    # General Superintendent indicators
    if any(kw in action_lower for kw in ["superintendent", "validate", "approve", "resource", "major project"]):
        return "General Superintendent"

    # Default based on SOP series
    if sop_id.startswith("9.2."):
        return "Project Manager"
    elif sop_id.startswith("9.3."):
        return "Foreman / Field Supervisor"
    elif sop_id.startswith("9.4."):
        return "Project Manager"

    return "Project Manager"


def main():
    print("=" * 60)
    print("UPDATING COVERAGE MATRIX FROM REVISED SOPs")
    print("=" * 60)

    # Read all revised SOPs
    sop_data_list = []
    for filename in sorted(os.listdir(REVISED_FOLDER)):
        if filename.endswith(".docx"):
            filepath = os.path.join(REVISED_FOLDER, filename)
            print(f"  Reading: {filename}")
            sop_data = extract_sop_data(filepath)
            if sop_data["sop_id"]:
                sop_data_list.append(sop_data)
                print(f"    Found {len(sop_data['actions'])} actions")

    print(f"\nExtracted data from {len(sop_data_list)} SOPs")

    # Build the new coverage matrix data
    coverage_data = []
    for sop in sop_data_list:
        sop_id = sop["sop_id"]
        sop_title = f"{sop_id} - {sop['title']}"

        for action_text in sop["actions"]:
            responsible_role = determine_responsible_role(action_text, sop_id)

            coverage_data.append({
                "SOP_ID": sop_id,
                "SOP Title": sop_title,
                "SOP Action": action_text,
                "Responsible Role": responsible_role,
                "Department": sop["department"]
            })

    print(f"Generated {len(coverage_data)} action items")

    # Load and update the Excel file
    print(f"\nUpdating Excel file: {EXCEL_FILE}")
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Coverage Matrix"]

    # Clear existing data (keep header row)
    for row in range(ws.max_row, 1, -1):
        ws.delete_rows(row)

    # Get column indices from header
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}

    # Add new data
    for i, item in enumerate(coverage_data, 2):
        if "SOP_ID" in headers:
            ws.cell(row=i, column=headers["SOP_ID"], value=item["SOP_ID"])
        if "SOP Title" in headers:
            ws.cell(row=i, column=headers["SOP Title"], value=item["SOP Title"])
        if "SOP Action" in headers:
            ws.cell(row=i, column=headers["SOP Action"], value=item["SOP Action"])
        if "Responsible Role" in headers:
            ws.cell(row=i, column=headers["Responsible Role"], value=item["Responsible Role"])

    # Save the workbook
    wb.save(EXCEL_FILE)
    print(f"Saved updated Coverage Matrix with {len(coverage_data)} rows")

    print("\n" + "=" * 60)
    print("COVERAGE MATRIX UPDATE COMPLETE")
    print("=" * 60)

    # Show sample of the data
    print("\nSample data (first 15 rows):")
    print(f"{'SOP_ID':<10} | {'SOP Title':<40} | {'Action':<40}")
    print("-" * 95)
    for item in coverage_data[:15]:
        title = item["SOP Title"][:38] + ".." if len(item["SOP Title"]) > 40 else item["SOP Title"]
        action = item["SOP Action"][:38] + ".." if len(item["SOP Action"]) > 40 else item["SOP Action"]
        print(f"{item['SOP_ID']:<10} | {title:<40} | {action:<40}")


if __name__ == "__main__":
    main()
